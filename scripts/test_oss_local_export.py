from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import stat
import shutil
import threading
import time
from collections import namedtuple
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterable
from urllib.error import HTTPError
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

import scripts.oss_local_export as exporter
from scripts.oss_local_export import (
    MAX_PENDING_DOWNLOADS,
    ExportReport,
    InsufficientDiskError,
    ManifestItem,
    download_manifest,
    run_export,
    run_export_stream,
    sanitize_error,
)


IMAGE_BYTES = b"verified-fake-oss-image"
IMAGE_SHA256 = hashlib.sha256(IMAGE_BYTES).hexdigest()
SCHEMA = "module-manager-oss-export/v1"


class FakeResponse(io.BytesIO):
    def __enter__(self) -> "FakeResponse":
        return self

    def __exit__(self, *_args: object) -> None:
        self.close()


class FakeHttpClient:
    def __init__(self, outcomes: Iterable[bytes | BaseException]) -> None:
        self.outcomes = iter(outcomes)
        self.sleeps: list[float] = []
        self.urls: list[str] = []

    def open(self, url: str, *, timeout: int) -> FakeResponse:
        assert timeout == 30
        self.urls.append(url)
        outcome = next(self.outcomes)
        if isinstance(outcome, BaseException):
            raise outcome
        return FakeResponse(outcome)

    def sleep(self, delay: float) -> None:
        self.sleeps.append(delay)


def valid_row(**overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "schema": SCHEMA,
        "kind": "item",
        "group_id": "group-001",
        "photo_id": "000123456789",
        "category": "现场照片",
        "relative_path": "T01/group-001/现场照片/photo-001.jpg",
        "storage_bucket": "photos-read-only",
        "storage_key": "photos/aa/object.jpg",
        "sha256": IMAGE_SHA256,
        "byte_size": len(IMAGE_BYTES),
        "content_type": "image/jpeg",
        "download_url": "https://oss.example.invalid/photos/aa/object.jpg?Expires=1&Signature=secret",
    }
    row.update(overrides)
    return row


def manifest_item(**overrides: Any) -> ManifestItem:
    return ManifestItem.from_mapping(valid_row(**overrides))


def manifest_jsonl(count: int, *, byte_size: int | None = None) -> str:
    item_size = len(IMAGE_BYTES) if byte_size is None else byte_size
    rows: list[dict[str, Any]] = [
        {
            "schema": SCHEMA,
            "kind": "manifest",
            "planned_count": count,
            "planned_bytes": count * item_size,
        }
    ]
    for index in range(count):
        rows.append(
            valid_row(
                group_id=f"group-{index:03d}",
                photo_id=f"{index:012d}",
                relative_path=f"T01/group-{index:03d}/photo-{index:03d}.jpg",
                storage_key=f"photos/{index:03d}.jpg",
                byte_size=item_size,
                download_url=(
                    f"https://oss.example.invalid/photos/{index:03d}.jpg"
                    f"?Expires=1&Signature=secret-{index}"
                ),
            )
        )
    return "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows)


def fake_oss_opener(_url: str, *, timeout: int) -> FakeResponse:
    assert timeout == 30
    return FakeResponse(IMAGE_BYTES)


@pytest.mark.parametrize(
    "relative_path",
    [
        "../escape.jpg",
        "/absolute.jpg",
        "C:/absolute.jpg",
        "\\\\server\\share\\photo.jpg",
        "safe/..\\escape.jpg",
        "safe//photo.jpg",
        "safe/./photo.jpg",
        "safe/NUL.jpg",
        "safe/com1.png",
        "safe/trailing. /photo.jpg",
        "safe/photo?.jpg",
    ],
)
def test_manifest_rejects_unsafe_windows_relative_paths(relative_path: str) -> None:
    """Removing a Windows path guard must allow one of these escape/alias paths."""
    with pytest.raises(ValueError, match="relative_path"):
        ManifestItem.from_mapping(valid_row(relative_path=relative_path))


@pytest.mark.parametrize(
    ("updates", "message"),
    [
        ({"sha256": "not-a-hash"}, "sha256"),
        ({"byte_size": 0}, "byte_size"),
        ({"byte_size": True}, "byte_size"),
        ({"download_url": "http://oss.invalid/photo.jpg"}, "HTTPS"),
        ({"schema": "module-manager-oss-export/v2"}, "schema"),
        ({"kind": "manifest"}, "kind"),
        ({"group_id": ""}, "group_id"),
        ({"storage_key": ""}, "storage_key"),
    ],
)
def test_manifest_item_strictly_validates_required_fields(
    updates: dict[str, Any], message: str
) -> None:
    """Weakening a required field check must admit a malformed download request."""
    with pytest.raises(ValueError, match=message):
        ManifestItem.from_mapping(valid_row(**updates))


def test_manifest_item_rejects_missing_or_unknown_fields() -> None:
    """Changing the fixed item schema must not be silently accepted."""
    missing = valid_row()
    missing.pop("photo_id")
    with pytest.raises(ValueError, match="fields"):
        ManifestItem.from_mapping(missing)
    with pytest.raises(ValueError, match="fields"):
        ManifestItem.from_mapping(valid_row(unexpected="value"))


def test_download_retries_to_part_then_atomically_replaces(tmp_path: Path) -> None:
    """Removing retry or atomic replacement must leave the target absent/partial."""
    client = FakeHttpClient([TimeoutError(), TimeoutError(), IMAGE_BYTES])
    report = download_manifest(
        [manifest_item()],
        tmp_path,
        max_workers=1,
        retry_delays=(1.0, 2.0, 4.0),
        opener=client.open,
        sleeper=client.sleep,
    )
    target = tmp_path / manifest_item().relative_path
    assert target.read_bytes() == IMAGE_BYTES
    assert not target.with_name(target.name + ".part").exists()
    assert client.sleeps == [1.0, 2.0]
    assert report.succeeded == 1
    assert report.failed == 0


@pytest.mark.parametrize(
    ("payload", "expected_error"),
    [(b"wrong-size", "byte size mismatch"), (b"verified-fake-oss-imagf", "sha256 mismatch")],
)
def test_verification_failure_keeps_no_final_or_part_file(
    tmp_path: Path, payload: bytes, expected_error: str
) -> None:
    """Removing size/hash verification must publish corrupt bytes as successful."""
    item = manifest_item()
    report = download_manifest(
        [item],
        tmp_path,
        max_workers=1,
        retry_delays=(),
        opener=lambda *_a, **_k: FakeResponse(payload),
    )
    target = tmp_path / item.relative_path
    assert report.failed == 1
    assert expected_error in (report.results[0].error or "")
    assert not target.exists()
    assert not target.with_name(target.name + ".part").exists()


def test_expired_signature_is_retried_and_error_is_sanitized(tmp_path: Path) -> None:
    """Leaking the HTTPError URL must expose its Signature in result/report output."""
    url = valid_row()["download_url"]
    error = HTTPError(url, 403, "Signature expired", {}, None)
    client = FakeHttpClient([error, error])
    report = download_manifest(
        [manifest_item()],
        tmp_path,
        max_workers=1,
        retry_delays=(1.0,),
        opener=client.open,
        sleeper=client.sleep,
    )
    assert report.failed == 1
    assert client.sleeps == [1.0]
    assert report.results[0].error == "HTTP 403: Signature expired"
    assert "https://" not in report.results[0].error
    assert "Signature=" not in report.results[0].error


def test_error_sanitizer_removes_bare_signature_query_fragments() -> None:
    """Redacting only the value must still leave a forbidden Signature= query marker."""
    message = sanitize_error(
        RuntimeError("request failed: Signature=top-secret&OSSAccessKeyId=key-secret")
    )
    assert "Signature=" not in message
    assert "OSSAccessKeyId=" not in message
    assert "top-secret" not in message
    assert "key-secret" not in message


@pytest.mark.parametrize("max_workers", [0, 5, -1, 1.5, "4", True])
def test_download_manifest_rejects_workers_outside_one_to_four(
    tmp_path: Path, max_workers: Any
) -> None:
    """Removing the worker limit must permit unsafe local/OSS concurrency."""
    with pytest.raises(ValueError, match="max_workers"):
        download_manifest([], tmp_path, max_workers=max_workers)


@pytest.mark.parametrize("max_pending", [0, 9, -1, 1.5, "8", True])
def test_download_manifest_rejects_pending_bound_outside_one_to_eight(
    tmp_path: Path, max_pending: Any
) -> None:
    """Removing the pending limit must permit unbounded signed-URL accumulation."""
    with pytest.raises(ValueError, match="max_pending"):
        download_manifest([], tmp_path, max_pending=max_pending)


def test_duplicate_and_case_colliding_output_paths_are_rejected(tmp_path: Path) -> None:
    """Removing canonical collision checks must let two items overwrite one Windows file."""
    exact = [manifest_item(), manifest_item(photo_id="2", storage_key="other")]
    with pytest.raises(ValueError, match="duplicate relative_path"):
        download_manifest(exact, tmp_path, max_workers=1, retry_delays=(), opener=fake_oss_opener)

    case_collision = [
        manifest_item(relative_path="T01/Folder/Photo.jpg"),
        manifest_item(
            photo_id="2",
            storage_key="other",
            relative_path="t01/folder/photo.JPG",
        ),
    ]
    with pytest.raises(ValueError, match="duplicate relative_path"):
        download_manifest(
            case_collision,
            tmp_path / "case",
            max_workers=1,
            retry_delays=(),
            opener=fake_oss_opener,
        )


@pytest.mark.parametrize(
    "relative_paths",
    [
        ("folder/photo.jpg", "folder/photo.jpg.part"),
        ("export-report.json",),
        ("photos.zip.part",),
    ],
)
def test_runtime_part_and_composer_output_collisions_are_rejected(
    tmp_path: Path, relative_paths: tuple[str, ...]
) -> None:
    """Ignoring runtime outputs must let a download overwrite a part/report artifact."""
    items = [
        manifest_item(
            group_id=str(index),
            photo_id=str(index),
            relative_path=relative_path,
            storage_key=f"p/{index}.jpg",
        )
        for index, relative_path in enumerate(relative_paths)
    ]
    with pytest.raises(ValueError, match="relative_path"):
        download_manifest(
            items,
            tmp_path,
            max_workers=1,
            retry_delays=(),
            opener=fake_oss_opener,
        )


def test_results_remain_in_manifest_order_and_partial_success_is_preserved(
    tmp_path: Path,
) -> None:
    """Recording completion order must make reports unstable across runs."""
    lock = threading.Lock()
    attempts: dict[str, int] = {}

    def opener(url: str, *, timeout: int) -> FakeResponse:
        assert timeout == 30
        with lock:
            attempts[url] = attempts.get(url, 0) + 1
        if "001.jpg" in url:
            raise TimeoutError("signed request timed out")
        if "000.jpg" in url:
            time.sleep(0.03)
        return FakeResponse(IMAGE_BYTES)

    items = [
        manifest_item(
            group_id=f"g-{i}",
            photo_id=str(i),
            relative_path=f"g-{i}/photo.jpg",
            storage_key=f"photos/{i:03d}.jpg",
            download_url=f"https://oss.invalid/{i:03d}.jpg?Signature=secret",
        )
        for i in range(3)
    ]
    report = download_manifest(
        items,
        tmp_path,
        max_workers=3,
        retry_delays=(),
        opener=opener,
    )
    assert [result.item.photo_id for result in report.results] == ["0", "1", "2"]
    assert report.succeeded == 2
    assert report.failed == 1
    assert (tmp_path / "g-0/photo.jpg").read_bytes() == IMAGE_BYTES
    assert not (tmp_path / "g-1/photo.jpg").exists()
    assert (tmp_path / "g-2/photo.jpg").read_bytes() == IMAGE_BYTES


def test_concurrency_never_exceeds_four(tmp_path: Path) -> None:
    """Using more executor workers than allowed must raise observed concurrency above four."""
    lock = threading.Lock()
    active = 0
    peak = 0

    def opener(_url: str, *, timeout: int) -> FakeResponse:
        nonlocal active, peak
        assert timeout == 30
        with lock:
            active += 1
            peak = max(peak, active)
        time.sleep(0.02)
        with lock:
            active -= 1
        return FakeResponse(IMAGE_BYTES)

    items = [
        manifest_item(
            group_id=str(i),
            photo_id=str(i),
            relative_path=f"g{i}/p.jpg",
            storage_key=f"p/{i}.jpg",
            download_url=f"https://oss.invalid/{i}?Signature=secret",
        )
        for i in range(20)
    ]
    report = download_manifest(items, tmp_path, max_workers=4, opener=opener)
    assert report.succeeded == 20
    assert 2 <= peak <= 4


def test_pending_futures_bound_applies_backpressure_to_lazy_items(tmp_path: Path) -> None:
    """Eagerly collecting items must consume more than eight signed URLs before progress."""
    release = threading.Event()
    started = threading.Event()
    consumed = 0
    consumed_lock = threading.Lock()

    def items() -> Iterable[ManifestItem]:
        nonlocal consumed
        for i in range(20):
            with consumed_lock:
                consumed += 1
            yield manifest_item(
                group_id=str(i),
                photo_id=str(i),
                relative_path=f"g{i}/p.jpg",
                storage_key=f"p/{i}.jpg",
                download_url=f"https://oss.invalid/{i}?Signature=secret",
            )

    def opener(_url: str, *, timeout: int) -> FakeResponse:
        assert timeout == 30
        started.set()
        assert release.wait(timeout=5)
        return FakeResponse(IMAGE_BYTES)

    outcome: list[ExportReport] = []
    worker = threading.Thread(
        target=lambda: outcome.append(
            download_manifest(items(), tmp_path, max_workers=4, opener=opener)
        )
    )
    worker.start()
    assert started.wait(timeout=2)
    time.sleep(0.05)
    with consumed_lock:
        assert consumed == MAX_PENDING_DOWNLOADS == 8
    release.set()
    worker.join(timeout=10)
    assert not worker.is_alive()
    assert outcome[0].succeeded == 20


def test_completed_report_releases_all_signed_urls_after_bounded_run(tmp_path: Path) -> None:
    """Keeping ManifestItem in report/result must retain every completed signed URL."""
    items = [
        manifest_item(
            group_id=str(index),
            photo_id=str(index),
            relative_path=f"g{index}/photo.jpg",
            storage_key=f"objects/{index}.jpg",
            download_url=(
                f"https://oss.invalid/{index}.jpg?Signature=secret-{index}"
            ),
        )
        for index in range(24)
    ]
    report = download_manifest(items, tmp_path, max_workers=4, opener=fake_oss_opener)

    rendered = repr(report) + json.dumps(asdict(report), default=str)
    assert "https://" not in rendered
    assert "Signature=" not in rendered
    assert "download_url" not in rendered
    assert all(not hasattr(item, "download_url") for item in report.items)
    assert all(not hasattr(result.item, "download_url") for result in report.results)
    assert report.succeeded == 24


@pytest.mark.parametrize(
    "relative_paths",
    [
        ("a", "a/child.jpg"),
        ("a/child.jpg", "A"),
        ("export-report.json/child.jpg",),
        ("photos.zip.part/child",),
    ],
)
def test_file_directory_prefix_collisions_are_rejected(
    tmp_path: Path, relative_paths: tuple[str, ...]
) -> None:
    """Comparing only complete path strings must miss file/directory aliases."""
    items = [
        manifest_item(
            group_id=str(index),
            photo_id=str(index),
            relative_path=relative_path,
            storage_key=f"objects/{index}.jpg",
        )
        for index, relative_path in enumerate(relative_paths)
    ]
    with pytest.raises(ValueError, match="relative_path"):
        download_manifest(
            items,
            tmp_path,
            max_workers=1,
            retry_delays=(),
            opener=fake_oss_opener,
        )


@pytest.mark.parametrize(
    "relative_path",
    [
        "safe/CONIN$/photo.jpg",
        "safe/CONOUT$.txt",
        "safe/COM¹.jpg",
        "safe/COM²",
        "safe/COM³.png",
        "safe/LPT¹.jpg",
        "safe/LPT²",
        "safe/LPT³.png",
    ],
)
def test_extended_windows_device_aliases_are_rejected(relative_path: str) -> None:
    """An incomplete device set must permit Windows console/device aliases."""
    with pytest.raises(ValueError, match="reserved name"):
        ManifestItem.from_mapping(valid_row(relative_path=relative_path))


def test_existing_symlink_component_is_rejected_before_outside_write(
    tmp_path: Path,
) -> None:
    """Resolving once must permit an existing in-root link that points outside."""
    output_root = tmp_path / "output"
    outside = tmp_path / "outside"
    output_root.mkdir()
    outside.mkdir()
    link = output_root / "linked"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"symlink creation is unavailable: {exc}")

    with pytest.raises(ValueError, match="symlink|reparse"):
        download_manifest(
            [manifest_item(relative_path="linked/escaped.jpg")],
            output_root,
            max_workers=1,
            retry_delays=(),
            opener=fake_oss_opener,
        )
    assert list(outside.iterdir()) == []


def test_windows_reparse_attribute_detection_boundary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ignoring st_file_attributes must miss Windows junction/reparse components."""
    reparse_flag = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
    monkeypatch.setattr(
        exporter.os,
        "lstat",
        lambda _path: SimpleNamespace(st_mode=stat.S_IFDIR, st_file_attributes=reparse_flag),
    )
    assert exporter._path_is_link_or_reparse(tmp_path / "junction") is True


def test_symlink_injected_before_replace_is_rejected_fail_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Removing replace-time revalidation must accept a newly injected reparse target."""
    outside = tmp_path / "outside-target.jpg"
    probe = tmp_path / "symlink-probe"
    try:
        probe.symlink_to(outside)
        probe.unlink()
    except OSError as exc:
        pytest.skip(f"symlink creation is unavailable: {exc}")

    item = manifest_item()
    target = tmp_path / item.relative_path
    real_fsync = os.fsync
    injected = False

    def inject_target_symlink(file_descriptor: int) -> None:
        nonlocal injected
        real_fsync(file_descriptor)
        if not injected:
            target.symlink_to(outside)
            injected = True

    monkeypatch.setattr(exporter.os, "fsync", inject_target_symlink)
    report = download_manifest(
        [item],
        tmp_path,
        max_workers=1,
        retry_delays=(),
        opener=fake_oss_opener,
    )
    assert injected is True
    assert report.failed == 1
    assert target.is_symlink()
    assert not outside.exists()
    assert "symlink" in (report.results[0].error or "")
    assert "https://" not in (report.results[0].error or "")


def test_empty_manifest_rejects_symlink_output_root_without_outside_write(
    tmp_path: Path,
) -> None:
    """Skipping download workers must not skip the output-root reparse boundary."""
    outside = tmp_path / "outside"
    outside.mkdir()
    output_root = tmp_path / "output-link"
    try:
        output_root.symlink_to(outside, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"symlink creation is unavailable: {exc}")
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 0,
        "planned_bytes": 0,
    }
    with pytest.raises(ValueError, match="symlink|reparse"):
        run_export(json.dumps(header) + "\n", output_root, opener=fake_oss_opener)
    assert list(outside.iterdir()) == []


def test_report_part_symlink_is_rejected_without_overwriting_outside_target(
    tmp_path: Path,
) -> None:
    """Opening composer .part with w must not follow a pre-existing symlink."""
    outside = tmp_path / "outside-report.json"
    outside.write_text("original", encoding="utf-8")
    part = tmp_path / "export-report.json.part"
    try:
        part.symlink_to(outside)
    except OSError as exc:
        pytest.skip(f"symlink creation is unavailable: {exc}")
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 0,
        "planned_bytes": 0,
    }
    with pytest.raises(ValueError, match="symlink|reparse"):
        run_export(json.dumps(header) + "\n", tmp_path, opener=fake_oss_opener)
    assert outside.read_text(encoding="utf-8") == "original"


@pytest.mark.parametrize(
    "text",
    [
        "",
        json.dumps(valid_row()) + "\n",
        json.dumps({"schema": SCHEMA, "kind": "manifest", "planned_count": 0}) + "\n",
        json.dumps(
            {
                "schema": SCHEMA,
                "kind": "manifest",
                "planned_count": 0,
                "planned_bytes": 0,
                "download_url": "https://forbidden.invalid/?Signature=secret",
            }
        )
        + "\n",
    ],
)
def test_stream_requires_exact_url_free_header_first(tmp_path: Path, text: str) -> None:
    """Weak header validation must admit missing fields, items first, or signed URLs."""
    with pytest.raises(ValueError, match="header"):
        run_export_stream(io.StringIO(text), tmp_path, opener=fake_oss_opener)


def test_stream_rejects_duplicate_or_late_header(tmp_path: Path) -> None:
    """Treating a later manifest row as an item must allow ambiguous stream boundaries."""
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 1,
        "planned_bytes": len(IMAGE_BYTES),
    }
    text = json.dumps(header) + "\n" + json.dumps(header) + "\n"
    with pytest.raises(ValueError, match="header"):
        run_export_stream(io.StringIO(text), tmp_path, opener=fake_oss_opener)


def test_json_header_rejects_duplicate_object_keys(tmp_path: Path) -> None:
    """Standard json.loads must not silently accept the last duplicate header key."""
    text = (
        '{"schema":"module-manager-oss-export/v1","kind":"manifest",'
        '"kind":"manifest","planned_count":0,"planned_bytes":0}\n'
    )
    with pytest.raises(ValueError, match="duplicate JSON key"):
        run_export_stream(io.StringIO(text), tmp_path, opener=fake_oss_opener)


def test_json_item_rejects_duplicate_object_keys(tmp_path: Path) -> None:
    """A duplicate item field must not be reduced to its last attacker value."""
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 1,
        "planned_bytes": len(IMAGE_BYTES),
    }
    item_text = json.dumps(valid_row())
    item_text = item_text.replace(
        '"photo_id": "000123456789"',
        '"photo_id":"first","photo_id":"000123456789"',
    )
    text = json.dumps(header) + "\n" + item_text + "\n"
    with pytest.raises(ValueError, match="duplicate JSON key"):
        run_export_stream(io.StringIO(text), tmp_path, opener=fake_oss_opener)


@pytest.mark.parametrize(
    ("planned_count", "planned_bytes", "message"),
    [(2, len(IMAGE_BYTES), "count"), (1, len(IMAGE_BYTES) + 1, "bytes")],
)
def test_stream_rejects_eof_count_or_byte_total_mismatch(
    tmp_path: Path, planned_count: int, planned_bytes: int, message: str
) -> None:
    """Skipping the EOF reconciliation must report an incomplete manifest as complete."""
    rows = [
        {
            "schema": SCHEMA,
            "kind": "manifest",
            "planned_count": planned_count,
            "planned_bytes": planned_bytes,
        },
        valid_row(),
    ]
    text = "".join(json.dumps(row) + "\n" for row in rows)
    with pytest.raises(ValueError, match=message):
        run_export_stream(io.StringIO(text), tmp_path, opener=fake_oss_opener)
    assert not (tmp_path / "export-report.json").exists()
    assert not (tmp_path / "photos.zip").exists()


def test_stream_consumes_only_bounded_item_lines_before_download_progress(
    tmp_path: Path,
) -> None:
    """Calling read/readlines must collect every signed URL before any download begins."""
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 20,
        "planned_bytes": 20 * len(IMAGE_BYTES),
    }
    lines = [json.dumps(header) + "\n"] + [
        json.dumps(
            valid_row(
                group_id=str(i),
                photo_id=str(i),
                relative_path=f"g{i}/p.jpg",
                storage_key=f"p/{i}.jpg",
                download_url=f"https://oss.invalid/{i}?Signature=secret",
            )
        )
        + "\n"
        for i in range(20)
    ]

    class TrackingStream:
        def __init__(self) -> None:
            self.index = 0
            self.lock = threading.Lock()

        def readline(self) -> str:
            with self.lock:
                if self.index == len(lines):
                    return ""
                line = lines[self.index]
                self.index += 1
                return line

        def __iter__(self) -> "TrackingStream":
            return self

        def __next__(self) -> str:
            line = self.readline()
            if not line:
                raise StopIteration
            return line

    stream = TrackingStream()
    release = threading.Event()
    started = threading.Event()

    def opener(_url: str, *, timeout: int) -> FakeResponse:
        assert timeout == 30
        started.set()
        assert release.wait(timeout=5)
        return FakeResponse(IMAGE_BYTES)

    outcome: list[ExportReport] = []
    worker = threading.Thread(
        target=lambda: outcome.append(
            run_export_stream(stream, tmp_path, max_workers=4, opener=opener)  # type: ignore[arg-type]
        )
    )
    worker.start()
    assert started.wait(timeout=2)
    time.sleep(0.05)
    with stream.lock:
        assert stream.index == 1 + MAX_PENDING_DOWNLOADS
    release.set()
    worker.join(timeout=10)
    assert not worker.is_alive()
    assert outcome[0].succeeded == 20


def test_disk_gate_runs_before_output_creation_or_download(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Moving the disk check after mkdir/network must cause observable side effects first."""
    Usage = namedtuple("Usage", "total used free")
    opened = False

    def opener(*_args: object, **_kwargs: object) -> FakeResponse:
        nonlocal opened
        opened = True
        return FakeResponse(IMAGE_BYTES)

    monkeypatch.setattr(shutil, "disk_usage", lambda _path: Usage(100, 99, 1))
    with pytest.raises(InsufficientDiskError, match="required"):
        run_export(
            stdin_text=manifest_jsonl(1),
            output_root=tmp_path / "new-output",
            formats={"files", "zip"},
            opener=opener,
        )
    assert not (tmp_path / "new-output").exists()
    assert opened is False


@pytest.mark.parametrize(
    ("planned_bytes", "formats", "required"),
    [
        (0, {"files"}, 268435456),
        (1024, {"files"}, 268436480),
        (1024, {"files", "zip"}, 268437504),
        (3221225472, {"files", "zip"}, 6764573491),
    ],
)
def test_disk_gate_uses_exact_formula_for_multiple_sizes_and_zip_modes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    planned_bytes: int,
    formats: set[str],
    required: int,
) -> None:
    """Changing any formula term must change the independently calculated threshold."""
    Usage = namedtuple("Usage", "total used free")
    monkeypatch.setattr(
        shutil,
        "disk_usage",
        lambda _path: Usage(required + 100, 101, required - 1),
    )
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 0,
        "planned_bytes": planned_bytes,
    }
    with pytest.raises(InsufficientDiskError, match=f"required {required} bytes"):
        run_export(
            json.dumps(header) + "\n",
            tmp_path / "output",
            formats=formats,
            opener=fake_oss_opener,
        )


def test_outputs_match_manifest_and_contain_no_signed_urls(tmp_path: Path) -> None:
    """Serializing ManifestItem directly must leak download_url/Signature into artifacts."""
    result = run_export(
        stdin_text=manifest_jsonl(3),
        output_root=tmp_path,
        formats={"files", "zip", "csv", "xlsx"},
        opener=fake_oss_opener,
    )
    assert result.planned == result.succeeded == 3
    assert result.failed == 0
    with ZipFile(tmp_path / "photos.zip") as archive:
        assert archive.namelist() == [item.relative_path for item in result.items]
        assert all(archive.read(name) == IMAGE_BYTES for name in archive.namelist())
    csv_bytes = (tmp_path / "manifest.csv").read_bytes()
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    report_text = (tmp_path / "export-report.json").read_text(encoding="utf-8")
    for artifact_text in (report_text, csv_bytes.decode("utf-8-sig")):
        assert "download_url" not in artifact_text
        assert "https://" not in artifact_text
        assert "Signature=" not in artifact_text
    assert not list(tmp_path.rglob("*.part"))


def test_all_report_metadata_is_sanitized_against_url_and_signature_fragments(
    tmp_path: Path,
) -> None:
    """Sanitizing only errors must allow hostile metadata to persist signed secrets."""
    row = valid_row(
        category="https://host.invalid/photo?Signature=category-secret",
        storage_key="objects/photo.jpg?Signature=key-secret",
    )
    header = {
        "schema": SCHEMA,
        "kind": "manifest",
        "planned_count": 1,
        "planned_bytes": len(IMAGE_BYTES),
    }
    text = json.dumps(header) + "\n" + json.dumps(row) + "\n"
    run_export(
        text,
        tmp_path,
        formats={"files", "csv", "xlsx"},
        opener=fake_oss_opener,
    )
    report_text = (tmp_path / "export-report.json").read_text(encoding="utf-8")
    csv_text = (tmp_path / "manifest.csv").read_text(encoding="utf-8-sig")
    workbook = load_workbook(tmp_path / "manifest.xlsx")
    workbook_text = " ".join(
        str(cell.value or "")
        for sheet in workbook.worksheets
        for row_cells in sheet.iter_rows()
        for cell in row_cells
    )
    for artifact_text in (report_text, csv_text, workbook_text):
        assert "https://" not in artifact_text
        assert "Signature=" not in artifact_text
        assert "category-secret" not in artifact_text
        assert "key-secret" not in artifact_text


def test_xlsx_has_manifest_formatting_and_failure_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Omitting planned failures from XLSX must hide partial-export evidence."""
    calls = 0

    def opener(_url: str, *, timeout: int) -> FakeResponse:
        nonlocal calls
        assert timeout == 30
        calls += 1
        if calls == 1:
            raise RuntimeError(
                "GET https://oss.invalid/p.jpg?OSSAccessKeyId=id&Signature=top-secret failed"
            )
        return FakeResponse(IMAGE_BYTES)

    monkeypatch.setattr(exporter, "RETRY_DELAYS", ())
    result = run_export(
        stdin_text=manifest_jsonl(2),
        output_root=tmp_path,
        formats={"files", "xlsx"},
        max_workers=1,
        opener=opener,
    )
    assert (result.succeeded, result.failed) == (1, 1)
    workbook = load_workbook(tmp_path / "manifest.xlsx")
    assert workbook.sheetnames == ["Manifest", "Failures"]
    manifest = workbook["Manifest"]
    failures = workbook["Failures"]
    assert manifest.freeze_panes == "A2"
    assert manifest.auto_filter.ref == manifest.dimensions
    assert manifest.max_row == 3
    headers = [cell.value for cell in manifest[1]]
    group_column = headers.index("group_id") + 1
    photo_column = headers.index("photo_id") + 1
    assert manifest.cell(2, group_column).number_format == "@"
    assert manifest.cell(2, photo_column).number_format == "@"
    assert failures.freeze_panes == "A2"
    assert failures.max_row == 2
    failure_text = " ".join(str(cell.value or "") for row in failures for cell in row)
    assert "https://" not in failure_text
    assert "Signature=" not in failure_text
    assert not (tmp_path / result.items[0].relative_path).exists()
    assert (tmp_path / result.items[1].relative_path).read_bytes() == IMAGE_BYTES


@pytest.mark.parametrize("dangerous", ["\t=1+1", "   +SUM(1,1)", "\t@WEBSERVICE(1)"])
def test_csv_and_xlsx_escape_formulas_after_leading_whitespace_or_controls(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    dangerous: str,
) -> None:
    """Checking only value[0] must leave whitespace-prefixed Excel formulas executable."""
    monkeypatch.setattr(exporter, "RETRY_DELAYS", ())

    def opener(_url: str, *, timeout: int) -> FakeResponse:
        assert timeout == 30
        raise RuntimeError(dangerous)

    report = run_export(
        manifest_jsonl(1),
        tmp_path,
        formats={"files", "csv", "xlsx"},
        max_workers=1,
        opener=opener,
    )
    assert report.failed == 1
    with (tmp_path / "manifest.csv").open(encoding="utf-8-sig", newline="") as handle:
        csv_error = next(csv.DictReader(handle))["error"]
    workbook = load_workbook(tmp_path / "manifest.xlsx", data_only=False)
    failure_sheet = workbook["Failures"]
    error_column = [cell.value for cell in failure_sheet[1]].index("error") + 1
    xlsx_error = failure_sheet.cell(2, error_column)
    assert csv_error.startswith("'")
    assert xlsx_error.value.startswith("'")
    assert xlsx_error.data_type == "s"


def test_existing_part_directory_becomes_sanitized_item_failure_with_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An unlink failure in cleanup must not escape its future or suppress reports."""
    monkeypatch.setattr(exporter, "RETRY_DELAYS", ())
    target = tmp_path / "T01/group-000/photo-000.jpg"
    part = target.with_name(target.name + ".part")
    part.mkdir(parents=True)

    report = run_export(
        manifest_jsonl(1),
        tmp_path,
        formats={"files", "xlsx"},
        max_workers=1,
        opener=fake_oss_opener,
    )
    assert report.failed == 1
    assert not target.exists()
    assert part.is_dir()
    persisted = (tmp_path / "export-report.json").read_text(encoding="utf-8")
    assert "https://" not in persisted
    assert "Signature=" not in persisted
    assert load_workbook(tmp_path / "manifest.xlsx")["Failures"].max_row == 2


def test_parent_mkdir_failure_becomes_item_failure_and_still_composes_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Keeping mkdir outside the item try block must upgrade one bad path to preflight 3."""
    monkeypatch.setattr(exporter, "RETRY_DELAYS", ())
    blocking_parent = tmp_path / "T01"
    blocking_parent.write_text("not a directory", encoding="utf-8")

    report = run_export(
        manifest_jsonl(1),
        tmp_path,
        formats={"files", "csv"},
        max_workers=1,
        opener=fake_oss_opener,
    )
    assert report.failed == 1
    assert (tmp_path / "export-report.json").is_file()
    assert (tmp_path / "manifest.csv").is_file()
    assert "https://" not in (report.results[0].error or "")


def test_part_is_visible_but_final_is_hidden_until_verified_replace(tmp_path: Path) -> None:
    """Writing directly to final must expose an incomplete file before verification."""
    blocked = threading.Event()
    release = threading.Event()

    class BlockingResponse:
        def __init__(self) -> None:
            self.calls = 0

        def __enter__(self) -> "BlockingResponse":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self, _size: int) -> bytes:
            self.calls += 1
            if self.calls == 1:
                return IMAGE_BYTES[:8]
            if self.calls == 2:
                blocked.set()
                assert release.wait(timeout=5)
                return IMAGE_BYTES[8:]
            return b""

    item = manifest_item()
    outcome: list[ExportReport] = []
    worker = threading.Thread(
        target=lambda: outcome.append(
            download_manifest(
                [item],
                tmp_path,
                max_workers=1,
                retry_delays=(),
                opener=lambda *_args, **_kwargs: BlockingResponse(),
            )
        )
    )
    worker.start()
    assert blocked.wait(timeout=2)
    target = tmp_path / item.relative_path
    part = target.with_name(target.name + ".part")
    assert part.is_file()
    assert not target.exists()
    release.set()
    worker.join(timeout=5)
    assert not worker.is_alive()
    assert outcome[0].succeeded == 1
    assert target.read_bytes() == IMAGE_BYTES
    assert not part.exists()


def test_download_calls_real_fsync_before_atomic_replace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Removing fsync or replace must be observable while the real file operation still runs."""
    real_fsync = os.fsync
    real_replace = os.replace
    fsync_calls: list[int] = []
    fsync_sizes: list[int] = []
    replace_calls: list[tuple[Path, Path]] = []

    def recording_fsync(file_descriptor: int) -> None:
        fsync_calls.append(file_descriptor)
        fsync_sizes.append(os.fstat(file_descriptor).st_size)
        real_fsync(file_descriptor)

    def recording_replace(source: str | os.PathLike[str], target: str | os.PathLike[str]) -> None:
        replace_calls.append((Path(source), Path(target)))
        real_replace(source, target)

    monkeypatch.setattr(exporter.os, "fsync", recording_fsync)
    monkeypatch.setattr(exporter.os, "replace", recording_replace)
    item = manifest_item()
    report = download_manifest(
        [item],
        tmp_path,
        max_workers=1,
        retry_delays=(),
        opener=fake_oss_opener,
    )
    target = tmp_path / item.relative_path
    assert report.succeeded == 1
    assert fsync_calls
    assert fsync_sizes == [len(IMAGE_BYTES)]
    assert replace_calls == [(target.with_name(target.name + ".part"), target)]
    assert target.read_bytes() == IMAGE_BYTES


def test_run_export_defaults_to_files_and_writes_atomic_report(tmp_path: Path) -> None:
    """Treating formats=None as every format must create unrequested artifacts."""
    result = run_export(manifest_jsonl(1), tmp_path, opener=fake_oss_opener)
    assert result.succeeded == 1
    assert (tmp_path / "export-report.json").is_file()
    assert not (tmp_path / "photos.zip").exists()
    assert not (tmp_path / "manifest.csv").exists()
    assert not (tmp_path / "manifest.xlsx").exists()
