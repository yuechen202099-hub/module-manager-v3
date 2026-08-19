#!/usr/bin/env python3
"""Download a signed OSS manifest to a Windows-local export directory.

The server-side companion writes a URL-free manifest header followed by signed
item rows to stdout.  This module consumes those rows lazily so no more than a
small bounded set of signed URLs is resident or waiting for download.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import shutil
import sys
import time
import unicodedata
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Mapping, Sequence, TextIO
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import urlopen
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


SCHEMA = "module-manager-oss-export/v1"
RETRY_DELAYS = (1.0, 2.0, 4.0)
MAX_DOWNLOAD_WORKERS = 4
MAX_PENDING_DOWNLOADS = 8
DISK_RESERVE_BYTES = 256 * 1024 * 1024
DOWNLOAD_TIMEOUT_SECONDS = 30
CHUNK_SIZE = 1024 * 1024

HEADER_FIELDS = frozenset({"schema", "kind", "planned_count", "planned_bytes"})
ITEM_FIELDS = frozenset(
    {
        "schema",
        "kind",
        "group_id",
        "photo_id",
        "category",
        "relative_path",
        "storage_bucket",
        "storage_key",
        "sha256",
        "byte_size",
        "content_type",
        "download_url",
    }
)
TABULAR_FIELDS = (
    "group_id",
    "photo_id",
    "category",
    "relative_path",
    "storage_bucket",
    "storage_key",
    "sha256",
    "byte_size",
    "content_type",
    "status",
    "downloaded_bytes",
    "error",
)
TEXT_TABULAR_FIELDS = frozenset(
    {
        "group_id",
        "photo_id",
        "category",
        "relative_path",
        "storage_bucket",
        "storage_key",
        "sha256",
        "content_type",
        "status",
        "error",
    }
)

_SHA256_RE = re.compile(r"^[0-9a-fA-F]{64}$")
_DRIVE_RE = re.compile(r"^[A-Za-z]:")
_URL_RE = re.compile(r"(?i)https?://[^\s\"'<>]+")
_SECRET_QUERY_RE = re.compile(
    r"(?i)(Signature|OSSAccessKeyId|Expires|SecurityToken|"
    r"x-oss-signature|x-oss-credential|x-oss-security-token)=([^&\s]+)"
)
_WINDOWS_INVALID_CHARS = frozenset('<>:"\\|?*')
_WINDOWS_RESERVED_NAMES = frozenset(
    {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        "CLOCK$",
        *(f"COM{index}" for index in range(1, 10)),
        *(f"LPT{index}" for index in range(1, 10)),
    }
)
_LOCAL_EXPORT_OUTPUTS = frozenset(
    {
        "export-report.json",
        "manifest.csv",
        "manifest.xlsx",
        "photos.zip",
    }
)


class InsufficientDiskError(RuntimeError):
    """Raised when the destination volume cannot safely hold the export."""


def _required_string(value: Mapping[str, Any], field: str) -> str:
    candidate = value.get(field)
    if not isinstance(candidate, str) or not candidate or candidate != candidate.strip():
        raise ValueError(f"{field} must be a non-empty string")
    if any(ord(character) < 32 for character in candidate):
        raise ValueError(f"{field} contains control characters")
    return candidate


def _required_positive_int(value: Mapping[str, Any], field: str) -> int:
    candidate = value.get(field)
    if isinstance(candidate, bool) or not isinstance(candidate, int) or candidate <= 0:
        raise ValueError(f"{field} must be a positive integer")
    return candidate


def _validated_relative_path(value: str) -> str:
    if not isinstance(value, str) or not value or value != value.strip():
        raise ValueError("relative_path must be a non-empty relative path")
    if value.startswith(("/", "\\")) or _DRIVE_RE.match(value):
        raise ValueError("relative_path must not be absolute")
    if "\\" in value:
        raise ValueError("relative_path must use forward slashes")
    segments = value.split("/")
    if any(segment in {"", ".", ".."} for segment in segments):
        raise ValueError("relative_path contains an unsafe segment")
    for segment in segments:
        if len(segment) > 255:
            raise ValueError("relative_path segment is too long")
        if segment.endswith((" ", ".")):
            raise ValueError("relative_path contains a Windows trailing alias")
        if any(ord(character) < 32 for character in segment):
            raise ValueError("relative_path contains control characters")
        if any(character in _WINDOWS_INVALID_CHARS for character in segment):
            raise ValueError("relative_path contains a Windows-invalid character")
        device_name = segment.split(".", 1)[0].rstrip(" .").upper()
        if device_name in _WINDOWS_RESERVED_NAMES:
            raise ValueError("relative_path contains a Windows reserved name")
    return value


def _collision_key(relative_path: str) -> str:
    return unicodedata.normalize("NFC", relative_path).casefold()


def _validated_download_url(value: str) -> str:
    if not isinstance(value, str) or not value or value != value.strip():
        raise ValueError("download_url must be a non-empty HTTPS URL")
    parsed = urlsplit(value)
    if parsed.scheme.lower() != "https" or not parsed.hostname:
        raise ValueError("download_url must use HTTPS")
    if parsed.username is not None or parsed.password is not None or parsed.fragment:
        raise ValueError("download_url must be a plain HTTPS signed URL")
    return value


@dataclass(frozen=True, slots=True)
class ManifestItem:
    schema: str
    kind: str
    group_id: str
    photo_id: str
    category: str
    relative_path: str
    storage_bucket: str
    storage_key: str
    sha256: str
    byte_size: int
    content_type: str
    download_url: str

    @classmethod
    def from_mapping(cls, value: Mapping[str, Any]) -> "ManifestItem":
        if not isinstance(value, Mapping) or frozenset(value.keys()) != ITEM_FIELDS:
            raise ValueError("manifest item fields do not match the required schema")
        if value.get("schema") != SCHEMA:
            raise ValueError("manifest item schema is unsupported")
        if value.get("kind") != "item":
            raise ValueError("manifest item kind must be item")
        sha256 = value.get("sha256")
        if not isinstance(sha256, str) or not _SHA256_RE.fullmatch(sha256):
            raise ValueError("sha256 must contain exactly 64 hexadecimal characters")
        relative_path = _validated_relative_path(_required_string(value, "relative_path"))
        return cls(
            schema=SCHEMA,
            kind="item",
            group_id=_required_string(value, "group_id"),
            photo_id=_required_string(value, "photo_id"),
            category=_required_string(value, "category"),
            relative_path=relative_path,
            storage_bucket=_required_string(value, "storage_bucket"),
            storage_key=_required_string(value, "storage_key"),
            sha256=sha256.lower(),
            byte_size=_required_positive_int(value, "byte_size"),
            content_type=_required_string(value, "content_type"),
            download_url=_validated_download_url(value.get("download_url")),
        )


@dataclass(frozen=True, slots=True)
class ItemResult:
    item: ManifestItem
    succeeded: bool
    local_path: str | None
    downloaded_bytes: int
    actual_sha256: str | None
    error: str | None

    @classmethod
    def success(cls, item: ManifestItem, target: Path, size: int) -> "ItemResult":
        return cls(
            item=item,
            succeeded=True,
            local_path=str(target.resolve()),
            downloaded_bytes=size,
            actual_sha256=item.sha256,
            error=None,
        )

    @classmethod
    def failure(cls, item: ManifestItem, error: str) -> "ItemResult":
        return cls(
            item=item,
            succeeded=False,
            local_path=None,
            downloaded_bytes=0,
            actual_sha256=None,
            error=error,
        )


@dataclass(frozen=True, slots=True)
class ExportReport:
    items: tuple[ManifestItem, ...]
    results: tuple[ItemResult, ...]
    planned_bytes: int
    output_root: Path
    generated_at: str

    @property
    def planned(self) -> int:
        return len(self.items)

    @property
    def succeeded(self) -> int:
        return sum(result.succeeded for result in self.results)

    @property
    def failed(self) -> int:
        return self.planned - self.succeeded

    @property
    def downloaded_bytes(self) -> int:
        return sum(result.downloaded_bytes for result in self.results)


def safe_output_path(output_root: Path, relative_path: str) -> Path:
    """Return a path beneath output_root after platform-independent validation."""
    safe_relative_path = _validated_relative_path(relative_path)
    root = output_root.resolve(strict=False)
    target = (root / Path(*safe_relative_path.split("/"))).resolve(strict=False)
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise ValueError("relative_path escapes the output root") from exc
    return target


def sanitize_error(error: BaseException | str) -> str:
    """Return a single-line diagnostic that cannot retain a signed URL."""
    if isinstance(error, HTTPError):
        raw = f"HTTP {error.code}: {error.reason}"
    elif isinstance(error, URLError):
        raw = f"URL error: {error.reason}"
    else:
        raw = str(error) or type(error).__name__
    sanitized = raw.replace("\r", " ").replace("\n", " ")
    sanitized = _URL_RE.sub("[redacted-url]", sanitized)
    sanitized = _SECRET_QUERY_RE.sub("[redacted-query]", sanitized)
    sanitized = re.sub(r"(?i)download_url", "signed request", sanitized)
    return sanitized[:1000]


def _download_one(
    item: ManifestItem,
    output_root: Path,
    *,
    opener: Callable[..., Any],
    sleeper: Callable[[float], None],
    retry_delays: tuple[float, ...],
) -> ItemResult:
    target = safe_output_path(output_root, item.relative_path)
    part = target.with_name(target.name + ".part")
    target.parent.mkdir(parents=True, exist_ok=True)
    for attempt in range(len(retry_delays) + 1):
        try:
            part.unlink(missing_ok=True)
            digest = hashlib.sha256()
            size = 0
            with opener(item.download_url, timeout=DOWNLOAD_TIMEOUT_SECONDS) as response:
                with part.open("wb") as output:
                    while True:
                        chunk = response.read(CHUNK_SIZE)
                        if not chunk:
                            break
                        if not isinstance(chunk, bytes):
                            raise ValueError("download response returned non-byte content")
                        size += len(chunk)
                        output.write(chunk)
                        digest.update(chunk)
                    output.flush()
                    os.fsync(output.fileno())
            if size != item.byte_size:
                raise ValueError(
                    f"byte size mismatch: expected {item.byte_size}, observed {size}"
                )
            observed_sha256 = digest.hexdigest()
            if observed_sha256 != item.sha256:
                raise ValueError("sha256 mismatch")
            os.replace(part, target)
            return ItemResult.success(item, target, size)
        except Exception as exc:
            part.unlink(missing_ok=True)
            if attempt == len(retry_delays):
                return ItemResult.failure(item, sanitize_error(exc))
            sleeper(retry_delays[attempt])
    raise AssertionError("unreachable")


def download_manifest(
    items: Iterable[ManifestItem],
    output_root: Path,
    max_workers: int = 4,
    retry_delays: tuple[float, ...] = RETRY_DELAYS,
    *,
    opener: Callable[..., Any] = urlopen,
    sleeper: Callable[[float], None] = time.sleep,
    max_pending: int = MAX_PENDING_DOWNLOADS,
) -> ExportReport:
    """Download items with stable results and bounded submission/backpressure."""
    if (
        isinstance(max_workers, bool)
        or not isinstance(max_workers, int)
        or not 1 <= max_workers <= MAX_DOWNLOAD_WORKERS
    ):
        raise ValueError("max_workers must be between 1 and 4")
    if (
        isinstance(max_pending, bool)
        or not isinstance(max_pending, int)
        or not 1 <= max_pending <= MAX_PENDING_DOWNLOADS
    ):
        raise ValueError("max_pending must be between 1 and 8")
    if not isinstance(retry_delays, tuple) or any(
        isinstance(delay, bool) or not isinstance(delay, (int, float)) or delay < 0
        for delay in retry_delays
    ):
        raise ValueError("retry_delays must be a tuple of non-negative numbers")

    iterator = iter(items)
    submitted_items: list[ManifestItem] = []
    results_by_index: dict[int, ItemResult] = {}
    used_runtime_paths: set[str] = {
        _collision_key(path)
        for output_name in _LOCAL_EXPORT_OUTPUTS
        for path in (output_name, output_name + ".part")
    }
    pending: dict[Future[ItemResult], int] = {}
    exhausted = False

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        while not exhausted or pending:
            while not exhausted and len(pending) < max_pending:
                try:
                    item = next(iterator)
                except StopIteration:
                    exhausted = True
                    break
                if not isinstance(item, ManifestItem):
                    raise ValueError("items must contain ManifestItem values")
                safe_output_path(output_root, item.relative_path)
                final_key = _collision_key(item.relative_path)
                part_key = _collision_key(item.relative_path + ".part")
                if final_key in used_runtime_paths or part_key in used_runtime_paths:
                    raise ValueError(
                        "duplicate relative_path or reserved runtime path is not allowed: "
                        f"{item.relative_path}"
                    )
                used_runtime_paths.update((final_key, part_key))
                index = len(submitted_items)
                submitted_items.append(item)
                future = executor.submit(
                    _download_one,
                    item,
                    output_root,
                    opener=opener,
                    sleeper=sleeper,
                    retry_delays=retry_delays,
                )
                pending[future] = index
            if not pending:
                continue
            completed, _ = wait(tuple(pending), return_when=FIRST_COMPLETED)
            for future in completed:
                index = pending.pop(future)
                results_by_index[index] = future.result()

    ordered_results = tuple(
        results_by_index[index] for index in range(len(submitted_items))
    )
    return ExportReport(
        items=tuple(submitted_items),
        results=ordered_results,
        planned_bytes=sum(item.byte_size for item in submitted_items),
        output_root=output_root.resolve(strict=False),
        generated_at=datetime.now(timezone.utc).isoformat(),
    )


def _parse_json_object(line: str, *, context: str) -> Mapping[str, Any]:
    try:
        value = json.loads(line)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"{context} must be valid JSON") from exc
    if not isinstance(value, Mapping):
        raise ValueError(f"{context} must be a JSON object")
    return value


def _parse_header(input_stream: TextIO) -> tuple[int, int]:
    first_line = input_stream.readline()
    if not first_line or not first_line.strip():
        raise ValueError("manifest header must be the first line")
    header = _parse_json_object(first_line, context="manifest header")
    if frozenset(header.keys()) != HEADER_FIELDS:
        raise ValueError("manifest header fields do not match the required schema")
    if header.get("schema") != SCHEMA or header.get("kind") != "manifest":
        raise ValueError("manifest header schema or kind is invalid")
    planned_count = header.get("planned_count")
    planned_bytes = header.get("planned_bytes")
    if (
        isinstance(planned_count, bool)
        or not isinstance(planned_count, int)
        or planned_count < 0
    ):
        raise ValueError("manifest header planned_count must be a non-negative integer")
    if (
        isinstance(planned_bytes, bool)
        or not isinstance(planned_bytes, int)
        or planned_bytes < 0
    ):
        raise ValueError("manifest header planned_bytes must be a non-negative integer")
    return planned_count, planned_bytes


def _iter_stream_items(
    input_stream: TextIO, *, planned_count: int, planned_bytes: int
) -> Iterator[ManifestItem]:
    observed_count = 0
    observed_bytes = 0
    for line_number, line in enumerate(input_stream, start=2):
        if not line.strip():
            raise ValueError(f"manifest item line {line_number} must not be blank")
        value = _parse_json_object(line, context=f"manifest item line {line_number}")
        if value.get("kind") == "manifest":
            raise ValueError("manifest header must be unique and appear only first")
        item = ManifestItem.from_mapping(value)
        observed_count += 1
        observed_bytes += item.byte_size
        yield item
    if observed_count != planned_count:
        raise ValueError(
            f"manifest item count mismatch: expected {planned_count}, observed {observed_count}"
        )
    if observed_bytes != planned_bytes:
        raise ValueError(
            f"manifest bytes total mismatch: expected {planned_bytes}, observed {observed_bytes}"
        )


def _existing_disk_probe(path: Path) -> Path:
    probe = path.resolve(strict=False)
    while not probe.exists() and probe != probe.parent:
        probe = probe.parent
    return probe


def _check_disk_space(output_root: Path, planned_bytes: int, formats: set[str]) -> None:
    required = planned_bytes
    if "zip" in formats:
        required += planned_bytes
    required += max(planned_bytes // 10, DISK_RESERVE_BYTES)
    free = shutil.disk_usage(_existing_disk_probe(output_root)).free
    if free < required:
        raise InsufficientDiskError(
            f"insufficient disk space: required {required} bytes, available {free} bytes"
        )


def _sanitized_row(result: ItemResult) -> dict[str, Any]:
    item = result.item
    row = {
        "group_id": item.group_id,
        "photo_id": item.photo_id,
        "category": item.category,
        "relative_path": item.relative_path,
        "storage_bucket": item.storage_bucket,
        "storage_key": item.storage_key,
        "sha256": item.sha256,
        "byte_size": item.byte_size,
        "content_type": item.content_type,
        "status": "succeeded" if result.succeeded else "failed",
        "downloaded_bytes": result.downloaded_bytes,
        "error": result.error or "",
    }
    return {
        key: _sanitize_artifact_text(value) if isinstance(value, str) else value
        for key, value in row.items()
    }


def _sanitize_artifact_text(value: str) -> str:
    sanitized = value.replace("\r", " ").replace("\n", " ")
    sanitized = _URL_RE.sub("[redacted-url]", sanitized)
    return _SECRET_QUERY_RE.sub("[redacted-query]", sanitized)


def _spreadsheet_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _fsync_file(path: Path) -> None:
    with path.open("r+b") as handle:
        handle.flush()
        os.fsync(handle.fileno())


def _atomic_csv(report: ExportReport, target: Path) -> None:
    part = target.with_name(target.name + ".part")
    try:
        with part.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=TABULAR_FIELDS)
            writer.writeheader()
            for result in report.results:
                writer.writerow(
                    {
                        key: _spreadsheet_text(value)
                        for key, value in _sanitized_row(result).items()
                    }
                )
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(part, target)
    finally:
        part.unlink(missing_ok=True)


def _format_sheet(sheet: Any, headers: Sequence[str], row_count: int) -> None:
    sheet.freeze_panes = "A2"
    last_column = sheet.cell(row=1, column=len(headers)).column_letter
    sheet.auto_filter.ref = f"A1:{last_column}{max(row_count + 1, 1)}"
    for column_index, header in enumerate(headers, start=1):
        if header in TEXT_TABULAR_FIELDS:
            for row_index in range(2, row_count + 2):
                sheet.cell(row=row_index, column=column_index).number_format = "@"


def _atomic_xlsx(report: ExportReport, target: Path) -> None:
    part = target.with_name(target.name + ".part")
    try:
        workbook = Workbook()
        manifest_sheet = workbook.active
        manifest_sheet.title = "Manifest"
        manifest_sheet.append(list(TABULAR_FIELDS))
        rows = [_sanitized_row(result) for result in report.results]
        for row in rows:
            manifest_sheet.append([_spreadsheet_text(row[field]) for field in TABULAR_FIELDS])
        _format_sheet(manifest_sheet, TABULAR_FIELDS, len(rows))

        failure_headers = (
            "group_id",
            "photo_id",
            "relative_path",
            "storage_key",
            "sha256",
            "byte_size",
            "error",
        )
        failures = [row for row in rows if row["status"] == "failed"]
        failure_sheet = workbook.create_sheet("Failures")
        failure_sheet.append(list(failure_headers))
        for row in failures:
            failure_sheet.append(
                [_spreadsheet_text(row[field]) for field in failure_headers]
            )
        _format_sheet(failure_sheet, failure_headers, len(failures))
        workbook.save(part)
        _fsync_file(part)
        os.replace(part, target)
    finally:
        part.unlink(missing_ok=True)


def _atomic_zip(report: ExportReport, target: Path) -> None:
    part = target.with_name(target.name + ".part")
    try:
        with ZipFile(part, mode="w", compression=ZIP_DEFLATED, allowZip64=True) as archive:
            for result in report.results:
                if not result.succeeded or not result.local_path:
                    continue
                archive.write(result.local_path, arcname=result.item.relative_path)
        _fsync_file(part)
        os.replace(part, target)
    finally:
        part.unlink(missing_ok=True)


def _atomic_report(report: ExportReport, target: Path) -> None:
    payload = {
        "schema": "module-manager-local-export-report/v1",
        "generated_at": report.generated_at,
        "output_root": str(report.output_root),
        "planned": report.planned,
        "planned_bytes": report.planned_bytes,
        "succeeded": report.succeeded,
        "failed": report.failed,
        "downloaded_bytes": report.downloaded_bytes,
        "complete": report.failed == 0,
        "items": [_sanitized_row(result) for result in report.results],
    }
    part = target.with_name(target.name + ".part")
    try:
        with part.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(part, target)
    finally:
        part.unlink(missing_ok=True)


def _validated_formats(formats: set[str] | None) -> set[str]:
    normalized = {"files"} if formats is None else set(formats)
    unknown = normalized - {"files", "zip", "csv", "xlsx"}
    if unknown:
        raise ValueError(f"unsupported export formats: {', '.join(sorted(unknown))}")
    return normalized


def run_export_stream(
    input_stream: TextIO,
    output_root: Path,
    formats: set[str] | None = None,
    *,
    max_workers: int = 4,
    opener: Callable[..., Any] = urlopen,
) -> ExportReport:
    """Consume the manifest header eagerly and item lines only under backpressure."""
    selected_formats = _validated_formats(formats)
    planned_count, planned_bytes = _parse_header(input_stream)
    _check_disk_space(output_root, planned_bytes, selected_formats)
    item_iterator = _iter_stream_items(
        input_stream,
        planned_count=planned_count,
        planned_bytes=planned_bytes,
    )
    report = download_manifest(
        item_iterator,
        output_root,
        max_workers=max_workers,
        retry_delays=RETRY_DELAYS,
        opener=opener,
        max_pending=MAX_PENDING_DOWNLOADS,
    )
    # The lazy iterator reconciles these values at EOF before this point.
    if report.planned != planned_count or report.planned_bytes != planned_bytes:
        raise ValueError("manifest preflight totals do not match downloaded items")

    output_root.mkdir(parents=True, exist_ok=True)
    if "zip" in selected_formats:
        _atomic_zip(report, output_root / "photos.zip")
    if "csv" in selected_formats:
        _atomic_csv(report, output_root / "manifest.csv")
    if "xlsx" in selected_formats:
        _atomic_xlsx(report, output_root / "manifest.xlsx")
    _atomic_report(report, output_root / "export-report.json")
    return report


def run_export(
    stdin_text: str,
    output_root: Path,
    formats: set[str] | None = None,
    *,
    max_workers: int = 4,
    opener: Callable[..., Any] = urlopen,
) -> ExportReport:
    return run_export_stream(
        io.StringIO(stdin_text),
        output_root,
        formats,
        max_workers=max_workers,
        opener=opener,
    )


def _default_output_root() -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return Path.home() / "Downloads" / "module-manager-exports" / stamp


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Download a streamed OSS export manifest to this computer."
    )
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument(
        "--format",
        dest="formats",
        action="append",
        choices=("files", "zip", "csv", "xlsx"),
    )
    parser.add_argument(
        "--max-workers", type=int, choices=range(1, MAX_DOWNLOAD_WORKERS + 1), default=4
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    output_root = args.output or _default_output_root()
    formats = set(args.formats) if args.formats else None
    try:
        report = run_export_stream(
            sys.stdin,
            output_root,
            formats,
            max_workers=args.max_workers,
        )
    except Exception as exc:
        print(f"local export preflight failed: {sanitize_error(exc)}", file=sys.stderr)
        return 3
    print(
        json.dumps(
            {
                "output_root": str(report.output_root),
                "planned": report.planned,
                "succeeded": report.succeeded,
                "failed": report.failed,
            },
            ensure_ascii=False,
        )
    )
    return 2 if report.failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
