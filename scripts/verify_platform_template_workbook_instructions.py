from __future__ import annotations

from io import BytesIO
import sys
from pathlib import Path
from tempfile import TemporaryDirectory


ROOT = Path(__file__).resolve().parents[1]
API_ROOT = ROOT / "v2-api"
sys.path.insert(0, str(API_ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from app.services.platform.catalog import configure_project_draft_store_path, reset_project_drafts  # noqa: E402


TERMINAL_SCHEMA = {
    "primary_field": {
        "key": "terminal_no",
        "label": "Terminal",
        "data_type": "text",
        "source": "import",
        "capture_method": "manual",
        "required": True,
        "show_in_construction_panel": True,
        "relation_role": "task_object",
    },
    "aggregate_field": {
        "key": "station_area",
        "label": "Station area",
        "data_type": "text",
        "source": "import",
        "capture_method": "manual",
        "required": True,
        "show_in_construction_panel": True,
        "relation_role": "aggregate",
    },
    "custom_fields": [
        {
            "key": "terminal_address",
            "label": "Terminal address",
            "data_type": "text",
            "source": "import",
            "capture_method": "manual",
            "required": True,
            "parent_key": "terminal_no",
            "show_in_construction_panel": True,
            "relation_role": "task_detail",
        },
        {
            "key": "new_terminal_no",
            "label": "New terminal",
            "data_type": "text",
            "source": "field_collection",
            "capture_method": "scan",
            "required": True,
            "parent_key": "terminal_no",
            "show_in_construction_panel": True,
            "relation_role": "replacement_device",
        },
        {
            "key": "communication_module_replace_confirm",
            "label": "Communication module replacement",
            "data_type": "enum",
            "source": "field_collection",
            "capture_method": "select",
            "required": True,
            "parent_key": "terminal_no",
            "show_in_construction_panel": True,
            "options": ["replace", "keep", "pending"],
            "relation_role": "accessory_replace_confirm",
        },
        {
            "key": "communication_module_no",
            "label": "New communication module",
            "data_type": "text",
            "source": "field_collection",
            "capture_method": "scan",
            "required": False,
            "parent_key": "terminal_no",
            "show_in_construction_panel": True,
            "relation_role": "accessory_new_device",
            "required_when": {"field_key": "communication_module_replace_confirm", "equals": "replace"},
        },
    ],
}


MODULE_SCHEMA = {
    "primary_field": {
        "key": "meter_no",
        "label": "Meter",
        "data_type": "text",
        "source": "import",
        "capture_method": "manual",
        "required": True,
        "show_in_construction_panel": True,
        "relation_role": "task_object",
    },
    "aggregate_field": {
        "key": "area_no",
        "label": "Station area",
        "data_type": "text",
        "source": "import",
        "capture_method": "manual",
        "required": True,
        "show_in_construction_panel": True,
        "relation_role": "aggregate",
    },
    "custom_fields": [
        {
            "key": "old_module_no",
            "label": "Old module",
            "data_type": "text",
            "source": "field_collection",
            "capture_method": "scan",
            "required": True,
            "parent_key": "meter_no",
            "show_in_construction_panel": True,
            "relation_role": "old_device",
        },
        {
            "key": "new_module_no",
            "label": "New module",
            "data_type": "text",
            "source": "field_collection",
            "capture_method": "scan",
            "required": True,
            "parent_key": "meter_no",
            "show_in_construction_panel": True,
            "relation_role": "accessory_new_device",
        },
        {
            "key": "collector_replace_confirm",
            "label": "Collector replacement",
            "data_type": "enum",
            "source": "field_collection",
            "capture_method": "select",
            "required": True,
            "parent_key": "meter_no",
            "show_in_construction_panel": True,
            "options": ["replace", "keep", "pending"],
            "relation_role": "accessory_replace_confirm",
        },
        {
            "key": "new_collector_no",
            "label": "New collector",
            "data_type": "text",
            "source": "field_collection",
            "capture_method": "scan",
            "required": False,
            "parent_key": "meter_no",
            "show_in_construction_panel": True,
            "relation_role": "accessory_new_device",
            "required_when": {"field_key": "collector_replace_confirm", "equals": "replace"},
        },
    ],
}


def fail(message: str) -> None:
    raise SystemExit(f"[FAIL] {message}")


def sheet_values(workbook_content: bytes, sheet_name: str) -> list[str]:
    workbook = load_workbook(BytesIO(workbook_content), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        fail(f"workbook must include {sheet_name!r} sheet; got {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    values: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                values.append(str(value))
    return values


def main() -> None:
    with TemporaryDirectory() as tmpdir:
        configure_project_draft_store_path(Path(tmpdir) / "platform-project-drafts.json")
        reset_project_drafts(remove_store=True)
        client = TestClient(app)

        scenarios = [
            (
                "Template Instructions Terminal Demo",
                TERMINAL_SCHEMA,
                [
                    "主设备更换后确认附属设备",
                    "先记录旧主设备和新主设备",
                    "Terminal",
                    "Communication module replacement",
                    "New communication module",
                ],
            ),
            (
                "Template Instructions Module Demo",
                MODULE_SCHEMA,
                [
                    "任务对象下更换附属设备",
                    "模块、采集器等附属设备挂在任务对象下",
                    "Meter",
                    "Old module",
                    "New module",
                    "Collector replacement",
                ],
            ),
        ]

        base_required_tokens = [
            "模板填写说明",
            "字段层级",
            "父字段",
            "条件采集",
            "上传时平台生成",
            "replace",
            "uploaded_at",
            "photo_count",
        ]
        for project_name, schema, scenario_tokens in scenarios:
            create_response = client.post(
                "/projects",
                json={
                    "name": project_name,
                    "module_ids": ["progress", "field", "review"],
                    "work_item_schema": schema,
                },
            )
            if create_response.status_code != 200:
                fail(f"project creation failed: {create_response.status_code} {create_response.text}")
            project_id = create_response.json()["data"]["id"]

            response = client.get(f"/projects/{project_id}/templates/external_completed")
            if response.status_code != 200:
                fail(f"template download failed: {response.status_code} {response.text}")

            values = sheet_values(response.content, "instructions")
            joined = "\n".join(values)
            for token in [*base_required_tokens, *scenario_tokens]:
                if token not in joined:
                    fail(f"instructions sheet missing token for {project_name}: {token}")

        reset_project_drafts(remove_store=True)
        configure_project_draft_store_path(None)

    print("[OK] platform template workbooks include hierarchy instructions.")


if __name__ == "__main__":
    main()
