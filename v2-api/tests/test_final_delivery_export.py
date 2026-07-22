from __future__ import annotations

import hashlib
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from threading import Event, Thread
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.services.final_delivery_export import (
    DeliveryPackageValidationError,
    build_delivery_package,
    build_delivery_workbook,
    cleanup_delivery_cache,
    delivery_evidence_fingerprint,
    get_or_build_delivery_package,
    release_delivery_cache_path,
    reserve_delivery_cache_path,
)


STANDARD_PHOTOS = (
    ("before_box", "表箱整体改造前"),
    ("collector_barcode", "采集器条形码"),
    ("module_meter", "模块与电能表"),
    ("after_box", "表箱整体改造后"),
)


def delivery_group(
    group_id: str = "group-001",
    *,
    terminal: str = "00123456",
    meter_no: str = "000012345678",
    module_no: str = "000098765432",
    collector: str = "00007777",
    address: str = "江宁路 1 号",
    completed_at: str = "2026-07-21T16:30:00Z",
) -> dict:
    photos = []
    for index, (category, _label) in enumerate(STANDARD_PHOTOS, start=1):
        content = f"photo-{group_id}-{category}".encode()
        photos.append(
            {
                "id": f"{group_id}-photo-{index}",
                "category": category,
                "sha256": hashlib.sha256(content).hexdigest(),
                "delivery_cache_content_sha256": hashlib.sha256(content).hexdigest(),
                "delivery_cache_path": f"objects/{index:02d}/{category}.jpeg",
                "delivery_cache_status": "ready",
                "delivery_cache_version": f"version-{index}",
                "original_filename": f"capture-{category}.jpeg",
                "is_active": True,
                "_content": content,
            }
        )
    return {
        "id": group_id,
        "terminal": terminal,
        "meter_no": meter_no,
        "module_asset_no": module_no,
        "collector": collector,
        "address": address,
        "client_completed_at": completed_at,
        "status": "archived",
        "photos": photos,
    }


def read_photo(photo: dict) -> bytes:
    return photo["_content"]


def test_delivery_workbook_has_exact_dual_sheet_contract_and_text_cells() -> None:
    group = delivery_group()

    workbook = load_workbook(BytesIO(build_delivery_workbook([group, dict(group)])))

    assert workbook.sheetnames == ["老设备", "新设备"]
    old_sheet = workbook["老设备"]
    new_sheet = workbook["新设备"]
    assert [cell.value for cell in old_sheet[1]] == [
        "终端地址",
        "终端厂家",
        "终端安装地址",
        "表号",
        "采集器设备号",
        "供服中心",
        "改造工程队",
        "旧设备拆除时间",
        "旧设备状态",
    ]
    assert [cell.value for cell in new_sheet[1]] == ["改造日期", "模块对应的表号", "新模块编号", "备注"]
    assert old_sheet.max_row == 2
    assert new_sheet.max_row == 2
    assert [cell.value for cell in old_sheet[2]] == [
        "00123456",
        None,
        "江宁路 1 号",
        "000012345678",
        "00007777",
        "南大供电服务中心",
        "奕福",
        None,
        None,
    ]
    assert [cell.value for cell in new_sheet[2]] == [
        "2026-07-22",
        "000012345678",
        "000098765432",
        None,
    ]
    for sheet in (old_sheet, new_sheet):
        for row in sheet.iter_rows(min_row=2):
            assert all(cell.number_format == "@" for cell in row)
    assert all(old_sheet.cell(2, column).data_type == "s" for column in (1, 4, 5))
    assert all(new_sheet.cell(2, column).data_type == "s" for column in (1, 2, 3))


def test_delivery_workbook_is_pure_and_preserves_replacement_remark() -> None:
    group = delivery_group()
    group["replacement_old_meter_no"] = "000000001111"
    for photo in group["photos"]:
        photo.pop("delivery_cache_path")
        photo.pop("delivery_cache_status")

    workbook = load_workbook(BytesIO(build_delivery_workbook([group])))

    assert workbook["新设备"]["D2"].value == "换表：旧表号 000000001111"


def test_delivery_package_uses_exact_four_category_paths_and_original_extensions() -> None:
    group = delivery_group()

    package = build_delivery_package([group], read_photo)

    with ZipFile(BytesIO(package)) as archive:
        assert archive.namelist() == [
            "设备清单.xlsx",
            "00123456/000012345678-000098765432-江宁路 1 号/000098765432-表箱整体改造前.jpeg",
            "00123456/000012345678-000098765432-江宁路 1 号/000098765432-采集器条形码.jpeg",
            "00123456/000012345678-000098765432-江宁路 1 号/000098765432-模块与电能表.jpeg",
            "00123456/000012345678-000098765432-江宁路 1 号/000098765432-表箱整体改造后.jpeg",
        ]
        for photo, name in zip(group["photos"], archive.namelist()[1:], strict=True):
            assert archive.read(name) == photo["_content"]


def test_delivery_package_sanitizes_windows_paths_and_suffixes_sanitized_collisions() -> None:
    first = delivery_group(
        "stable/group:one",
        terminal="T:01.",
        meter_no="M:01",
        module_no="MOD:01",
        collector="C:01",
        address="一号楼? ",
    )
    second = delivery_group(
        "stable-group-two",
        terminal="T?01.",
        meter_no="M?01",
        module_no="MOD?01",
        collector="C?01",
        address="一号楼* ",
    )

    package = build_delivery_package([first, second], read_photo)

    with ZipFile(BytesIO(package)) as archive:
        photo_names = archive.namelist()[1:]
    assert len(photo_names) == 8
    assert len(set(photo_names)) == 8
    assert all(not any(character in name for character in '<>:"\\|?*\x00') for name in photo_names)
    assert any("stable-group-two" in name for name in photo_names)
    assert all(not part.endswith((" ", ".")) for name in photo_names for part in name.split("/"))


@pytest.mark.parametrize(
    ("mutate", "expected_code"),
    [
        (lambda group: group.update(terminal=""), "missing_terminal"),
        (lambda group: group.update(meter_no=""), "missing_meter_no"),
        (lambda group: group.update(module_asset_no=""), "missing_module_no"),
        (lambda group: group.update(collector=""), "missing_collector"),
        (lambda group: group.update(address=""), "missing_address"),
        (lambda group: group.update(client_completed_at="not-a-date"), "invalid_completed_at"),
        (lambda group: group["photos"].pop(), "invalid_photo_count"),
        (lambda group: group["photos"][1].update(category="before_box"), "invalid_photo_categories"),
        (lambda group: group.update(status="approved"), "not_archived"),
        (lambda group: group.update(terminal="00000000"), "placeholder_identity"),
        (lambda group: group["photos"][0].update(delivery_cache_status="pending"), "delivery_cache_pending"),
    ],
)
def test_delivery_package_blocks_invalid_groups_with_structured_errors(mutate, expected_code: str) -> None:
    group = delivery_group()
    mutate(group)
    reads = 0

    def guarded_reader(photo: dict) -> bytes:
        nonlocal reads
        reads += 1
        return read_photo(photo)

    with pytest.raises(DeliveryPackageValidationError) as captured:
        build_delivery_package([group], guarded_reader)

    assert reads == 0
    assert captured.value.errors
    assert captured.value.errors[0]["group_id"] == "group-001"
    assert expected_code in {error["code"] for error in captured.value.errors}


def test_delivery_package_reports_cross_group_device_conflicts_without_partial_output() -> None:
    first = delivery_group("group-001")
    second = delivery_group("group-002", address="江宁路 2 号")

    with pytest.raises(DeliveryPackageValidationError) as captured:
        build_delivery_package([first, second], read_photo)

    conflicts = [error for error in captured.value.errors if error["code"] == "device_conflict"]
    assert {error["group_id"] for error in conflicts} == {"group-001", "group-002"}
    assert {error["field"] for error in conflicts} == {"terminal", "meter_no", "module_asset_no", "collector"}


def test_delivery_package_rejects_an_empty_formal_scope() -> None:
    with pytest.raises(DeliveryPackageValidationError) as captured:
        build_delivery_package([], read_photo)

    assert captured.value.errors == [
        {
            "group_id": "",
            "code": "no_groups",
            "field": "scope",
            "message": "No archived groups are eligible for formal delivery",
        }
    ]


def test_get_or_build_delivery_package_reuses_same_fingerprint_for_seven_days(tmp_path: Path) -> None:
    group = delivery_group()
    reads = 0
    builds = 0

    def counting_reader(photo: dict) -> bytes:
        nonlocal reads
        reads += 1
        return read_photo(photo)

    def counting_builder(groups, reader) -> bytes:
        nonlocal builds
        builds += 1
        return build_delivery_package(groups, reader)

    now = datetime(2026, 7, 22, 8, tzinfo=UTC)
    first = get_or_build_delivery_package(
        "team-a/task-1",
        "a" * 64,
        groups=[group],
        photo_reader=counting_reader,
        cache_root=tmp_path,
        now=now,
        package_builder=counting_builder,
    )
    second = get_or_build_delivery_package(
        "team-a/task-1",
        "a" * 64,
        groups=[group],
        photo_reader=counting_reader,
        cache_root=tmp_path,
        now=now + timedelta(days=6, hours=23),
        package_builder=counting_builder,
    )

    assert first == second
    assert first.is_file()
    assert builds == 1
    assert reads == 4


def test_get_or_build_delivery_package_rebuilds_after_ttl_or_fingerprint_change(tmp_path: Path) -> None:
    group = delivery_group()
    builds = 0

    def counting_builder(groups, reader) -> bytes:
        nonlocal builds
        builds += 1
        return build_delivery_package(groups, reader)

    now = datetime(2026, 7, 22, 8, tzinfo=UTC)
    first = get_or_build_delivery_package(
        "task-1",
        "a" * 64,
        groups=[group],
        photo_reader=read_photo,
        cache_root=tmp_path,
        now=now,
        package_builder=counting_builder,
    )
    changed = get_or_build_delivery_package(
        "task-1",
        "b" * 64,
        groups=[group],
        photo_reader=read_photo,
        cache_root=tmp_path,
        now=now + timedelta(hours=1),
        package_builder=counting_builder,
    )
    rebuilt = get_or_build_delivery_package(
        "task-1",
        "a" * 64,
        groups=[group],
        photo_reader=read_photo,
        cache_root=tmp_path,
        now=now + timedelta(days=8),
        package_builder=counting_builder,
    )

    assert first != changed
    assert rebuilt == first
    assert builds == 3


def test_cached_package_never_bypasses_current_group_validation(tmp_path: Path) -> None:
    group = delivery_group()
    now = datetime(2026, 7, 22, 8, tzinfo=UTC)
    path = get_or_build_delivery_package(
        "task-1",
        "a" * 64,
        groups=[group],
        photo_reader=read_photo,
        cache_root=tmp_path,
        now=now,
    )
    group["photos"][0]["delivery_cache_status"] = "pending"

    with pytest.raises(DeliveryPackageValidationError) as captured:
        get_or_build_delivery_package(
            "task-1",
            "a" * 64,
            groups=[group],
            photo_reader=lambda _photo: (_ for _ in ()).throw(AssertionError("must not read photos")),
            cache_root=tmp_path,
            now=now + timedelta(hours=1),
        )

    assert path.is_file()
    assert {error["code"] for error in captured.value.errors} == {"delivery_cache_pending"}


def test_identity_category_and_photo_changes_invalidate_delivery_fingerprint() -> None:
    group = delivery_group()
    original = delivery_evidence_fingerprint([group])
    group["module_asset_no"] = "000098765433"
    identity_changed = delivery_evidence_fingerprint([group])
    group["photos"][0]["category"] = "collector_barcode"
    category_changed = delivery_evidence_fingerprint([group])
    group["photos"][0]["sha256"] = "f" * 64
    group["photos"][0]["delivery_cache_content_sha256"] = "f" * 64
    photo_changed = delivery_evidence_fingerprint([group])

    assert len({original, identity_changed, category_changed, photo_changed}) == 4


def test_delivery_lru_keeps_referenced_and_active_paths(tmp_path: Path) -> None:
    now = datetime(2026, 7, 22, 8, tzinfo=UTC)
    referenced = tmp_path / "objects" / "aa" / "referenced.jpg"
    active_object = tmp_path / "objects" / "bb" / "active.jpg"
    removable = tmp_path / "objects" / "cc" / "removable.jpg"
    active_package = tmp_path / "packages" / "scope" / "active.zip"
    for path, content in (
        (referenced, b"referenced"),
        (active_object, b"active"),
        (removable, b"remove-me"),
        (active_package, b"package"),
    ):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        old = (now - timedelta(days=8)).timestamp()
        path.touch()
        import os

        os.utime(path, (old, old))
    group = {
        "delivery_cache_status": "ready",
        "photos": [
            {
                "is_active": True,
                "delivery_cache_status": "ready",
                "delivery_cache_path": str(referenced.relative_to(tmp_path)).replace("\\", "/"),
            }
        ],
    }
    object_lease = reserve_delivery_cache_path(active_object)
    package_lease = reserve_delivery_cache_path(active_package)
    try:
        result = cleanup_delivery_cache(tmp_path, max_object_bytes=0, groups=[group], now=now)

        assert referenced.exists()
        assert active_object.exists()
        assert active_package.exists()
        assert not removable.exists()
        assert result["protected_objects"] == 2
    finally:
        release_delivery_cache_path(object_lease)
        release_delivery_cache_path(package_lease)

    cleanup_delivery_cache(tmp_path, max_object_bytes=0, groups=[], now=now)

    assert not referenced.exists()
    assert not active_object.exists()
    assert not active_package.exists()


def test_delivery_content_builder_blocks_lru_for_its_object_store(tmp_path: Path) -> None:
    from app.services.delivery_cache import cache_group_photos

    old_object = tmp_path / "objects" / "old" / "old.jpg"
    old_object.parent.mkdir(parents=True, exist_ok=True)
    old_object.write_bytes(b"old")
    content = b"new-content"
    started = Event()
    finish = Event()
    group = {
        "id": "building-group",
        "photos": [
            {
                "id": "building-photo",
                "sha256": hashlib.sha256(content).hexdigest(),
                "sha256_source": "declared",
                "category": "before_box",
                "original_filename": "before.jpg",
                "is_active": True,
            }
        ],
    }

    def fetch_photo(_photo):
        started.set()
        assert finish.wait(5)
        return content, ".jpg", "image/jpeg"

    worker = Thread(
        target=cache_group_photos,
        kwargs={"group": group, "cache_root": tmp_path, "fetch_photo": fetch_photo},
    )
    worker.start()
    assert started.wait(5)
    try:
        cleanup_delivery_cache(tmp_path, max_object_bytes=0, groups=[])
        assert old_object.exists()
    finally:
        finish.set()
        worker.join(5)
    assert not worker.is_alive()
