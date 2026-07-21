from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.services.construction_priority_import import (
    PriorityImportError,
    build_priority_template,
    parse_priority_workbook,
)


def workbook_bytes(rows, *, headers=("终端号", "优先施工"), formula=False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for terminal, priority in rows:
        sheet.append([terminal, priority])
    if formula:
        sheet["A2"] = "=CONCAT(\"350000000001\")"
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def test_priority_template_uses_exact_chinese_headers_and_text_validation() -> None:
    workbook = load_workbook(BytesIO(build_priority_template()), data_only=False)
    sheet = workbook.active

    assert [sheet.cell(1, 1).value, sheet.cell(1, 2).value] == ["终端号", "优先施工"]
    assert sheet["A2"].number_format == "@"
    assert any("是,否" in str(rule.formula1) for rule in sheet.data_validations.dataValidation)


def test_priority_parser_merges_same_values_and_blocks_conflicts() -> None:
    rows = parse_priority_workbook(
        workbook_bytes(
            [
                ("350000000001", "是"),
                ("350000000001", "是"),
                ("350000000002", "是"),
                ("350000000002", "否"),
            ]
        )
    )

    assert [row.status for row in rows] == ["valid", "duplicate", "conflict", "conflict"]
    assert [row.terminal for row in rows] == ["350000000001", "350000000001", "350000000002", "350000000002"]


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (workbook_bytes([], headers=("终端", "优先施工")), "headers"),
        (b"not an xlsx", "workbook"),
    ],
)
def test_priority_parser_rejects_malformed_workbooks(content: bytes, message: str) -> None:
    with pytest.raises(PriorityImportError, match=message):
        parse_priority_workbook(content)


def test_priority_parser_marks_formula_and_invalid_rows_malformed() -> None:
    rows = parse_priority_workbook(
        workbook_bytes(
            [("", "是"), ("350000000001", "maybe"), ("350000000002", "是")],
            formula=True,
        )
    )

    assert [row.status for row in rows] == ["malformed", "malformed", "valid"]


def test_priority_parser_rejects_more_than_5000_rows_without_coercing_terminals() -> None:
    content = workbook_bytes([(f"350000{i:06d}", "是") for i in range(5001)])

    with pytest.raises(PriorityImportError, match="5000"):
        parse_priority_workbook(content)
