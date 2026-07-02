import json
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import app
from app.services.platform import catalog
from app.services.platform.adapters.replacement import build_replacement_project_overview as build_adapter_overview
from app.services.platform.catalog import (
    ProjectConfigurationError,
    ProjectDefinition,
    ProjectNotFound,
    configure_project_draft_store_path,
    get_project_module_definition,
    get_project_overview,
    get_project_section,
    list_project_definitions,
    list_project_modules,
    reset_project_drafts,
)
from app.services.platform.delivery import build_delivery_summary
from app.services.platform.field import build_field_summary
from app.services.platform.overview import build_replacement_project_overview
from app.services.platform.progress import build_progress_summary
from app.services.platform.projects import build_project_overview
from app.services.platform.review import build_review_summary
from app.services.platform.risks import build_risk_summary
from app.services.platform.tasks import build_task_summary


@pytest.fixture(autouse=True)
def clear_project_drafts(tmp_path):
    configure_project_draft_store_path(tmp_path / "platform-project-drafts.json")
    reset_project_drafts(remove_store=True)
    yield
    reset_project_drafts(remove_store=True)
    configure_project_draft_store_path(None)


def test_platform_catalog_lists_replacement_project_definition():
    definitions = list_project_definitions()

    assert [item["id"] for item in definitions] == ["replacement-project"]
    assert definitions[0]["name"] == "更换模块项目"
    assert definitions[0]["status"] == "active"


def test_platform_catalog_loads_project_overview_and_sections():
    overview = get_project_overview("replacement-project")
    progress = get_project_section("replacement-project", "progress")
    tasks = get_project_section("replacement-project", "tasks")

    assert overview["id"] == "replacement-project"
    assert progress["stage"] == overview["stage"]
    assert tasks["total"] == overview["tasks"]["total"]


def test_platform_catalog_rejects_unknown_project_and_section():
    with pytest.raises(ProjectNotFound):
        get_project_overview("unknown-project")
    with pytest.raises(KeyError):
        get_project_section("replacement-project", "unknown-section")


def test_project_scoped_business_requests_accept_registered_project():
    client = TestClient(app)
    login = client.post("/auth/login", json={"username": "admin", "password": "admin123"})
    headers = {"Authorization": f"bearer {login.json()['data']['access_token']}"}

    response = client.get("/local-test/system/status?project_id=replacement-project", headers=headers)

    assert response.status_code == 200
    assert response.json()["data"]["project_id"] == "replacement-project"


def test_project_scoped_business_requests_reject_unknown_project():
    client = TestClient(app)
    login = client.post("/auth/login", json={"username": "admin", "password": "admin123"})
    headers = {"Authorization": f"bearer {login.json()['data']['access_token']}"}

    response = client.get("/local-test/system/status?project_id=unknown-project", headers=headers)

    assert response.status_code == 404
    assert response.json()["detail"] == "Project not found"


def test_platform_progress_module_calculates_stage_and_progress():
    progress = build_progress_summary(
        groups=10,
        reviewed_groups=6,
        risk_total=5,
        unconstructed_groups=3,
        reviewing_tasks=1,
    )

    assert progress["stage"] == "审阅中"
    assert progress["system_progress"] == 60
    assert progress["management_progress"] == 60
    assert progress["management_locked"] is False


def test_platform_summary_modules_keep_independent_boundaries():
    tasks = build_task_summary(
        {
            "total": 4,
            "uploaded": 2,
            "reviewing": 1,
            "archived": 1,
            "avg_upload_rate": 0.5,
            "avg_review_rate": 0.6,
        }
    )
    field = build_field_summary(photo_rows_linked=20, unconstructed_groups=3, exception_groups=2)
    review = build_review_summary(groups=10, reviewed_groups=6, progress=60)
    risks = build_risk_summary(exception_groups=2, unconstructed_groups=3, delivery_ready=False)
    delivery = build_delivery_summary(progress=60, risk_total=risks["total"])

    assert tasks["upload_rate"] == 50
    assert field["exception_count"] == 2
    assert review["pending_groups"] == 4
    assert risks["delivery_blockers"] == 5
    assert delivery["completed_items"] == 2


def test_project_module_builds_overview_envelope_from_focused_summaries():
    overview = build_project_overview(
        project_id="demo-project",
        name="演示项目",
        status="active",
        progress={"stage": "施工中", "system_progress": 60, "management_progress": 60, "management_locked": False},
        total_groups=10,
        completed_groups=6,
        exception_groups=2,
        tasks={"total": 4},
        delivery={"status": "preparing"},
        field={"exception_count": 2},
        review={"pending_groups": 4},
        risks={"total": 5},
        updated_at="2026-06-29T00:00:00+00:00",
    )

    assert overview["id"] == "demo-project"
    assert overview["name"] == "演示项目"
    assert overview["stage"] == "施工中"
    assert overview["delivery"]["status"] == "preparing"
    assert overview["risks"]["total"] == 5


def test_replacement_adapter_matches_public_overview_facade():
    summary = {"groups": 5, "reviewed_groups": 5, "exception_groups": 0, "unconstructed_groups": 0}
    task_status = {"total": 2, "uploaded": 2, "reviewing": 0, "archived": 2, "avg_upload_rate": 1, "avg_review_rate": 1}

    adapter_overview = build_adapter_overview(summary=summary, task_status=task_status)
    facade_overview = build_replacement_project_overview(summary=summary, task_status=task_status)

    adapter_overview["updated_at"] = "stable"
    facade_overview["updated_at"] = "stable"
    assert adapter_overview == facade_overview


def test_build_replacement_project_overview_maps_summary_to_platform_fields():
    overview = build_replacement_project_overview(
        summary={
            "groups": 10,
            "reviewed_groups": 6,
            "exception_groups": 2,
            "unconstructed_groups": 3,
            "photo_rows_linked": 20,
        },
        task_status={
            "total": 4,
            "uploaded": 2,
            "reviewing": 1,
            "archived": 1,
            "avg_upload_rate": 0.5,
            "avg_review_rate": 0.6,
        },
    )

    assert overview["id"] == "replacement-project"
    assert overview["name"] == "更换模块项目"
    assert overview["stage"] == "审阅中"
    assert overview["system_progress"] == 60
    assert overview["management_progress"] == 60
    assert overview["delivery"]["total_items"] == 4
    assert overview["field"]["exception_count"] == 2
    assert overview["review"]["reviewed_groups"] == 6
    assert overview["risks"]["total"] == 5
    assert overview["tasks"]["total"] == 4


def test_build_replacement_project_overview_marks_delivered_when_review_complete_and_no_risk():
    overview = build_replacement_project_overview(
        summary={
            "groups": 5,
            "reviewed_groups": 5,
            "exception_groups": 0,
            "unconstructed_groups": 0,
        },
        task_status={
            "total": 2,
            "uploaded": 2,
            "reviewing": 0,
            "archived": 2,
            "avg_upload_rate": 1,
            "avg_review_rate": 1,
        },
    )

    assert overview["stage"] == "待验收"
    assert overview["system_progress"] == 100
    assert overview["delivery"]["status"] == "ready"
    assert overview["risks"]["total"] == 0


def test_projects_list_returns_replacement_platform_project():
    client = TestClient(app)
    response = client.get("/projects")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["total"] == 1
    assert payload["items"][0]["id"] == "replacement-project"
    assert "delivery" in payload["items"][0]
    assert "risks" in payload["items"][0]


def test_project_detail_returns_platform_project():
    client = TestClient(app)
    response = client.get("/projects/replacement-project")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["id"] == "replacement-project"
    assert "tasks" in payload
    assert [module["id"] for module in payload["modules"][:4]] == [
        "progress",
        "delivery",
        "field",
        "review",
    ]
    assert payload["modules"][0]["endpoint"] == "/projects/replacement-project/modules/progress"
    assert payload["modules"][0]["route_path"] == "/project-board"


def test_project_modules_endpoint_lists_registered_sections_by_priority():
    client = TestClient(app)
    response = client.get("/projects/replacement-project/modules")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["project_id"] == "replacement-project"
    assert [module["id"] for module in payload["items"]] == [
        "progress",
        "delivery",
        "field",
        "review",
        "risks",
        "tasks",
    ]
    assert {module["id"]: module["route_path"] for module in payload["items"]} == {
        "progress": "/project-board",
        "delivery": "/project-board",
        "field": "/construction",
        "review": "/task-hall",
        "risks": "/project-board",
        "tasks": "/claim-tasks",
    }


def test_project_modules_are_limited_by_project_definition(monkeypatch):
    monkeypatch.setitem(
        catalog._PROJECT_BY_ID,
        "limited-project",
        ProjectDefinition(
            id="limited-project",
            name="模块裁剪项目",
            status="active",
            adapter="replacement",
            module_ids=("review", "progress"),
        ),
    )

    modules = list_project_modules("limited-project")
    overview = get_project_overview("limited-project")

    assert [module["id"] for module in modules] == ["progress", "review"]
    assert overview["id"] == "limited-project"
    assert overview["name"] == "模块裁剪项目"
    assert [module["id"] for module in overview["modules"]] == ["progress", "review"]
    assert "review" in overview
    assert "delivery" not in overview
    assert "field" not in overview
    assert "risks" not in overview
    assert "tasks" not in overview
    assert get_project_module_definition("limited-project", "review").id == "review"
    with pytest.raises(KeyError):
        get_project_module_definition("limited-project", "delivery")


def test_project_modules_report_invalid_project_configuration(monkeypatch):
    monkeypatch.setitem(
        catalog._PROJECT_BY_ID,
        "invalid-module-project",
        ProjectDefinition(
            id="invalid-module-project",
            name="错误模块项目",
            status="active",
            adapter="replacement",
            module_ids=("progress", "unknown-module"),
        ),
    )

    with pytest.raises(ProjectConfigurationError, match="unknown-module"):
        list_project_modules("invalid-module-project")

    client = TestClient(app)
    response = client.get("/projects/invalid-module-project/modules")

    assert response.status_code == 500
    assert response.json()["detail"] == "Project configuration invalid"


def test_projects_list_reports_invalid_registered_project_configuration(monkeypatch):
    invalid_project = ProjectDefinition(
        id="invalid-listed-project",
        name="列表错误项目",
        status="active",
        adapter="replacement",
        module_ids=("progress", "unknown-module"),
    )
    monkeypatch.setattr(catalog, "_PROJECTS", (invalid_project,))
    monkeypatch.setitem(catalog._PROJECT_BY_ID, invalid_project.id, invalid_project)

    client = TestClient(app)
    response = client.get("/projects")

    assert response.status_code == 500
    assert response.json()["detail"] == "Project configuration invalid"


def test_project_overview_can_disable_progress_module(monkeypatch):
    monkeypatch.setitem(
        catalog._PROJECT_BY_ID,
        "review-only-project",
        ProjectDefinition(
            id="review-only-project",
            name="只审阅项目",
            status="active",
            adapter="replacement",
            module_ids=("review",),
        ),
    )

    overview = get_project_overview("review-only-project")

    assert [module["id"] for module in overview["modules"]] == ["review"]
    assert "review" in overview
    assert "stage" not in overview
    assert "system_progress" not in overview
    assert "management_progress" not in overview
    assert "management_locked" not in overview
    with pytest.raises(KeyError):
        get_project_section("review-only-project", "progress")


@pytest.mark.parametrize(
    ("section", "expected_keys"),
    [
        ("progress", {"stage", "system_progress", "management_progress", "management_locked"}),
        ("delivery", {"status", "total_items", "completed_items"}),
        ("field", {"photo_rows_linked", "unconstructed_groups", "exception_count"}),
        ("review", {"reviewed_groups", "review_rate", "pending_groups"}),
        ("risks", {"total", "field_exceptions", "delivery_blockers"}),
        ("tasks", {"total", "uploaded", "reviewing", "archived"}),
    ],
)
def test_project_module_endpoints_return_focused_sections(section, expected_keys):
    client = TestClient(app)
    response = client.get(f"/projects/replacement-project/{section}")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert set(payload).issuperset(expected_keys)


def test_project_module_endpoint_returns_section_through_unified_modules_path():
    client = TestClient(app)
    response = client.get("/projects/replacement-project/modules/delivery")

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["status"] in {"preparing", "ready"}
    assert {"total_items", "completed_items"}.issubset(payload)


def test_project_module_endpoints_reject_unknown_project():
    client = TestClient(app)
    response = client.get("/projects/unknown-project/progress")

    assert response.status_code == 404


def test_create_project_draft_registers_project_with_selected_modules():
    client = TestClient(app)
    response = client.post(
        "/projects",
        json={
            "name": "  线路巡检项目  ",
            "description": "用于巡检流程接入",
            "module_ids": ["review", "progress", "field", "progress"],
        },
    )

    assert response.status_code == 200
    draft = response.json()["data"]
    assert draft["id"] == "xian-lu-xun-jian-xiang-mu"
    assert draft["name"] == "线路巡检项目"
    assert draft["description"] == "用于巡检流程接入"
    assert draft["status"] == "draft"
    assert [module["id"] for module in draft["modules"]] == ["progress", "field", "review"]
    assert draft["stage"] == "准备中"
    assert draft["system_progress"] == 0
    assert draft["field"]["photo_rows_linked"] == 0
    assert draft["review"]["pending_groups"] == 0
    assert "delivery" not in draft

    list_response = client.get("/projects")
    assert list_response.status_code == 200
    assert "xian-lu-xun-jian-xiang-mu" in [
        item["id"] for item in list_response.json()["data"]["items"]
    ]


def test_project_draft_module_endpoints_follow_selected_modules():
    client = TestClient(app)
    create_response = client.post(
        "/projects",
        json={"name": "审阅专项", "module_ids": ["review"]},
    )
    project_id = create_response.json()["data"]["id"]

    modules_response = client.get(f"/projects/{project_id}/modules")
    assert modules_response.status_code == 200
    assert [module["id"] for module in modules_response.json()["data"]["items"]] == ["review"]

    review_response = client.get(f"/projects/{project_id}/modules/review")
    assert review_response.status_code == 200
    assert review_response.json()["data"]["reviewed_groups"] == 0

    progress_response = client.get(f"/projects/{project_id}/modules/progress")
    assert progress_response.status_code == 404


def test_project_draft_registry_persists_created_projects(tmp_path):
    store_path = tmp_path / "platform-project-drafts.json"
    configure_project_draft_store_path(store_path)
    client = TestClient(app)

    response = client.post(
        "/projects",
        json={"name": "Persistent Draft", "description": "kept on disk", "module_ids": ["progress", "review"]},
    )

    assert response.status_code == 200
    payload = json.loads(store_path.read_text(encoding="utf-8"))
    assert payload["version"] == 1
    assert payload["projects"][0]["id"] == "persistent-draft"
    assert payload["projects"][0]["name"] == "Persistent Draft"
    assert payload["projects"][0]["description"] == "kept on disk"
    assert payload["projects"][0]["module_ids"] == ["progress", "review"]


def test_create_project_draft_keeps_work_item_schema_with_required_platform_fields():
    client = TestClient(app)
    response = client.post(
        "/projects",
        json={
            "name": "Meter Replacement Field Model",
            "module_ids": ["progress", "field", "review"],
            "work_item_schema": {
                "primary_field": {
                    "key": "meter_no",
                    "label": "Electric meter",
                    "data_type": "text",
                    "source": "import",
                    "capture_method": "scan",
                    "required": True,
                },
                "aggregate_field": {
                    "key": "terminal",
                    "label": "Terminal",
                    "data_type": "text",
                    "source": "import",
                    "capture_method": "manual",
                    "required": True,
                },
                "custom_fields": [
                    {
                        "key": "module_asset_no",
                        "label": "Module asset number",
                        "data_type": "text",
                        "source": "field_collection",
                        "capture_method": "scan",
                        "required": True,
                        "parent_key": "meter_no",
                    },
                    {
                        "key": "collector_no",
                        "label": "Collector number",
                        "data_type": "text",
                        "source": "field_collection",
                        "capture_method": "scan",
                        "required": False,
                        "parent_key": "meter_no",
                    },
                ],
            },
        },
    )

    assert response.status_code == 200
    schema = response.json()["data"]["work_item_schema"]
    assert schema["primary_field"]["key"] == "meter_no"
    assert schema["aggregate_field"]["key"] == "terminal"
    assert [field["key"] for field in schema["custom_fields"]] == ["module_asset_no", "collector_no"]
    required_keys = {field["key"] for field in schema["platform_required_fields"]}
    assert {"installer", "completed_at", "online_duration_minutes", "photo_count"}.issubset(required_keys)
    required_by_key = {field["key"]: field for field in schema["platform_required_fields"]}
    assert required_by_key["installer"]["label"] == "安装人员"
    assert required_by_key["completed_at"]["label"] == "完成时间"
    assert required_by_key["online_duration_minutes"]["label"] == "在线时长（分钟）"
    assert required_by_key["photo_count"]["label"] == "照片数量"
    kpi_keys = {field["key"] for field in schema["platform_required_fields"] if field["kpi_enabled"]}
    assert {"installer", "completed_at", "online_duration_minutes"}.issubset(kpi_keys)


def test_project_draft_registry_persists_work_item_schema(tmp_path):
    store_path = tmp_path / "platform-project-drafts.json"
    configure_project_draft_store_path(store_path)
    client = TestClient(app)

    response = client.post(
        "/projects",
        json={
            "name": "Terminal Swap Field Model",
            "module_ids": ["progress", "field"],
            "work_item_schema": {
                "primary_field": {"key": "terminal_no", "label": "Terminal", "source": "import"},
                "aggregate_field": {"key": "area", "label": "Area", "source": "import"},
                "custom_fields": [
                    {"key": "carrier_module", "label": "Carrier module", "source": "field_collection"}
                ],
            },
        },
    )

    assert response.status_code == 200
    payload = json.loads(store_path.read_text(encoding="utf-8"))
    stored_schema = payload["projects"][0]["work_item_schema"]
    assert stored_schema["primary_field"]["key"] == "terminal_no"
    assert stored_schema["aggregate_field"]["key"] == "area"
    assert stored_schema["custom_fields"][0]["key"] == "carrier_module"
    assert any(field["key"] == "installer" for field in stored_schema["platform_required_fields"])


def test_project_draft_work_item_schema_can_be_updated_and_persisted(tmp_path):
    store_path = tmp_path / "platform-project-drafts.json"
    configure_project_draft_store_path(store_path)
    client = TestClient(app)
    create_response = client.post(
        "/projects",
        json={
            "name": "Editable Field Model",
            "module_ids": ["progress", "field", "review"],
            "work_item_schema": {
                "primary_field": {"key": "meter_no", "label": "电能表", "source": "import"},
                "aggregate_field": {"key": "area_no", "label": "台区", "source": "import"},
                "custom_fields": [
                    {"key": "module_asset_no", "label": "模块", "source": "field_collection"}
                ],
            },
        },
    )
    project_id = create_response.json()["data"]["id"]

    update_response = client.patch(
        f"/projects/{project_id}/work-item-schema",
        json={
            "primary_field": {"key": "terminal_no", "label": "终端", "source": "import"},
            "aggregate_field": {"key": "station_area", "label": "台区", "source": "import"},
            "custom_fields": [
                {
                    "key": "communication_module",
                    "label": "通讯模块",
                    "source": "field_collection",
                    "capture_method": "scan",
                    "required": True,
                    "parent_key": "terminal_no",
                },
                {
                    "key": "carrier_module",
                    "label": "载波模块",
                    "source": "field_collection",
                    "capture_method": "photo",
                    "data_type": "image",
                    "parent_key": "terminal_no",
                },
            ],
        },
    )

    assert update_response.status_code == 200
    updated_schema = update_response.json()["data"]["work_item_schema"]
    assert updated_schema["primary_field"]["key"] == "terminal_no"
    assert updated_schema["primary_field"]["label"] == "终端"
    assert [field["key"] for field in updated_schema["custom_fields"]] == [
        "communication_module",
        "carrier_module",
    ]
    assert any(field["key"] == "installer" for field in updated_schema["platform_required_fields"])
    assert get_project_overview(project_id)["work_item_schema"]["primary_field"]["label"] == "终端"

    payload = json.loads(store_path.read_text(encoding="utf-8"))
    persisted_schema = payload["projects"][0]["work_item_schema"]
    assert persisted_schema["primary_field"]["key"] == "terminal_no"
    assert persisted_schema["custom_fields"][1]["capture_method"] == "photo"


def test_project_draft_registry_recovers_projects_after_memory_reset(tmp_path):
    store_path = tmp_path / "platform-project-drafts.json"
    configure_project_draft_store_path(store_path)
    client = TestClient(app)
    create_response = client.post(
        "/projects",
        json={"name": "Restart Safe Draft", "module_ids": ["field", "review"]},
    )
    project_id = create_response.json()["data"]["id"]

    reset_project_drafts()

    detail_response = client.get(f"/projects/{project_id}")
    modules_response = client.get(f"/projects/{project_id}/modules")

    assert detail_response.status_code == 200
    assert detail_response.json()["data"]["id"] == project_id
    assert [module["id"] for module in modules_response.json()["data"]["items"]] == ["field", "review"]


@pytest.mark.parametrize(
    ("template_type", "expected_headers"),
    [
        ("initial_work_orders", ["Terminal", "Station area", "Address"]),
        ("external_completed", ["Terminal", "Station area", "Address", "Communication module", "Installed at"]),
    ],
)
def test_project_templates_download_schema_driven_workbooks(template_type, expected_headers):
    client = TestClient(app)
    create_response = client.post(
        "/projects",
        json={
            "name": "Terminal Field Template Demo",
            "module_ids": ["progress", "field", "review"],
            "work_item_schema": {
                "primary_field": {"key": "terminal_no", "label": "Terminal", "source": "import", "required": True},
                "aggregate_field": {"key": "station_area", "label": "Station area", "source": "import", "required": True},
                "custom_fields": [
                    {"key": "address", "label": "Address", "source": "import", "required": False},
                    {
                        "key": "communication_module",
                        "label": "Communication module",
                        "source": "field_collection",
                        "capture_method": "scan",
                        "required": True,
                    },
                    {
                        "key": "installed_at",
                        "label": "Installed at",
                        "source": "field_collection",
                        "capture_method": "datetime",
                        "data_type": "datetime",
                    },
                ],
            },
        },
    )
    project_id = create_response.json()["data"]["id"]

    response = client.get(f"/projects/{project_id}/templates/{template_type}")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["template", "fields"]
    template_sheet = workbook["template"]
    headers = [cell.value for cell in template_sheet[1] if cell.value]
    for expected_header in expected_headers:
        assert expected_header in headers
    assert template_sheet.max_row >= 3
    fields_sheet = workbook["fields"]
    field_rows = {
        row[0]: row
        for row in fields_sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    }
    assert "terminal_no" in field_rows
    assert "station_area" in field_rows
    if template_type == "external_completed":
        assert "uploaded_at" in field_rows
        assert field_rows["uploaded_at"][6] == "平台上传时补齐"
        assert field_rows["completed_at"][6] == "缺失时按上传时间补齐"
        assert field_rows["installer"][6] == "缺失时按上传人补齐"


def test_unknown_project_template_type_returns_404():
    client = TestClient(app)

    response = client.get("/projects/replacement-project/templates/unknown")

    assert response.status_code == 404


def test_field_collection_project_template_returns_404():
    client = TestClient(app)

    response = client.get("/projects/replacement-project/templates/field_collection")

    assert response.status_code == 404


@pytest.mark.parametrize(
    ("payload", "expected_detail"),
    [
        ({"name": "   ", "module_ids": ["progress"]}, "Project name is required"),
        ({"name": "空模块", "module_ids": []}, "At least one project module is required"),
        ({"name": "未知模块", "module_ids": ["unknown"]}, "Unknown project module"),
    ],
)
def test_create_project_draft_rejects_invalid_payload(payload, expected_detail):
    client = TestClient(app)
    response = client.post("/projects", json=payload)

    assert response.status_code == 400
    assert expected_detail in response.json()["detail"]
