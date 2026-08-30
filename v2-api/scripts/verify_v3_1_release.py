from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import re
import subprocess
import sys
from collections import deque
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlsplit


EXPECTED_VERSION = "3.2.19"
EXPECTED_MIGRATION = "v2-api/alembic/versions/0006_group_barcode_verification.py"


def fail(message: str) -> None:
    raise AssertionError(message)


def read(root: Path, relative_path: str) -> str:
    return (root / relative_path).read_text(encoding="utf-8")


def sample_group() -> dict[str, Any]:
    categories = ("before_box", "collector_barcode", "module_meter", "after_box")
    return {
        "id": "release-sample",
        "status": "approved",
        "archive_status": "archived",
        "terminal": "00123456",
        "meter_no": "000012345678",
        "module_asset_no": "000098765432",
        "collector": "00007777",
        "address": "Release verification address",
        "client_completed_at": "2026-07-23T09:00:00+08:00",
        "photos": [
            {
                "id": f"release-photo-{index}",
                "sha256": f"{index:x}" * 64,
                "category": category,
                "is_active": True,
                "upload_status": "uploaded",
            }
            for index, category in enumerate(categories, start=1)
        ],
    }


class _MigrationRecorder:
    def __init__(self) -> None:
        self.tables: dict[str, tuple[Any, ...]] = {}
        self.indexes: set[str] = set()

    def create_table(self, name: str, *elements: Any, **_kwargs: Any) -> None:
        self.tables[name] = elements

    def create_index(self, name: str, *_args: Any, **_kwargs: Any) -> None:
        self.indexes.add(name)


def verify_paused_defaults(root: Path) -> dict[str, bool]:
    api_root = root / "v2-api"
    if str(api_root) not in sys.path:
        sys.path.insert(0, str(api_root))

    from app.models import BarcodeMaintenanceControl

    paused_column = BarcodeMaintenanceControl.__table__.c.paused
    model_python_default = paused_column.default is not None and paused_column.default.arg is True
    model_server_default = (
        paused_column.server_default is not None
        and str(paused_column.server_default.arg).strip().lower() == "true"
    )
    model_paused = model_python_default and model_server_default

    migration_path = root / EXPECTED_MIGRATION
    spec = importlib.util.spec_from_file_location("v3_1_release_migration", migration_path)
    if spec is None or spec.loader is None:
        fail(f"unable to import migration: {migration_path}")
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    recorder = _MigrationRecorder()
    migration.op = recorder
    migration.upgrade()
    if set(recorder.tables) != {"group_barcode_verifications", "barcode_maintenance_controls"}:
        fail("migration must create both durable verification tables")
    if recorder.indexes != {
        "ix_group_barcode_verifications_pending",
        "ix_group_barcode_verifications_lease",
    }:
        fail("migration must create pending and lease indexes")
    migration_paused_column = next(
        (
            element
            for element in recorder.tables["barcode_maintenance_controls"]
            if getattr(element, "name", None) == "paused"
        ),
        None,
    )
    migration_paused = (
        migration_paused_column is not None
        and migration_paused_column.server_default is not None
        and str(migration_paused_column.server_default.arg).strip().lower() == "true"
    )
    if not migration_paused or not model_paused:
        fail("migration and model must both default the worker to paused")
    return {"migration_paused": migration_paused, "model_paused": model_paused}


def verify_worker_behavior() -> dict[str, Any]:
    from app.services.barcode_maintenance_worker import MaintenanceJob, run_worker_batch

    pending = deque(
        MaintenanceJob(kind="verification", team_id="release-team", group_id=f"group-{index}")
        for index in range(21)
    )
    active = 0
    maximum_active = 0
    processed: list[str] = []
    sleeps: list[float] = []

    def claim_next() -> MaintenanceJob | None:
        return pending.popleft() if pending else None

    def process_job(job: MaintenanceJob) -> None:
        nonlocal active, maximum_active
        active += 1
        maximum_active = max(maximum_active, active)
        processed.append(job.group_id)
        active -= 1

    report = run_worker_batch(
        batch_size=20,
        batch_pause_seconds=5,
        claim_next=claim_next,
        process_job=process_job,
        can_claim=lambda: True,
        load_too_high=lambda: False,
        sleeper=sleeps.append,
    )
    if report != {"processed": 20, "failed": 0, "status": "complete"}:
        fail("worker must process exactly one successful 20-group batch")
    if processed != [f"group-{index}" for index in range(20)]:
        fail("worker must process groups in strict serial claim order")
    result = {
        "processed": len(processed),
        "remaining": len(pending),
        "maximum_active": maximum_active,
        "sleeps": sleeps,
    }
    if result != {"processed": 20, "remaining": 1, "maximum_active": 1, "sleeps": [5.0]}:
        fail("worker must leave task 21 queued and sleep once for five seconds")
    return result


def verify_eligibility_and_ocr_behavior() -> dict[str, str]:
    from app.services.group_barcode_verification import evaluate_group_eligibility, scan_group_evidence

    eligible = sample_group()
    eligible_result = evaluate_group_eligibility(eligible)
    ineligible = copy.deepcopy(eligible)
    ineligible["photos"][1]["category"] = "before_box"
    ineligible_result = evaluate_group_eligibility(ineligible)

    ocr_only = copy.deepcopy(eligible)
    expected_values = (ocr_only["meter_no"], ocr_only["module_asset_no"], ocr_only["collector"])
    for photo, value in zip(ocr_only["photos"], expected_values, strict=False):
        photo["ocr_candidate_values"] = [value]
    ocr_result = scan_group_evidence(ocr_only, ocr_only["photos"])

    if eligible_result.status != "pending" or ineligible_result.status != "not_eligible":
        fail("eligibility must accept exactly four unique categories and reject duplicates")
    if ocr_result.status == "passed" or ocr_result.passed_count != 0 or len(ocr_result.matched_ocr_candidates) != 3:
        fail("OCR-only evidence must remain a candidate and cannot pass")
    return {
        "eligible_status": eligible_result.status,
        "ineligible_status": ineligible_result.status,
        "ocr_only_status": ocr_result.status,
    }


def verify_export_behavior() -> dict[str, Any]:
    from openpyxl import load_workbook

    from app.services.final_delivery_export import (
        build_delivery_workbook,
        collect_delivery_validation_errors,
    )

    group = sample_group()
    workbook = load_workbook(BytesIO(build_delivery_workbook([group])))
    if workbook.sheetnames != ["老设备", "新设备"]:
        fail("export sample must generate the two required sheets")
    old_sheet = workbook["老设备"]
    new_sheet = workbook["新设备"]
    old_row = [cell.value for cell in old_sheet[2]]
    new_row = [cell.value for cell in new_sheet[2]]
    if old_row != [
        "00123456",
        None,
        "Release verification address",
        "000012345678",
        "00007777",
        "南大供电服务中心",
        "奕福",
        None,
        None,
    ]:
        fail("old-device sheet row content changed")
    if new_row != ["2026-07-23", "000012345678", "000098765432", None]:
        fail("new-device sheet row content changed")
    all_text_formatted = all(
        cell.number_format == "@"
        for sheet in (old_sheet, new_sheet)
        for row in sheet.iter_rows(min_row=2)
        for cell in row
    )
    if not all_text_formatted:
        fail("delivery workbook data cells must use text format")

    for field in ("terminal", "meter_no", "module_asset_no", "collector"):
        for placeholder in ("0", "00000000", "000000000000", "test-release"):
            invalid = copy.deepcopy(group)
            invalid[field] = placeholder
            errors = collect_delivery_validation_errors([invalid], require_cache=False)
            if not any(error["code"] == "placeholder_identity" and error["field"] == field for error in errors):
                fail(f"formal export must reject placeholder {field}: {placeholder}")
    return {
        "sheetnames": workbook.sheetnames,
        "old_fixed_values": old_row[5:7],
        "all_text_formatted": all_text_formatted,
    }


def _required_mapping(value: Any, name: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        fail(f"performance report {name} is missing")
    return value


def verify_performance_report(report: dict[str, Any], *, expected_source_commit: str) -> dict[str, Any]:
    from scripts.verify_task_review_performance import observed_build_count, verify_measurements

    expected_source_commit = str(expected_source_commit or "").strip().lower()
    report_commit = str(report.get("source_commit") or "").strip().lower()
    if not re.fullmatch(r"[0-9a-f]{40}", expected_source_commit) or report_commit != expected_source_commit:
        fail("performance report source commit does not match the expected source commit")

    base_url = str(report.get("base_url") or "").strip()
    parsed_url = urlsplit(base_url)
    if parsed_url.scheme not in {"http", "https"} or parsed_url.hostname not in {"localhost", "127.0.0.1"}:
        fail("performance report base_url must use localhost or 127.0.0.1")
    if parsed_url.username or parsed_url.password:
        fail("performance report base_url must not contain credentials")

    task_id = str(report.get("task_id") or "").strip()
    if not task_id:
        fail("performance report task_id must be non-empty")
    if int(report.get("build_sample_seconds") or 0) < 60:
        fail("performance report must sample cache builds for at least 60 seconds")

    start_instance = str(report.get("start_cache_instance_id") or "").strip()
    end_instance = str(report.get("end_cache_instance_id") or "").strip()
    try:
        observed_builds = observed_build_count(
            int(report.get("start_build_count") or 0),
            int(report.get("end_build_count") or 0),
            start_instance,
            end_instance,
        )
    except (TypeError, ValueError) as error:
        fail(f"performance report cache instance evidence is invalid: {error}")

    measurements = _required_mapping(report.get("measurements"), "measurements")
    required_measurements = {
        "task_snapshot_ms",
        "review_groups_ms",
        "review_group_count",
        "task_builds_in_60_seconds",
    }
    if not required_measurements <= measurements.keys():
        fail("performance report measurements are incomplete")
    if int(measurements["task_builds_in_60_seconds"]) != observed_builds:
        fail("performance report task_builds_in_60_seconds does not match cache counters")
    recomputed = verify_measurements(
        task_snapshot_ms=float(measurements["task_snapshot_ms"]),
        review_groups_ms=float(measurements["review_groups_ms"]),
        review_group_count=int(measurements["review_group_count"]),
        task_builds_in_60_seconds=observed_builds,
    )
    if not recomputed["ok"]:
        fail(f"performance report exceeds thresholds: {', '.join(recomputed['failures'])}")
    if report.get("ok") is not True or list(report.get("failures") or []):
        fail("performance report did not record a successful real API verification")

    routes = _required_mapping(report.get("routes"), "routes")
    task_route = _required_mapping(routes.get("task_snapshot"), "task snapshot route")
    review_route = _required_mapping(routes.get("review_groups"), "review groups route")
    if task_route.get("path") != "/local-test/tasks/snapshot":
        fail("performance report task snapshot route is invalid")
    if int(task_route.get("item_count") or 0) < 1:
        fail("performance report must contain non-empty task snapshot data")
    if str(task_route.get("selected_task_id") or "").strip() != task_id:
        fail("performance report selected task_id does not match route evidence")
    if int(task_route.get("serialized_bytes") or 0) < 1:
        fail("performance report task snapshot serialization evidence is empty")

    expected_review_path = (
        f"/local-test/tasks/{quote(task_id, safe='')}/review-groups"
        "?limit=20&offset=0&review_status=all&query="
    )
    if review_route.get("path") != expected_review_path:
        fail("performance report review groups route is invalid")
    review_count = int(review_route.get("item_count") or 0)
    if review_count < 1:
        fail("performance report must contain non-empty review group data")
    if review_count != int(measurements["review_group_count"]):
        fail("performance report review group count does not match measurements")
    if int(review_route.get("total") or 0) < review_count:
        fail("performance report review total is inconsistent")
    if int(review_route.get("limit") or 0) != 20 or int(review_route.get("offset", -1)) != 0:
        fail("performance report review pagination evidence is invalid")
    if int(review_route.get("serialized_bytes") or 0) < 1:
        fail("performance report review serialization evidence is empty")
    return recomputed


def verify_repository_source_commit(root: Path, expected_source_commit: str) -> str:
    try:
        current_commit = subprocess.check_output(
            ["git", "rev-parse", "HEAD"],
            cwd=root,
            text=True,
            encoding="utf-8",
            stderr=subprocess.STDOUT,
        ).strip().lower()
    except (OSError, subprocess.CalledProcessError) as error:
        fail(f"unable to resolve current HEAD for performance report binding: {error}")
    if current_commit != str(expected_source_commit or "").strip().lower():
        fail("expected performance source commit does not match current HEAD")
    return current_commit


def verify(root: Path, *, performance_report: dict[str, Any], expected_source_commit: str) -> list[str]:
    api_root = root / "v2-api"
    if str(api_root) not in sys.path:
        sys.path.insert(0, str(api_root))

    version_markers = {
        "v2-api/app/main.py": f'version="{EXPECTED_VERSION}"',
        "v2-api/app/services/ops_status.py": f'return "{EXPECTED_VERSION}"',
        "v2-api/pyproject.toml": f'version = "{EXPECTED_VERSION}"',
        "v2-web/package.json": f'"version": "{EXPECTED_VERSION}"',
        "v2-web/index.html": f"Module Manager V{EXPECTED_VERSION}",
        "v2-web/src/version.json": f'"version":"{EXPECTED_VERSION}"',
        "v2-web/src/components/AppLayout.vue": f"V{EXPECTED_VERSION}",
        "scripts/build-client-release.ps1": f'[string]$Version = "{EXPECTED_VERSION}"',
    }
    for path, marker in version_markers.items():
        if marker not in read(root, path):
            fail(f"runtime version source is not {EXPECTED_VERSION}: {path}")

    verify_repository_source_commit(root, expected_source_commit)
    verify_paused_defaults(root)
    verify_worker_behavior()
    verify_eligibility_and_ocr_behavior()
    verify_export_behavior()
    verify_performance_report(performance_report, expected_source_commit=expected_source_commit)

    return [
        "migration",
        "4-photo eligibility",
        "OCR-only blocked",
        "worker paused",
        "20 groups serial",
        "5 seconds",
        "placeholder zero",
        "export sample",
        "task claim",
        "review performance",
    ]


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Verify the offline V3.1 release candidate contract.")
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--performance-report", type=Path, required=True)
    parser.add_argument("--expected-source-commit", required=True)
    parser.add_argument("--json", action="store_true", help="Print the verified gate names as JSON.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        performance_report = json.loads(args.performance_report.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        fail(f"unable to read performance report: {error}")
    if not isinstance(performance_report, dict):
        fail("performance report root must be an object")
    gates = verify(
        args.repo_root.resolve(),
        performance_report=performance_report,
        expected_source_commit=args.expected_source_commit,
    )
    if args.json:
        print(json.dumps({"version": EXPECTED_VERSION, "gates": gates}, ensure_ascii=False))
    else:
        for gate in gates:
            print(f"[OK] {gate}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AssertionError as exc:
        print(f"[FAIL] {exc}", file=sys.stderr)
        raise SystemExit(1)
