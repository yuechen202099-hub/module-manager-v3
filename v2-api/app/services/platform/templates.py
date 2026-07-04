from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, time
from io import BytesIO
import json
from pathlib import Path
import threading
from typing import Any, Literal
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.services.platform.catalog import get_project_overview

TemplateType = Literal["initial_work_orders", "external_completed"]

SUPPORTED_TEMPLATE_TYPES = {"initial_work_orders", "external_completed"}

_IMPORT_BATCHES_LOCK = threading.RLock()
_IMPORT_BATCHES_LOADED = False
_IMPORT_BATCHES: dict[str, dict[str, Any]] = {}
_WORK_ORDER_TASKS_LOCK = threading.RLock()
_WORK_ORDER_TASKS_LOADED = False
_WORK_ORDER_TASKS: dict[str, dict[str, Any]] = {}
_PLATFORM_WORK_ORDERS_LOCK = threading.RLock()
_PLATFORM_WORK_ORDERS_LOADED = False
_PLATFORM_WORK_ORDERS: dict[str, dict[str, Any]] = {}

PLATFORM_FILL_RULES = {
    "uploaded_at": "平台上传时补齐",
    "completed_at": "缺失时按上传时间补齐",
    "installer": "缺失时按上传人补齐",
    "work_order_id": "缺失时由平台生成",
    "photo_count": "缺失时按0或附件数补齐",
    "online_duration_minutes": "外部无法提供时可留空",
}

_DELIVERY_ARCHIVE_PLATFORM_KPI_KEYS = (
    "installer",
    "started_at",
    "completed_at",
    "uploaded_at",
    "online_duration_minutes",
    "photo_count",
    "old_device_recovered",
)

_DELIVERY_ARCHIVE_COMPUTED_KPI_FIELDS: dict[str, dict[str, Any]] = {
    "old_device_recovered": {
        "key": "old_device_recovered",
        "label": "旧设备回收",
        "source": "system",
        "capture_method": "system",
        "data_type": "boolean",
        "required": True,
        "relation_role": "old_device",
    },
}


def build_project_template_workbook(project_id: str, template_type: str) -> bytes:
    if template_type not in SUPPORTED_TEMPLATE_TYPES:
        raise KeyError(template_type)
    project = get_project_overview(project_id)
    schema = project.get("work_item_schema") or {}
    fields = _fields_for_template(schema, template_type)
    field_rows = _field_rows_for_template(schema, fields, template_type)

    workbook = Workbook()
    template_sheet = workbook.active
    template_sheet.title = "template"
    fields_sheet = workbook.create_sheet("fields")
    instructions_sheet = workbook.create_sheet("instructions")

    headers = [field["label"] for field in fields]
    keys = [field["key"] for field in fields]
    template_sheet.append(headers)
    template_sheet.append([_example_value(field, index) for index, field in enumerate(fields, start=1)])
    template_sheet.append([f"字段编码: {key}" for key in keys])

    fields_sheet.append(["字段编码", "字段名称", "来源", "采集方式", "数据格式", "是否必填", "层级角色", "父字段", "条件采集", "平台补齐规则"])
    for field in field_rows:
        fields_sheet.append(
            [
                field["key"],
                field["label"],
                field.get("source", ""),
                field.get("capture_method", ""),
                field.get("data_type", ""),
                "是" if field.get("required") else "否",
                _template_hierarchy_role(field),
                _template_parent_label(schema, field),
                _template_condition_hint(schema, field),
                _platform_fill_rule(field, template_type),
            ]
        )

    _append_template_instruction_sheet(instructions_sheet, schema, fields, field_rows, template_type)

    for sheet in workbook.worksheets:
        _style_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_project_template_preview(project_id: str, work_item_schema: dict[str, Any] | None = None) -> dict[str, Any]:
    project = get_project_overview(project_id)
    saved_schema = project.get("work_item_schema") or {}
    schema = deepcopy(work_item_schema) if isinstance(work_item_schema, dict) else deepcopy(saved_schema)
    if not schema.get("platform_required_fields"):
        schema["platform_required_fields"] = deepcopy(saved_schema.get("platform_required_fields") or [])

    templates: list[dict[str, Any]] = []
    for template_type in ("initial_work_orders", "external_completed"):
        fields = _fields_for_template(schema, template_type)
        field_rows = _field_rows_for_template(schema, fields, template_type)
        templates.append(
            {
                "template_type": template_type,
                "headers": [field["label"] for field in fields],
                "field_rows": [
                    {
                        "key": field["key"],
                        "label": field["label"],
                        "source": field.get("source", ""),
                        "capture_method": field.get("capture_method", ""),
                        "data_type": field.get("data_type", ""),
                        "required": bool(field.get("required")),
                        "parent_key": field.get("parent_key", ""),
                        "relation_role": field.get("relation_role", ""),
                        "required_when": field.get("required_when") or None,
                        "show_in_construction_panel": bool(field.get("show_in_construction_panel", False)),
                        "template_hierarchy_role": _template_hierarchy_role(field),
                        "template_parent_label": _template_parent_label(schema, field),
                        "template_condition_hint": _template_condition_hint(schema, field),
                        "platform_fill_rule": _platform_fill_rule(field, template_type),
                    }
                    for field in field_rows
                ],
            }
        )

    site_required_fields = [
        field["label"]
        for field in _dedupe_fields(
            [
                _clean_field(field)
                for field in schema.get("custom_fields", [])
                if isinstance(field, dict)
                and field.get("required")
                and str(field.get("source") or "").strip() == "field_collection"
            ]
        )
    ]
    return {
        "project_id": project_id,
        "templates": templates,
        "site_required_fields": site_required_fields,
    }


def validate_project_template_workbook(project_id: str, template_type: str, workbook_content: bytes) -> dict[str, Any]:
    if template_type not in SUPPORTED_TEMPLATE_TYPES:
        raise KeyError(template_type)
    project = get_project_overview(project_id)
    schema = project.get("work_item_schema") or {}
    fields = _fields_for_template(schema, template_type)
    workbook = load_workbook(BytesIO(workbook_content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return _validation_report(project_id, template_type, fields, [], [], [_report_item("error", "missing_header", None, "", "", "模板没有表头")])

    headers = [_clean_cell(value) for value in rows[0]]
    header_indexes = {header: index for index, header in enumerate(headers) if header}
    data_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows[1:], start=2)
        if not _is_empty_row(row) and not _is_field_code_row(row)
    ]
    items: list[dict[str, Any]] = []
    expected_labels = [field["label"] for field in fields]
    primary = _clean_field(schema.get("primary_field"))
    aggregate = _clean_field(schema.get("aggregate_field"))

    for field in _required_headers(fields, primary, aggregate, template_type):
        if field["label"] not in header_indexes:
            items.append(
                _report_item(
                    "error",
                    "missing_required_header",
                    None,
                    field["key"],
                    field["label"],
                    f"缺少必填列：{field['label']}",
                )
            )

    for field in _recommended_headers(fields, primary, aggregate, template_type):
        if field["label"] not in header_indexes:
            items.append(
                _report_item(
                    "warning",
                    "missing_recommended_header",
                    None,
                    field["key"],
                    field["label"],
                    f"建议补充列：{field['label']}",
                )
            )

    seen_primary_values: dict[str, int] = {}
    fields_by_label = {field["label"]: field for field in fields}
    for row_number, row in data_rows:
        row_values = {header: _value_at(row, index) for header, index in header_indexes.items()}
        primary_label = primary.get("label", "")
        primary_value = _clean_cell(row_values.get(primary_label))
        if primary_label and primary_label in header_indexes:
            if not primary_value:
                items.append(
                    _report_item("error", "missing_primary_field", row_number, primary.get("key", ""), primary_label, f"主字段不能为空：{primary_label}")
                )
            elif primary_value in seen_primary_values:
                items.append(
                    _report_item(
                        "error",
                        "duplicate_primary_field",
                        row_number,
                        primary.get("key", ""),
                        primary_label,
                        f"主字段重复：{primary_value}",
                        primary_value,
                    )
                )
            else:
                seen_primary_values[primary_value] = row_number

        for field in _recommended_headers(fields, primary, aggregate, template_type):
            if field["label"] in header_indexes and not _clean_cell(row_values.get(field["label"])):
                items.append(
                    _report_item(
                        "warning",
                        "missing_recommended_field",
                        row_number,
                        field["key"],
                        field["label"],
                        f"建议补充：{field['label']}",
                    )
                )

        _validate_conditional_template_fields(
            items=items,
            fields=fields,
            row_number=row_number,
            row_values=row_values,
            header_indexes=header_indexes,
            template_type=template_type,
        )

        for label, value in row_values.items():
            field = fields_by_label.get(label)
            if not field or value in (None, ""):
                continue
            data_type = str(field.get("data_type") or "")
            if data_type in {"number", "duration"} and not _is_valid_number(value):
                items.append(_report_item("error", "invalid_number", row_number, field["key"], label, f"数字格式错误：{label}", _clean_cell(value)))
            if data_type == "datetime" and not _is_valid_datetime(value):
                items.append(_report_item("error", "invalid_datetime", row_number, field["key"], label, f"时间格式错误：{label}", _clean_cell(value)))

    return _validation_report(project_id, template_type, fields, expected_labels, data_rows, items)


def create_project_template_import_draft(
    project_id: str,
    template_type: str,
    workbook_content: bytes,
    filename: str = "",
) -> dict[str, Any]:
    if template_type not in SUPPORTED_TEMPLATE_TYPES:
        raise KeyError(template_type)
    project = get_project_overview(project_id)
    schema = project.get("work_item_schema") or {}
    fields = _fields_for_template(schema, template_type)
    validation = validate_project_template_workbook(project_id, template_type, workbook_content)
    headers, data_rows = _template_data_rows(workbook_content)
    header_indexes = {header: index for index, header in enumerate(headers) if header}
    fields_by_label = {field["label"]: field for field in fields}
    primary = _clean_field(schema.get("primary_field"))
    aggregate = _clean_field(schema.get("aggregate_field"))
    warnings_by_row: dict[int, list[dict[str, Any]]] = {}
    for item in validation.get("items", []):
        if item.get("severity") != "warning" or item.get("row") is None:
            continue
        warnings_by_row.setdefault(int(item["row"]), []).append(item)

    preview_rows: list[dict[str, Any]] = []
    for row_number, row in data_rows[:10]:
        values: dict[str, str] = {}
        for label, index in header_indexes.items():
            field = fields_by_label.get(label)
            if not field:
                continue
            values[field["key"]] = _clean_cell(_value_at(row, index))
        preview_rows.append(
            {
                "row": row_number,
                "primary_value": values.get(primary.get("key", ""), ""),
                "aggregate_value": values.get(aggregate.get("key", ""), ""),
                "values": values,
                "warnings": warnings_by_row.get(row_number, []),
            }
        )

    error_count = int(validation.get("summary", {}).get("error_count", 0))
    warning_count = int(validation.get("summary", {}).get("warning_count", 0))
    hierarchy_gap_items = [
        {
            "row": item.get("row"),
            "field_key": item.get("field_key", ""),
            "field_label": item.get("field_label", ""),
            "message": item.get("message", ""),
            "value": item.get("value", ""),
        }
        for item in validation.get("items", [])
        if item.get("code") == "missing_conditional_field"
    ]
    status = "blocked" if error_count else "ready_with_warnings" if warning_count else "ready"
    return {
        "job_id": f"draft-{uuid4().hex}",
        "project_id": project_id,
        "template_type": template_type,
        "filename": filename,
        "status": status,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "progress": {
            "phase": "blocked" if error_count else "draft_ready",
            "processed_rows": len(data_rows),
            "total_rows": len(data_rows),
        },
        "result": {
            "mode": "dry_run",
            "validation": validation,
            "summary": {
                "total_rows": len(data_rows),
                "ready_rows": 0 if error_count else len(data_rows),
                "preview_rows": len(preview_rows),
                "hierarchy_gap_count": len(hierarchy_gap_items),
            },
            "field_mappings": {
                label: fields_by_label[label]["key"]
                for label in headers
                if label in fields_by_label
            },
            "platform_fill_policy": _platform_fill_policies(schema, fields, template_type),
            "preview_rows": preview_rows,
            "hierarchy_gap_items": hierarchy_gap_items[:20],
        },
        "error": "",
    }


def create_project_template_import_batch(
    project_id: str,
    template_type: str,
    workbook_content: bytes,
    filename: str = "",
    actor: str = "",
) -> dict[str, Any]:
    draft = create_project_template_import_draft(project_id, template_type, workbook_content, filename)
    validation_summary = draft.get("result", {}).get("validation", {}).get("summary", {})
    if int(validation_summary.get("error_count") or 0):
        raise ValueError("Template import batch has validation errors")

    result = deepcopy(draft.get("result", {}))
    summary = result.setdefault("summary", {})
    total_rows = int(summary.get("total_rows") or 0)
    summary["created_work_orders"] = 0
    result["mode"] = "confirmed_dry_run_batch"
    result["rollback"] = {
        "strategy": "delete_import_batch_record",
        "created_work_orders": 0,
        "created_groups": 0,
        "note": "Current step records only the import batch. No business work orders are created yet.",
    }

    batch_id = f"batch-{uuid4().hex}"
    record = {
        "job_id": batch_id,
        "project_id": project_id,
        "template_type": template_type,
        "filename": filename,
        "actor": actor,
        "status": "confirmed",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "progress": {
            "phase": "batch_recorded",
            "processed_rows": total_rows,
            "total_rows": total_rows,
        },
        "result": result,
        "error": "",
    }
    with _IMPORT_BATCHES_LOCK:
        _load_import_batches_unlocked()
        _IMPORT_BATCHES[batch_id] = record
        _save_import_batches_unlocked()
    return record


def get_project_template_import_batch(project_id: str, batch_id: str) -> dict[str, Any]:
    with _IMPORT_BATCHES_LOCK:
        _load_import_batches_unlocked()
        record = _IMPORT_BATCHES.get(batch_id)
        if not record or record.get("project_id") != project_id:
            raise KeyError(batch_id)
        return deepcopy(record)


def create_import_work_order_task(project_id: str, batch_id: str, actor: str = "", mode: str = "safe_record_only") -> dict[str, Any]:
    if mode != "safe_record_only":
        raise ValueError("Only safe_record_only mode is supported")
    batch = get_project_template_import_batch(project_id, batch_id)
    batch_result = deepcopy(batch.get("result") or {})
    batch_summary = batch_result.get("summary") if isinstance(batch_result.get("summary"), dict) else {}
    total_rows = int(batch_summary.get("total_rows") or 0)
    ready_rows = int(batch_summary.get("ready_rows") or total_rows)
    task_id = f"task-{uuid4().hex}"
    result = {
        "mode": mode,
        "source_batch_id": batch_id,
        "summary": {
            "total_rows": total_rows,
            "planned_work_orders": ready_rows,
            "created_work_orders": 0,
            "skipped_rows": max(total_rows - ready_rows, 0),
        },
        "field_mappings": batch_result.get("field_mappings") or {},
        "platform_fill_policy": batch_result.get("platform_fill_policy") or [],
        "preview_rows": batch_result.get("preview_rows") or [],
        "rollback": {
            "strategy": "delete_import_work_order_task_record",
            "created_work_orders": 0,
            "created_groups": 0,
            "note": "Current step plans work order creation only. No business work orders are created yet.",
        },
    }
    record = {
        "job_id": task_id,
        "project_id": project_id,
        "batch_id": batch_id,
        "template_type": batch.get("template_type", ""),
        "filename": batch.get("filename", ""),
        "actor": actor,
        "status": "planned",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "progress": {
            "phase": "work_order_task_planned",
            "processed_rows": 0,
            "total_rows": total_rows,
        },
        "result": result,
        "error": "",
    }
    with _WORK_ORDER_TASKS_LOCK:
        _load_work_order_tasks_unlocked()
        _WORK_ORDER_TASKS[task_id] = record
        _save_work_order_tasks_unlocked()
    return record


def get_import_work_order_task(project_id: str, task_id: str) -> dict[str, Any]:
    with _WORK_ORDER_TASKS_LOCK:
        _load_work_order_tasks_unlocked()
        record = _WORK_ORDER_TASKS.get(task_id)
        if not record or record.get("project_id") != project_id:
            raise KeyError(task_id)
        return deepcopy(record)


def execute_import_work_order_task(project_id: str, task_id: str, actor: str = "", mode: str = "local_platform_store") -> dict[str, Any]:
    if mode != "local_platform_store":
        raise ValueError("Only local_platform_store mode is supported")
    construction_field_keys: set[str] = set()
    photo_slot_keys: set[str] = set()
    with _WORK_ORDER_TASKS_LOCK, _PLATFORM_WORK_ORDERS_LOCK:
        _load_work_order_tasks_unlocked()
        _load_platform_work_orders_unlocked()
        task = _WORK_ORDER_TASKS.get(task_id)
        if not task or task.get("project_id") != project_id:
            raise KeyError(task_id)
        template_type = str(task.get("template_type") or "")
        if template_type == "external_completed":
            construction_schema = _platform_construction_schema(project_id)
            construction_field_keys = {field["key"] for field in construction_schema["construction_fields"]}
            photo_slot_keys = {field["key"] for field in construction_schema["photo_slots"]}
        if task.get("status") == "rolled_back":
            raise ValueError("Rolled back work order task cannot be executed")
        if task.get("status") == "completed":
            return deepcopy(task)

        result = task.setdefault("result", {})
        preview_rows = result.get("preview_rows") if isinstance(result.get("preview_rows"), list) else []
        created_ids: list[str] = []
        created_at = datetime.now().isoformat(timespec="seconds")
        for row in preview_rows:
            if not isinstance(row, dict):
                continue
            work_order_id = f"wo-{uuid4().hex}"
            values = row.get("values") if isinstance(row.get("values"), dict) else {}
            values = {
                str(key): "" if value is None else str(value).strip()
                for key, value in values.items()
                if str(key)
            }
            covered_photo_slots: list[str] = []
            collection_field_values: dict[str, str] = {}
            if template_type == "external_completed":
                covered_photo_slots = [
                    key
                    for key in sorted(photo_slot_keys)
                    if str(values.get(key) or "").strip()
                ]
                values = _external_completed_platform_values(
                    values,
                    work_order_id=work_order_id,
                    actor=actor,
                    imported_at=created_at,
                    photo_count=len(covered_photo_slots),
                )
                collection_field_values = {
                    key: values[key]
                    for key in sorted(construction_field_keys)
                    if str(values.get(key) or "").strip()
                }
            created_ids.append(work_order_id)
            work_order = {
                "id": work_order_id,
                "project_id": project_id,
                "source_task_id": task_id,
                "source_batch_id": task.get("batch_id", ""),
                "source_template_type": template_type,
                "primary_value": row.get("primary_value", ""),
                "aggregate_value": row.get("aggregate_value", ""),
                "values": values,
                "status": "created",
                "created_by": actor,
                "created_at": created_at,
            }
            if template_type == "external_completed":
                work_order.update(
                    {
                        "collection_status": "submitted",
                        "collection_field_values": collection_field_values,
                        "covered_photo_slots": covered_photo_slots,
                        "client_batch_id": task.get("batch_id", ""),
                        "collected_by": actor,
                        "collected_at": created_at,
                        "review_hierarchy_gap_items": _platform_hierarchy_gap_items_from_warnings(row),
                    }
                )
            _PLATFORM_WORK_ORDERS[work_order_id] = work_order

        summary = result.setdefault("summary", {})
        summary["created_work_orders"] = len(created_ids)
        summary["planned_work_orders"] = int(summary.get("planned_work_orders") or len(created_ids))
        result["created_work_order_ids"] = created_ids
        result["rollback"] = {
            "strategy": "delete_platform_work_orders_by_task",
            "created_work_orders": len(created_ids),
            "created_groups": 0,
            "note": "Rollback deletes only local platform work orders created by this task.",
        }
        task["status"] = "completed"
        task["executed_at"] = created_at
        task["executed_by"] = actor
        task["progress"] = {
            "phase": "work_orders_created",
            "processed_rows": len(created_ids),
            "total_rows": int(summary.get("total_rows") or len(created_ids)),
        }
        _WORK_ORDER_TASKS[task_id] = task
        _save_platform_work_orders_unlocked()
        _save_work_order_tasks_unlocked()
        return deepcopy(task)


def _external_completed_platform_values(
    values: dict[str, str],
    *,
    work_order_id: str,
    actor: str,
    imported_at: str,
    photo_count: int,
) -> dict[str, str]:
    filled = dict(values)
    filled["work_order_id"] = filled.get("work_order_id") or work_order_id
    filled["uploaded_at"] = filled.get("uploaded_at") or imported_at
    filled["completed_at"] = filled.get("completed_at") or imported_at
    filled["installer"] = filled.get("installer") or actor
    filled["photo_count"] = filled.get("photo_count") or str(photo_count)
    filled.setdefault("online_duration_minutes", "")
    return filled


def rollback_import_work_order_task(project_id: str, task_id: str, actor: str = "") -> dict[str, Any]:
    with _WORK_ORDER_TASKS_LOCK, _PLATFORM_WORK_ORDERS_LOCK:
        _load_work_order_tasks_unlocked()
        _load_platform_work_orders_unlocked()
        task = _WORK_ORDER_TASKS.get(task_id)
        if not task or task.get("project_id") != project_id:
            raise KeyError(task_id)
        deleted_ids = [
            work_order_id
            for work_order_id, work_order in _PLATFORM_WORK_ORDERS.items()
            if work_order.get("project_id") == project_id and work_order.get("source_task_id") == task_id
        ]
        for work_order_id in deleted_ids:
            _PLATFORM_WORK_ORDERS.pop(work_order_id, None)

        result = task.setdefault("result", {})
        summary = result.setdefault("summary", {})
        summary["created_work_orders"] = 0
        result["created_work_order_ids"] = []
        result["rollback"] = {
            "strategy": "delete_platform_work_orders_by_task",
            "created_work_orders": 0,
            "deleted_work_orders": len(deleted_ids),
            "created_groups": 0,
            "note": "Local platform work orders created by this task have been deleted.",
        }
        task["status"] = "rolled_back"
        task["rolled_back_at"] = datetime.now().isoformat(timespec="seconds")
        task["rolled_back_by"] = actor
        task["progress"] = {
            "phase": "work_orders_rolled_back",
            "processed_rows": len(deleted_ids),
            "total_rows": int(summary.get("total_rows") or len(deleted_ids)),
        }
        _WORK_ORDER_TASKS[task_id] = task
        _save_platform_work_orders_unlocked()
        _save_work_order_tasks_unlocked()
        return deepcopy(task)


def summarize_platform_work_orders(project_id: str) -> dict[str, int]:
    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_orders = [
            deepcopy(work_order)
            for work_order in _PLATFORM_WORK_ORDERS.values()
            if work_order.get("project_id") == project_id
        ]
    total = len(work_orders)
    uploaded = 0
    reviewing = 0
    archived = 0
    returned = 0
    exceptions = 0
    not_ready = 0
    initial_work_orders = 0
    external_completed = 0
    kpi_ready = 0
    photo_total = 0
    old_device_recovered = 0
    online_durations: list[int] = []
    installers: set[str] = set()
    for work_order in work_orders:
        template_type = str(work_order.get("source_template_type") or "").strip()
        if template_type == "external_completed":
            external_completed += 1
        else:
            initial_work_orders += 1
        kpi_values = _platform_kpi_values_for_work_order(work_order)
        installer = str(kpi_values.get("installer") or "").strip()
        completed_at = str(kpi_values.get("completed_at") or "").strip()
        if installer and completed_at:
            kpi_ready += 1
        if installer:
            installers.add(installer)
        photo_total += _safe_int(kpi_values.get("photo_count"))
        old_device_recovered += 1 if _safe_int(kpi_values.get("old_device_recovered")) > 0 else 0
        online_duration = _safe_int(kpi_values.get("online_duration_minutes"))
        if online_duration > 0:
            online_durations.append(online_duration)
        has_collection = _has_platform_collection(work_order)
        if has_collection:
            uploaded += 1
        status = _platform_review_status(work_order)
        if status == "approved":
            archived += 1
            continue
        if status == "pending_review":
            reviewing += 1
            continue
        if status == "returned":
            returned += 1
            continue
        if status == "exception":
            exceptions += 1
            continue
        if not has_collection:
            not_ready += 1
    return {
        "total": total,
        "uploaded": uploaded,
        "reviewing": reviewing,
        "archived": archived,
        "completed": archived,
        "approved": archived,
        "returned": returned,
        "initial_work_orders": initial_work_orders,
        "external_completed": external_completed,
        "pending_review": reviewing,
        "returned_rework": returned,
        "approved_archive": archived,
        "exceptions": returned + exceptions,
        "exception": exceptions,
        "not_ready": not_ready,
        "kpi_ready": kpi_ready,
        "photo_total": photo_total,
        "old_device_recovered": old_device_recovered,
        "average_online_duration_minutes": int(round(sum(online_durations) / len(online_durations))) if online_durations else 0,
        "installer_count": len(installers),
    }


def build_platform_delivery_archive_readiness(project_id: str) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    construction_fields = construction_schema["construction_fields"]
    photo_slots = construction_schema["photo_slots"]

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_orders = [
            deepcopy(work_order)
            for work_order in _PLATFORM_WORK_ORDERS.values()
            if work_order.get("project_id") == project_id
        ]

    work_orders.sort(key=lambda item: (str(item.get("primary_value") or ""), str(item.get("id") or "")))
    counts = {
        "total": len(work_orders),
        "ready_for_archive": 0,
        "approved_archive": 0,
        "pending_review": 0,
        "returned_rework": 0,
        "evidence_gap": 0,
        "not_ready": 0,
        "exception": 0,
    }
    blockers: list[dict[str, Any]] = []
    ready_items: list[dict[str, str]] = []

    for work_order in work_orders:
        review_status = _platform_review_status(work_order)
        item = _delivery_archive_work_order_item(work_order)
        if review_status == "approved":
            counts["ready_for_archive"] += 1
            counts["approved_archive"] += 1
            ready_items.append({**item, "reason": "approved_archive"})
            continue
        if review_status == "returned":
            counts["returned_rework"] += 1
            blockers.append({**item, "reason": "returned_rework", "detail": str(work_order.get("review_reason") or work_order.get("review_note") or "")})
            continue
        if review_status == "exception":
            counts["exception"] += 1
            blockers.append({**item, "reason": "exception", "detail": str(work_order.get("review_reason") or work_order.get("review_note") or "")})
            continue
        if not _has_platform_collection(work_order):
            counts["not_ready"] += 1
            blockers.append({**item, "reason": "not_ready", "detail": "No construction collection has been submitted."})
            continue

        try:
            _validate_platform_required_collection(
                work_order,
                construction_fields,
                photo_slots,
                _stored_platform_collection_field_values(work_order),
                _stored_platform_covered_photo_slots(work_order),
            )
        except ValueError as exc:
            counts["evidence_gap"] += 1
            blockers.append({**item, "reason": "evidence_gap", "detail": str(exc)})
            continue

        if str(work_order.get("collection_status") or "").strip() != "submitted":
            counts["not_ready"] += 1
            blockers.append({**item, "reason": "not_ready", "detail": "Collection is still cached and has not been submitted."})
            continue

        counts["pending_review"] += 1
        blockers.append({**item, "reason": "pending_review", "detail": "Submitted collection is waiting for review."})

    blocked = (
        counts["pending_review"]
        + counts["returned_rework"]
        + counts["evidence_gap"]
        + counts["not_ready"]
        + counts["exception"]
    )
    ready = counts["total"] > 0 and blocked == 0 and counts["ready_for_archive"] == counts["total"]
    return {
        "project_id": project_id,
        **counts,
        "blocked": blocked,
        "ready": ready,
        "status": "ready" if ready else "blocked" if blocked else "empty",
        "next_actions": _delivery_archive_next_actions(counts),
        "blockers": blockers,
        "ready_items": ready_items,
    }


def build_platform_delivery_archive_manifest(project_id: str) -> dict[str, Any]:
    readiness = build_platform_delivery_archive_readiness(project_id)
    construction_schema = _platform_construction_schema(project_id)
    required_evidence = _delivery_archive_required_evidence(
        construction_schema["construction_fields"],
        construction_schema["photo_slots"],
        construction_schema["platform_required_fields"],
    )
    return {
        "project_id": project_id,
        "manifest_id": f"{project_id}:delivery-archive-preview",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "status": readiness["status"],
        "ready": readiness["ready"],
        "can_export": bool(readiness["ready"]),
        "total": readiness["total"],
        "ready_count": readiness["ready_for_archive"],
        "blocked_count": readiness["blocked"],
        "next_actions": readiness["next_actions"],
        "required_evidence": required_evidence,
        "sections": _delivery_archive_manifest_sections(readiness),
    }


def _delivery_archive_work_order_item(work_order: dict[str, Any]) -> dict[str, str]:
    return {
        "work_order_id": str(work_order.get("id") or ""),
        "primary_value": str(work_order.get("primary_value") or ""),
        "aggregate_value": str(work_order.get("aggregate_value") or ""),
    }


def _delivery_archive_next_actions(counts: dict[str, int]) -> list[str]:
    actions: list[str] = []
    if counts.get("ready_for_archive", 0) > 0:
        actions.append("archive_approved_work_orders")
    if counts.get("pending_review", 0) > 0:
        actions.append("review_pending_work_orders")
    if counts.get("returned_rework", 0) > 0:
        actions.append("resolve_returned_rework")
    if counts.get("evidence_gap", 0) > 0:
        actions.append("complete_evidence")
    if counts.get("not_ready", 0) > 0:
        actions.append("collect_not_ready_work_orders")
    if counts.get("exception", 0) > 0:
        actions.append("handle_exceptions")
    return actions


def _delivery_archive_manifest_sections(readiness: dict[str, Any]) -> list[dict[str, Any]]:
    ready_items = [
        item
        for item in readiness.get("ready_items", [])
        if isinstance(item, dict)
    ]
    blockers = [
        item
        for item in readiness.get("blockers", [])
        if isinstance(item, dict)
    ]
    blocker_by_reason: dict[str, list[dict[str, Any]]] = {}
    for blocker in blockers:
        reason = str(blocker.get("reason") or "").strip() or "blocked"
        blocker_by_reason.setdefault(reason, []).append(blocker)

    section_specs = [
        ("approved_archive", "已通过，可纳入交付包", ready_items),
        ("pending_review", "待审阅", blocker_by_reason.get("pending_review", [])),
        ("evidence_gap", "证据缺口", blocker_by_reason.get("evidence_gap", [])),
        ("returned_rework", "返工", blocker_by_reason.get("returned_rework", [])),
        ("not_ready", "未施工", blocker_by_reason.get("not_ready", [])),
        ("exception", "异常", blocker_by_reason.get("exception", [])),
    ]
    return [
        {
            "id": section_id,
            "title": title,
            "count": len(items),
            "items": deepcopy(items),
        }
        for section_id, title, items in section_specs
    ]


def _delivery_archive_required_evidence(
    construction_fields: list[dict[str, Any]],
    photo_slots: list[dict[str, Any]],
    platform_required_fields: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    return {
        "fields": [
            _delivery_archive_evidence_item(field)
            for field in _dedupe_fields(
                [
                    *[field for field in construction_fields if _field_is_archive_required_evidence(field)],
                    *_delivery_archive_platform_kpi_fields(platform_required_fields),
                ]
            )
        ],
        "photos": [
            _delivery_archive_evidence_item(slot)
            for slot in photo_slots
            if _field_is_archive_required_evidence(slot)
        ],
    }


def _field_is_archive_required_evidence(field: dict[str, Any]) -> bool:
    return bool(field.get("required") or field.get("required_when"))


def _delivery_archive_platform_kpi_fields(platform_required_fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
    platform_by_key = {
        str(field.get("key") or ""): field
        for field in platform_required_fields
        if isinstance(field, dict) and str(field.get("key") or "")
    }
    evidence_fields: list[dict[str, Any]] = []
    for key in _DELIVERY_ARCHIVE_PLATFORM_KPI_KEYS:
        field = deepcopy(platform_by_key.get(key) or _DELIVERY_ARCHIVE_COMPUTED_KPI_FIELDS.get(key) or {})
        if not field:
            continue
        field["required"] = True
        field["relation_role"] = str(field.get("relation_role") or "task_detail")
        evidence_fields.append(field)
    return evidence_fields


def _delivery_archive_evidence_item(field: dict[str, Any]) -> dict[str, Any]:
    return {
        "key": str(field.get("key") or ""),
        "label": str(field.get("label") or field.get("key") or ""),
        "capture_method": str(field.get("capture_method") or ""),
        "relation_role": str(field.get("relation_role") or ""),
        "required": bool(field.get("required")),
        "required_when": deepcopy(field.get("required_when") or None),
    }


def list_platform_construction_work_orders(project_id: str) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    primary = construction_schema["primary_field"]
    aggregate = construction_schema["aggregate_field"]
    display_fields = construction_schema["display_fields"]
    construction_fields = construction_schema["construction_fields"]
    photo_slots = construction_schema["photo_slots"]

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_orders = [
            deepcopy(work_order)
            for work_order in _PLATFORM_WORK_ORDERS.values()
            if work_order.get("project_id") == project_id
        ]

    work_orders.sort(key=lambda item: (str(item.get("primary_value") or ""), str(item.get("id") or "")))
    return {
        "project_id": project_id,
        "total": len(work_orders),
        "field_schema": {
            "primary_field": primary,
            "aggregate_field": aggregate,
            "display_fields": display_fields,
            "construction_fields": construction_fields,
            "photo_slots": photo_slots,
        },
        "items": [
            _construction_work_order_payload(work_order, construction_fields, photo_slots)
            for work_order in work_orders
        ],
    }


def list_platform_review_work_orders(project_id: str) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    construction_fields = construction_schema["construction_fields"]
    photo_slots = construction_schema["photo_slots"]

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_orders = [
            deepcopy(work_order)
            for work_order in _PLATFORM_WORK_ORDERS.values()
            if work_order.get("project_id") == project_id and _has_platform_collection(work_order)
        ]

    work_orders.sort(key=lambda item: (str(item.get("primary_value") or ""), str(item.get("id") or "")))
    items = [
        _platform_review_work_order_payload(work_order, construction_fields, photo_slots)
        for work_order in work_orders
    ]
    status_counts = {
        "pending_review": 0,
        "approved": 0,
        "returned": 0,
        "exception": 0,
        "not_ready": 0,
    }
    for item in items:
        status = str(item.get("review_status") or "not_ready")
        if status not in status_counts:
            status = "not_ready"
        status_counts[status] += 1
    return {
        "project_id": project_id,
        "total": len(items),
        "status_counts": status_counts,
        "field_schema": {
            "primary_field": construction_schema["primary_field"],
            "aggregate_field": construction_schema["aggregate_field"],
            "display_fields": construction_schema["display_fields"],
            "construction_fields": construction_fields,
            "photo_slots": photo_slots,
        },
        "items": items,
    }


def review_platform_work_order(project_id: str, work_order_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    construction_fields = construction_schema["construction_fields"]
    photo_slots = construction_schema["photo_slots"]
    action = str(payload.get("action") or "").strip()
    if action not in {"approved", "returned", "exception"}:
        raise ValueError("Unsupported platform review action")

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_order = _PLATFORM_WORK_ORDERS.get(work_order_id)
        if not work_order or work_order.get("project_id") != project_id:
            raise KeyError(work_order_id)
        if not _has_platform_collection(work_order):
            raise ValueError("Platform work order has no collection to review")
        if action == "approved":
            _validate_platform_required_collection(
                work_order,
                construction_fields,
                photo_slots,
                _stored_platform_collection_field_values(work_order),
                _stored_platform_covered_photo_slots(work_order),
            )
        reviewed_at = datetime.now().isoformat(timespec="seconds")
        actor = str(payload.get("actor") or "").strip()
        note = str(payload.get("note") or "").strip()
        reason = str(payload.get("reason") or "").strip()
        if action == "returned" and not reason:
            reason = _platform_review_return_reason_suggestion(work_order)
            if not note:
                note = reason
        review_history = work_order.get("review_history") if isinstance(work_order.get("review_history"), list) else []
        review_history = [event for event in review_history if isinstance(event, dict)]
        review_history.append(
            {
                "id": f"review-action-{uuid4().hex[:12]}",
                "action": action,
                "actor": actor,
                "reviewed_at": reviewed_at,
                "note": note,
                "reason": reason,
            }
        )
        work_order["review_status"] = action
        work_order["reviewed_by"] = actor
        work_order["reviewed_at"] = reviewed_at
        work_order["review_note"] = note
        work_order["review_reason"] = reason
        work_order["review_history"] = review_history
        _PLATFORM_WORK_ORDERS[work_order_id] = work_order
        _save_platform_work_orders_unlocked()
        updated = deepcopy(work_order)

    return _platform_review_work_order_payload(updated, construction_fields, photo_slots)


def save_platform_construction_work_order_collection(project_id: str, work_order_id: str, payload: dict[str, Any]) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    construction_fields = construction_schema["construction_fields"]
    photo_slots = construction_schema["photo_slots"]
    allowed_field_keys = {field["key"] for field in construction_fields}
    allowed_photo_keys = {field["key"] for field in photo_slots}
    status = str(payload.get("status") or "cached").strip() or "cached"
    if status not in {"cached", "submitted"}:
        raise ValueError("Unsupported platform construction collection status")

    raw_field_values = payload.get("field_values") if isinstance(payload.get("field_values"), dict) else {}
    collection_field_values = {
        str(key): "" if value is None else str(value).strip()
        for key, value in raw_field_values.items()
        if str(key) in allowed_field_keys
    }
    covered_photo_slots: list[str] = []
    for raw_slot in payload.get("covered_photo_slots", []) if isinstance(payload.get("covered_photo_slots"), list) else []:
        slot = str(raw_slot or "").strip()
        if slot and slot in allowed_photo_keys and slot not in covered_photo_slots:
            covered_photo_slots.append(slot)

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_order = _PLATFORM_WORK_ORDERS.get(work_order_id)
        if not work_order or work_order.get("project_id") != project_id:
            raise KeyError(work_order_id)
        if status == "submitted":
            _validate_platform_required_collection(
                work_order,
                construction_fields,
                photo_slots,
                collection_field_values,
                covered_photo_slots,
            )
        collected_at = datetime.now().isoformat(timespec="seconds")
        work_order["collection_status"] = status
        work_order["collection_field_values"] = collection_field_values
        work_order["covered_photo_slots"] = covered_photo_slots
        work_order["client_batch_id"] = str(payload.get("client_batch_id") or "").strip()
        work_order["collected_by"] = str(payload.get("actor") or "").strip()
        work_order["collected_at"] = collected_at
        work_order["kpi_values"] = _build_platform_kpi_values(
            work_order,
            raw_field_values=raw_field_values,
            covered_photo_slots=covered_photo_slots,
            actor=work_order["collected_by"],
            collected_at=collected_at,
        )
        if status == "submitted" and str(work_order.get("review_status") or "").strip() == "returned":
            review_history = work_order.get("review_history") if isinstance(work_order.get("review_history"), list) else []
            review_history = [event for event in review_history if isinstance(event, dict)]
            review_history.append(
                {
                    "id": f"review-action-{uuid4().hex[:12]}",
                    "action": "rework_submitted",
                    "actor": work_order["collected_by"],
                    "reviewed_at": collected_at,
                    "note": "\u8fd4\u5de5\u8865\u91c7\u540e\u91cd\u65b0\u63d0\u4ea4\u5ba1\u9605",
                    "reason": "returned_rework_resubmitted",
                }
            )
            work_order["review_history"] = review_history
            work_order["review_status"] = "pending_review"
            work_order["reviewed_by"] = ""
            work_order["reviewed_at"] = ""
            work_order["review_note"] = ""
            work_order["review_reason"] = ""
        _PLATFORM_WORK_ORDERS[work_order_id] = work_order
        _save_platform_work_orders_unlocked()
        updated = deepcopy(work_order)

    return _construction_work_order_payload(updated, construction_fields, photo_slots)


def _validate_platform_required_collection(
    work_order: dict[str, Any],
    construction_fields: list[dict[str, Any]],
    photo_slots: list[dict[str, Any]],
    collection_field_values: dict[str, str],
    covered_photo_slots: list[str],
) -> None:
    initial_values = {
        str(key): "" if value is None else str(value).strip()
        for key, value in (work_order.get("values") if isinstance(work_order.get("values"), dict) else {}).items()
    }
    merged_values = {**initial_values, **collection_field_values}
    missing_fields = [
        _field_label(field)
        for field in construction_fields
        if _is_platform_collection_item_required(field, merged_values)
        and not str(merged_values.get(str(field.get("key") or ""), "")).strip()
    ]
    covered_slots = set(covered_photo_slots)
    missing_photos = [
        _field_label(slot)
        for slot in photo_slots
        if _is_platform_collection_item_required(slot, merged_values)
        and str(slot.get("key") or "") not in covered_slots
    ]
    if not missing_fields and not missing_photos:
        return
    details: list[str] = []
    if missing_fields:
        details.append(f"缺少必填字段：{'、'.join(missing_fields)}")
    if missing_photos:
        details.append(f"缺少必填照片：{'、'.join(missing_photos)}")
    raise ValueError("；".join(details))


def _field_label(field: dict[str, Any]) -> str:
    return str(field.get("label") or field.get("key") or "").strip() or "未命名字段"


def _stored_platform_collection_field_values(work_order: dict[str, Any]) -> dict[str, str]:
    values = work_order.get("collection_field_values") if isinstance(work_order.get("collection_field_values"), dict) else {}
    return {str(key): "" if value is None else str(value).strip() for key, value in values.items()}


def _stored_platform_covered_photo_slots(work_order: dict[str, Any]) -> list[str]:
    slots: list[str] = []
    stored_slots = work_order.get("covered_photo_slots") if isinstance(work_order.get("covered_photo_slots"), list) else []
    for raw_slot in stored_slots:
        slot = str(raw_slot or "").strip()
        if slot and slot not in slots:
            slots.append(slot)
    photos = work_order.get("collection_photos") if isinstance(work_order.get("collection_photos"), list) else []
    for slot in _covered_slots_from_photos([photo for photo in photos if isinstance(photo, dict)]):
        if slot not in slots:
            slots.append(slot)
    return slots


def _is_platform_collection_item_required(field: dict[str, Any], values: dict[str, str]) -> bool:
    required_when = field.get("required_when")
    if isinstance(required_when, dict) and str(required_when.get("field_key") or "").strip():
        return _platform_required_when_matches(required_when, values)
    return bool(field.get("required"))


def _platform_required_when_matches(required_when: dict[str, Any], values: dict[str, str]) -> bool:
    field_key = str(required_when.get("field_key") or "").strip()
    current = str(values.get(field_key) or "").strip()
    if not field_key or not current:
        return False
    expected = required_when.get("equals")
    if isinstance(expected, list):
        return current in {str(item).strip() for item in expected if str(item).strip()}
    return current == str(expected or "").strip()


def upload_platform_construction_work_order_photo(
    project_id: str,
    work_order_id: str,
    *,
    slot: str,
    filename: str,
    content_type: str,
    content: bytes,
    actor: str = "",
    client_photo_id: str = "",
) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    photo_slots = construction_schema["photo_slots"]
    allowed_photo_keys = {field["key"] for field in photo_slots}
    clean_slot = str(slot or "").strip()
    if clean_slot not in allowed_photo_keys:
        raise ValueError("Unknown platform construction photo slot")
    if not content:
        raise ValueError("Photo file is empty")

    photo_id = f"photo-{uuid4().hex}"
    suffix = _safe_photo_suffix(filename, content_type)
    storage_key = f"platform-construction-photos/{project_id}/{work_order_id}/{photo_id}{suffix}"
    storage_path = _platform_photo_storage_root() / project_id / work_order_id / f"{photo_id}{suffix}"
    storage_path.parent.mkdir(parents=True, exist_ok=True)
    storage_path.write_bytes(content)

    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_order = _PLATFORM_WORK_ORDERS.get(work_order_id)
        if not work_order or work_order.get("project_id") != project_id:
            if storage_path.exists():
                storage_path.unlink()
            raise KeyError(work_order_id)
        photos = work_order.get("collection_photos") if isinstance(work_order.get("collection_photos"), list) else []
        photos = [photo for photo in photos if isinstance(photo, dict)]
        photos.append(
            {
                "id": photo_id,
                "slot": clean_slot,
                "client_photo_id": str(client_photo_id or "").strip(),
                "filename": Path(filename or "photo").name or f"{photo_id}{suffix}",
                "content_type": content_type or "application/octet-stream",
                "size": len(content),
                "storage": "local_platform_store",
                "storage_key": storage_key,
                "uploaded_by": str(actor or "").strip(),
                "uploaded_at": datetime.now().isoformat(timespec="seconds"),
            }
        )
        work_order["collection_photos"] = photos
        work_order["covered_photo_slots"] = _covered_slots_from_photos(photos)
        work_order["collection_status"] = work_order.get("collection_status") or "cached"
        _PLATFORM_WORK_ORDERS[work_order_id] = work_order
        _save_platform_work_orders_unlocked()
        updated = deepcopy(work_order)

    return _construction_work_order_payload(updated, construction_schema["construction_fields"], photo_slots)


def delete_platform_construction_work_order_photo(project_id: str, work_order_id: str, photo_id: str) -> dict[str, Any]:
    construction_schema = _platform_construction_schema(project_id)
    with _PLATFORM_WORK_ORDERS_LOCK:
        _load_platform_work_orders_unlocked()
        work_order = _PLATFORM_WORK_ORDERS.get(work_order_id)
        if not work_order or work_order.get("project_id") != project_id:
            raise KeyError(work_order_id)
        photos = work_order.get("collection_photos") if isinstance(work_order.get("collection_photos"), list) else []
        kept_photos: list[dict[str, Any]] = []
        deleted_photo: dict[str, Any] | None = None
        for photo in photos:
            if not isinstance(photo, dict):
                continue
            if str(photo.get("id") or "") == photo_id:
                deleted_photo = photo
            else:
                kept_photos.append(photo)
        if deleted_photo is None:
            raise KeyError(photo_id)
        storage_key = str(deleted_photo.get("storage_key") or "")
        storage_path = _platform_photo_storage_root() / storage_key.replace("platform-construction-photos/", "", 1)
        if storage_path.exists() and _is_relative_to(storage_path.resolve(), _platform_photo_storage_root().resolve()):
            storage_path.unlink()
        work_order["collection_photos"] = kept_photos
        work_order["covered_photo_slots"] = _covered_slots_from_photos(kept_photos)
        _PLATFORM_WORK_ORDERS[work_order_id] = work_order
        _save_platform_work_orders_unlocked()
        updated = deepcopy(work_order)

    return _construction_work_order_payload(updated, construction_schema["construction_fields"], construction_schema["photo_slots"])


def project_template_filename(project_id: str, template_type: str) -> str:
    return f"{project_id}-{template_type}.xlsx"


def _platform_construction_schema(project_id: str) -> dict[str, Any]:
    project = get_project_overview(project_id)
    schema = project.get("work_item_schema") or {}
    primary = _clean_field(schema.get("primary_field"))
    aggregate = _clean_field(schema.get("aggregate_field"))
    custom_fields = [_clean_field(field) for field in schema.get("custom_fields", []) if isinstance(field, dict)]
    platform_required_fields = [_clean_field(field) for field in schema.get("platform_required_fields", []) if isinstance(field, dict)]
    display_fields = _dedupe_fields(
        [
            {**primary, "show_in_construction_panel": True},
            {**aggregate, "show_in_construction_panel": True},
            *[
                field
                for field in custom_fields
                if field.get("show_in_construction_panel")
                and field.get("source") != "field_collection"
                and field.get("data_type") != "image"
            ],
        ]
    )
    construction_fields = [
        field
        for field in custom_fields
        if field.get("show_in_construction_panel")
        and field.get("source") == "field_collection"
        and field.get("data_type") != "image"
    ]
    photo_slots = [
        field
        for field in custom_fields
        if field.get("show_in_construction_panel")
        and field.get("source") == "field_collection"
        and field.get("data_type") == "image"
    ]
    return {
        "primary_field": primary,
        "aggregate_field": aggregate,
        "display_fields": display_fields,
        "construction_fields": construction_fields,
        "photo_slots": photo_slots,
        "platform_required_fields": platform_required_fields,
    }


def _construction_work_order_payload(
    work_order: dict[str, Any],
    construction_fields: list[dict[str, Any]],
    photo_slots: list[dict[str, Any]],
) -> dict[str, Any]:
    values = work_order.get("values") if isinstance(work_order.get("values"), dict) else {}
    kpi_values = _platform_kpi_values_for_work_order(work_order, photo_slots=photo_slots)
    return {
        "id": str(work_order.get("id") or ""),
        "project_id": str(work_order.get("project_id") or ""),
        "source_task_id": str(work_order.get("source_task_id") or ""),
        "source_batch_id": str(work_order.get("source_batch_id") or ""),
        "primary_value": str(work_order.get("primary_value") or ""),
        "aggregate_value": str(work_order.get("aggregate_value") or ""),
        "status": str(work_order.get("status") or "created"),
        "created_at": str(work_order.get("created_at") or ""),
        "created_by": str(work_order.get("created_by") or ""),
        "field_values": {str(key): "" if value is None else str(value) for key, value in values.items()},
        "required_fields": [field for field in construction_fields if field.get("required")],
        "photo_slots": photo_slots,
        "collection_status": str(work_order.get("collection_status") or ""),
        "collection_field_values": {
            str(key): "" if value is None else str(value)
            for key, value in (work_order.get("collection_field_values") if isinstance(work_order.get("collection_field_values"), dict) else {}).items()
        },
        "covered_photo_slots": [
            str(slot)
            for slot in (work_order.get("covered_photo_slots") if isinstance(work_order.get("covered_photo_slots"), list) else [])
            if str(slot)
        ],
        "collection_photos": [
            _platform_photo_payload(photo)
            for photo in (work_order.get("collection_photos") if isinstance(work_order.get("collection_photos"), list) else [])
            if isinstance(photo, dict)
        ],
        "client_batch_id": str(work_order.get("client_batch_id") or ""),
        "collected_by": str(work_order.get("collected_by") or ""),
        "collected_at": str(work_order.get("collected_at") or ""),
        "review_status": _platform_review_status(work_order),
        "reviewed_by": str(work_order.get("reviewed_by") or ""),
        "reviewed_at": str(work_order.get("reviewed_at") or ""),
        "review_note": str(work_order.get("review_note") or ""),
        "review_reason": str(work_order.get("review_reason") or ""),
        "review_history": _platform_review_history_payload(work_order),
        "rework_evidence_gap_groups": _platform_rework_evidence_gap_groups(work_order, construction_fields, photo_slots),
        "kpi_values": kpi_values,
    }


def _build_platform_kpi_values(
    work_order: dict[str, Any],
    *,
    raw_field_values: dict[str, Any] | None = None,
    covered_photo_slots: list[str] | None = None,
    actor: str = "",
    collected_at: str = "",
) -> dict[str, str]:
    raw_values = raw_field_values if isinstance(raw_field_values, dict) else {}
    existing = work_order.get("kpi_values") if isinstance(work_order.get("kpi_values"), dict) else {}
    collected_values = work_order.get("collection_field_values") if isinstance(work_order.get("collection_field_values"), dict) else {}
    values = work_order.get("values") if isinstance(work_order.get("values"), dict) else {}

    def pick(*keys: str, fallback: str = "") -> str:
        for key in keys:
            for source in (raw_values, existing, collected_values, values):
                value = source.get(key) if isinstance(source, dict) else None
                clean = "" if value is None else str(value).strip()
                if clean:
                    return clean
        return fallback

    photo_count = pick("photo_count")
    if not photo_count:
        slots = covered_photo_slots
        if slots is None:
            slots = work_order.get("covered_photo_slots") if isinstance(work_order.get("covered_photo_slots"), list) else []
        photos = work_order.get("collection_photos") if isinstance(work_order.get("collection_photos"), list) else []
        photo_count = str(len({str(slot) for slot in slots if str(slot).strip()}) + len([photo for photo in photos if isinstance(photo, dict)]))

    old_device_recovered = "1" if _platform_old_device_recovered(raw_values, collected_values, covered_photo_slots or []) else "0"

    return {
        "installer": pick("installer", fallback=str(actor or "").strip()),
        "started_at": pick("started_at", "installed_at"),
        "completed_at": pick("completed_at", fallback=str(collected_at or "").strip()),
        "uploaded_at": pick("uploaded_at", fallback=str(collected_at or "").strip()),
        "online_duration_minutes": pick("online_duration_minutes", "online_minutes", "online_time"),
        "photo_count": photo_count,
        "old_device_recovered": old_device_recovered,
    }


def _platform_kpi_values_for_work_order(work_order: dict[str, Any], photo_slots: list[dict[str, Any]] | None = None) -> dict[str, str]:
    kpi_values = _build_platform_kpi_values(
        work_order,
        raw_field_values=work_order.get("kpi_values") if isinstance(work_order.get("kpi_values"), dict) else {},
        covered_photo_slots=work_order.get("covered_photo_slots") if isinstance(work_order.get("covered_photo_slots"), list) else [],
        actor=str(work_order.get("collected_by") or work_order.get("created_by") or ""),
        collected_at=str(work_order.get("collected_at") or work_order.get("created_at") or ""),
    )
    if photo_slots is not None and not kpi_values.get("old_device_recovered"):
        covered_slots = work_order.get("covered_photo_slots") if isinstance(work_order.get("covered_photo_slots"), list) else []
        kpi_values["old_device_recovered"] = "1" if _platform_old_device_recovered({}, {}, covered_slots) else "0"
    return kpi_values


def _platform_old_device_recovered(
    raw_values: dict[str, Any],
    collected_values: dict[str, Any],
    covered_photo_slots: list[str],
) -> bool:
    for source in (raw_values, collected_values):
        for key, value in source.items():
            normalized_key = str(key or "").lower()
            if ("old_device" in normalized_key or "old_terminal" in normalized_key or "old_meter" in normalized_key) and str(value or "").strip():
                return True
    return any("old_device" in str(slot or "").lower() or "old_terminal" in str(slot or "").lower() for slot in covered_photo_slots)


def _safe_int(value: Any) -> int:
    try:
        return int(float(str(value or "").strip()))
    except (TypeError, ValueError):
        return 0


def _platform_photo_payload(photo: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(photo.get("id") or ""),
        "slot": str(photo.get("slot") or ""),
        "client_photo_id": str(photo.get("client_photo_id") or ""),
        "filename": str(photo.get("filename") or ""),
        "content_type": str(photo.get("content_type") or ""),
        "size": int(photo.get("size") or 0),
        "storage": str(photo.get("storage") or ""),
        "storage_key": str(photo.get("storage_key") or ""),
        "uploaded_by": str(photo.get("uploaded_by") or ""),
        "uploaded_at": str(photo.get("uploaded_at") or ""),
    }


def _platform_review_work_order_payload(
    work_order: dict[str, Any],
    construction_fields: list[dict[str, Any]],
    photo_slots: list[dict[str, Any]],
) -> dict[str, Any]:
    construction_payload = _construction_work_order_payload(work_order, construction_fields, photo_slots)
    initial_values = construction_payload["field_values"]
    collected_values = construction_payload["collection_field_values"]
    covered_photo_slots = construction_payload["covered_photo_slots"]
    return {
        **construction_payload,
        "review_status": _platform_review_status(work_order),
        "reviewed_by": str(work_order.get("reviewed_by") or ""),
        "reviewed_at": str(work_order.get("reviewed_at") or ""),
        "review_note": str(work_order.get("review_note") or ""),
        "review_reason": str(work_order.get("review_reason") or ""),
        "review_history": _platform_review_history_payload(work_order),
        "suggested_review_return_reason": _platform_review_return_reason_suggestion(work_order),
        "review_hierarchy_gap_items": _platform_review_hierarchy_gap_items(work_order),
        "field_reviews": [
            {
                "key": field["key"],
                "label": field.get("label") or field["key"],
                "capture_method": field.get("capture_method") or "manual",
                "required": bool(field.get("required")),
                "required_when": field.get("required_when") or None,
                "relation_role": str(field.get("relation_role") or ""),
                "initial_value": initial_values.get(field["key"], ""),
                "collected_value": collected_values.get(field["key"], ""),
            }
            for field in construction_fields
        ],
        "photo_slot_reviews": [
            {
                "key": slot["key"],
                "label": slot.get("label") or slot["key"],
                "required": bool(slot.get("required")),
                "required_when": slot.get("required_when") or None,
                "relation_role": str(slot.get("relation_role") or ""),
                "covered": slot["key"] in covered_photo_slots,
                "photo_count": sum(
                    1
                    for photo in construction_payload["collection_photos"]
                    if photo.get("slot") == slot["key"]
                ),
            }
            for slot in photo_slots
        ],
    }


def _platform_review_history_payload(work_order: dict[str, Any]) -> list[dict[str, str]]:
    history = work_order.get("review_history") if isinstance(work_order.get("review_history"), list) else []
    return [
        {
            "id": str(event.get("id") or ""),
            "action": str(event.get("action") or ""),
            "actor": str(event.get("actor") or ""),
            "reviewed_at": str(event.get("reviewed_at") or ""),
            "note": str(event.get("note") or ""),
            "reason": str(event.get("reason") or ""),
        }
        for event in history
        if isinstance(event, dict)
    ]


def _platform_hierarchy_gap_items_from_warnings(row: dict[str, Any]) -> list[dict[str, Any]]:
    warnings = row.get("warnings") if isinstance(row.get("warnings"), list) else []
    return _platform_normalized_hierarchy_gap_items(warnings, require_conditional_code=True)


def _platform_review_hierarchy_gap_items(work_order: dict[str, Any]) -> list[dict[str, Any]]:
    raw_items = work_order.get("review_hierarchy_gap_items")
    if not isinstance(raw_items, list):
        return []
    return _platform_normalized_hierarchy_gap_items(raw_items)


def _platform_review_return_reason_suggestion(work_order: dict[str, Any]) -> str:
    hierarchy_gap_items = _platform_review_hierarchy_gap_items(work_order)
    if not hierarchy_gap_items:
        return ""
    labels: list[str] = []
    for item in hierarchy_gap_items:
        label = str(item.get("field_label") or item.get("field_key") or "").strip()
        if label and label not in labels:
            labels.append(label)
    if not labels:
        return ""
    visible_labels = labels[:6]
    suffix = "等" if len(labels) > len(visible_labels) else ""
    return f"导入层级缺口：缺少{'、'.join(visible_labels)}{suffix}，请补齐对应现场证据后重新提交。"


def _platform_normalized_hierarchy_gap_items(
    raw_items: list[Any],
    *,
    require_conditional_code: bool = False,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for raw_item in raw_items:
        if not isinstance(raw_item, dict):
            continue
        if require_conditional_code and raw_item.get("code") != "missing_conditional_field":
            continue
        row = _safe_optional_int(raw_item.get("row"))
        field_key = str(raw_item.get("field_key") or "").strip()
        field_label = str(raw_item.get("field_label") or field_key).strip()
        message = str(raw_item.get("message") or "").strip()
        value = str(raw_item.get("value") or "").strip()
        if not field_key and not field_label and not message:
            continue
        items.append(
            {
                "row": row,
                "field_key": field_key,
                "field_label": field_label,
                "message": message,
                "value": value,
            }
        )
    return items


def _safe_optional_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def _has_platform_collection(work_order: dict[str, Any]) -> bool:
    return bool(
        work_order.get("collection_status")
        or work_order.get("collection_field_values")
        or work_order.get("collection_photos")
        or work_order.get("covered_photo_slots")
    )


def _platform_review_status(work_order: dict[str, Any]) -> str:
    stored_status = str(work_order.get("review_status") or "").strip()
    if stored_status in {"approved", "returned", "exception"}:
        return stored_status
    if str(work_order.get("collection_status") or "") == "submitted":
        return "pending_review"
    if work_order.get("collection_photos") or work_order.get("collection_field_values"):
        return "pending_review"
    return "not_ready"


def _dedupe_non_empty_strings(values: list[str]) -> list[str]:
    deduped: list[str] = []
    for value in values:
        clean = str(value or "").strip()
        if clean and clean not in deduped:
            deduped.append(clean)
    return deduped


def _platform_rework_evidence_gap_groups(
    work_order: dict[str, Any],
    construction_fields: list[dict[str, Any]],
    photo_slots: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if _platform_review_status(work_order) != "returned":
        return []

    groups: list[dict[str, Any]] = []
    imported_gap_labels = _dedupe_non_empty_strings(
        [
            str(item.get("field_label") or item.get("field_key") or "").strip()
            for item in _platform_review_hierarchy_gap_items(work_order)
            if isinstance(item, dict)
        ]
    )
    if imported_gap_labels:
        groups.append({"label": "\u5bfc\u5165\u5c42\u7ea7\u7f3a\u53e3", "items": imported_gap_labels})

    initial_values = {
        str(key): "" if value is None else str(value).strip()
        for key, value in (work_order.get("values") if isinstance(work_order.get("values"), dict) else {}).items()
    }
    collected_values = _stored_platform_collection_field_values(work_order)
    merged_values = {**initial_values, **collected_values}
    missing_field_labels = _dedupe_non_empty_strings(
        [
            _field_label(field)
            for field in construction_fields
            if _is_platform_collection_item_required(field, merged_values)
            and not str(merged_values.get(str(field.get("key") or ""), "")).strip()
        ]
    )
    if missing_field_labels:
        groups.append({"label": "\u7f3a\u5c11\u5b57\u6bb5", "items": missing_field_labels})

    covered_slots = set(_stored_platform_covered_photo_slots(work_order))
    missing_photo_labels = _dedupe_non_empty_strings(
        [
            _field_label(slot)
            for slot in photo_slots
            if _is_platform_collection_item_required(slot, merged_values)
            and str(slot.get("key") or "") not in covered_slots
        ]
    )
    if missing_photo_labels:
        groups.append({"label": "\u7f3a\u5c11\u7167\u7247", "items": missing_photo_labels})

    return groups


def _covered_slots_from_photos(photos: list[dict[str, Any]]) -> list[str]:
    slots: list[str] = []
    for photo in photos:
        slot = str(photo.get("slot") or "").strip()
        if slot and slot not in slots:
            slots.append(slot)
    return slots


def _import_batches_store_path() -> Path:
    configured = settings.platform_project_drafts_path.strip()
    if configured:
        return Path(configured).with_name("platform-import-batches.json")
    return Path.cwd() / "data" / "platform-import-batches.json"


def _work_order_tasks_store_path() -> Path:
    configured = settings.platform_project_drafts_path.strip()
    if configured:
        return Path(configured).with_name("platform-import-work-order-tasks.json")
    return Path.cwd() / "data" / "platform-import-work-order-tasks.json"


def _platform_work_orders_store_path() -> Path:
    configured = settings.platform_project_drafts_path.strip()
    if configured:
        return Path(configured).with_name("platform-work-orders.json")
    return Path.cwd() / "data" / "platform-work-orders.json"


def _platform_photo_storage_root() -> Path:
    configured = settings.platform_project_drafts_path.strip()
    if configured:
        return Path(configured).with_name("platform-construction-photos")
    return Path.cwd() / "data" / "platform-construction-photos"


def _safe_photo_suffix(filename: str, content_type: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png", ".webp", ".heic"}:
        return suffix
    if content_type == "image/png":
        return ".png"
    if content_type == "image/webp":
        return ".webp"
    return ".jpg"


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def _load_import_batches_unlocked() -> None:
    global _IMPORT_BATCHES_LOADED
    if _IMPORT_BATCHES_LOADED:
        return
    path = _import_batches_store_path()
    if not path.exists():
        _IMPORT_BATCHES_LOADED = True
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    raw_batches = payload.get("batches") if isinstance(payload, dict) else None
    if not isinstance(raw_batches, list):
        raw_batches = []
    _IMPORT_BATCHES.clear()
    for raw_batch in raw_batches:
        if not isinstance(raw_batch, dict):
            continue
        batch_id = str(raw_batch.get("job_id") or "").strip()
        if batch_id:
            _IMPORT_BATCHES[batch_id] = raw_batch
    _IMPORT_BATCHES_LOADED = True


def _save_import_batches_unlocked() -> None:
    path = _import_batches_store_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 1,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "batches": [_IMPORT_BATCHES[batch_id] for batch_id in sorted(_IMPORT_BATCHES)],
    }
    tmp_path = path.with_name(f"{path.name}.tmp")
    tmp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp_path.replace(path)


def _load_work_order_tasks_unlocked() -> None:
    global _WORK_ORDER_TASKS_LOADED
    if _WORK_ORDER_TASKS_LOADED:
        return
    path = _work_order_tasks_store_path()
    if not path.exists():
        _WORK_ORDER_TASKS_LOADED = True
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    raw_tasks = payload.get("tasks") if isinstance(payload, dict) else None
    if not isinstance(raw_tasks, list):
        raw_tasks = []
    _WORK_ORDER_TASKS.clear()
    for raw_task in raw_tasks:
        if not isinstance(raw_task, dict):
            continue
        task_id = str(raw_task.get("job_id") or "").strip()
        if task_id:
            _WORK_ORDER_TASKS[task_id] = raw_task
    _WORK_ORDER_TASKS_LOADED = True


def _save_work_order_tasks_unlocked() -> None:
    path = _work_order_tasks_store_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 1,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "tasks": [_WORK_ORDER_TASKS[task_id] for task_id in sorted(_WORK_ORDER_TASKS)],
    }
    tmp_path = path.with_name(f"{path.name}.tmp")
    tmp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp_path.replace(path)


def _load_platform_work_orders_unlocked() -> None:
    global _PLATFORM_WORK_ORDERS_LOADED
    if _PLATFORM_WORK_ORDERS_LOADED:
        return
    path = _platform_work_orders_store_path()
    if not path.exists():
        _PLATFORM_WORK_ORDERS_LOADED = True
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    raw_work_orders = payload.get("work_orders") if isinstance(payload, dict) else None
    if not isinstance(raw_work_orders, list):
        raw_work_orders = []
    _PLATFORM_WORK_ORDERS.clear()
    for raw_work_order in raw_work_orders:
        if not isinstance(raw_work_order, dict):
            continue
        work_order_id = str(raw_work_order.get("id") or "").strip()
        if work_order_id:
            _PLATFORM_WORK_ORDERS[work_order_id] = raw_work_order
    _PLATFORM_WORK_ORDERS_LOADED = True


def _save_platform_work_orders_unlocked() -> None:
    path = _platform_work_orders_store_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 1,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "work_orders": [_PLATFORM_WORK_ORDERS[work_order_id] for work_order_id in sorted(_PLATFORM_WORK_ORDERS)],
    }
    tmp_path = path.with_name(f"{path.name}.tmp")
    tmp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp_path.replace(path)


def _fields_for_template(schema: dict[str, Any], template_type: str) -> list[dict[str, Any]]:
    primary = _clean_field(schema.get("primary_field"))
    aggregate = _clean_field(schema.get("aggregate_field"))
    custom_fields = [_clean_field(field) for field in schema.get("custom_fields", []) if isinstance(field, dict)]
    platform_fields = [_clean_field(field) for field in schema.get("platform_required_fields", []) if isinstance(field, dict)]

    import_fields = [field for field in custom_fields if field.get("source") == "import"]
    field_collection_fields = [field for field in custom_fields if field.get("source") == "field_collection"]
    platform_by_key = {field["key"]: field for field in platform_fields}

    if template_type == "initial_work_orders":
        return _dedupe_fields([primary, aggregate, *import_fields])
    return _dedupe_fields(
        [
            primary,
            aggregate,
            *import_fields,
            *field_collection_fields,
            platform_by_key.get("installer"),
            platform_by_key.get("completed_at"),
            {"key": "external_evidence", "label": "外部完成证明", "source": "import", "capture_method": "manual", "data_type": "text", "required": False},
        ]
    )


def _field_rows_for_template(schema: dict[str, Any], template_fields: list[dict[str, Any]], template_type: str) -> list[dict[str, Any]]:
    primary = _clean_field(schema.get("primary_field"))
    aggregate = _clean_field(schema.get("aggregate_field"))
    if template_type != "external_completed":
        return template_fields
    platform_fields = [_clean_field(field) for field in schema.get("platform_required_fields", []) if isinstance(field, dict)]
    platform_by_key = {field["key"]: field for field in platform_fields}
    return _dedupe_fields(
        [
            *template_fields,
            platform_by_key.get("work_order_id"),
            platform_by_key.get("uploaded_at"),
            platform_by_key.get("photo_count"),
            platform_by_key.get("online_duration_minutes"),
        ]
    )


def _clean_field(raw: Any) -> dict[str, Any]:
    field = raw if isinstance(raw, dict) else {}
    key = str(field.get("key") or "").strip()
    label = str(field.get("label") or key).strip() or key
    clean = {
        "key": key,
        "label": label,
        "source": str(field.get("source") or "").strip(),
        "capture_method": str(field.get("capture_method") or field.get("captureMethod") or "").strip(),
        "data_type": str(field.get("data_type") or field.get("dataType") or "").strip(),
        "required": bool(field.get("required", False)),
        "parent_key": str(field.get("parent_key") or field.get("parentKey") or "").strip(),
        "show_in_construction_panel": bool(
            field.get(
                "show_in_construction_panel",
                field.get("showInConstructionPanel", str(field.get("source") or "").strip() == "field_collection"),
            )
        ),
    }
    options = field.get("options")
    if isinstance(options, list):
        clean_options = [str(option).strip() for option in options if str(option).strip()]
        if clean_options:
            clean["options"] = clean_options
    required_when = _clean_required_when(field.get("required_when") or field.get("requiredWhen"))
    if required_when:
        clean["required_when"] = required_when
    relation_role = str(field.get("relation_role") or field.get("relationRole") or "").strip()
    if relation_role:
        clean["relation_role"] = relation_role
    return clean


def _clean_required_when(raw_required_when: Any) -> dict[str, Any] | None:
    raw = raw_required_when if isinstance(raw_required_when, dict) else {}
    field_key = str(raw.get("field_key") or raw.get("fieldKey") or "").strip()
    if not field_key:
        return None
    raw_equals = raw.get("equals")
    if isinstance(raw_equals, list):
        equals = [str(item).strip() for item in raw_equals if str(item).strip()]
        return {"field_key": field_key, "equals": equals} if equals else None
    equals = str(raw_equals or "").strip()
    return {"field_key": field_key, "equals": equals} if equals else None


def _dedupe_fields(fields: list[dict[str, Any] | None]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    result: list[dict[str, Any]] = []
    for field in fields:
        if not field:
            continue
        key = str(field.get("key") or "").strip()
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(field)
    return result


def _platform_fill_rule(field: dict[str, Any], template_type: str) -> str:
    if template_type != "external_completed":
        return ""
    return PLATFORM_FILL_RULES.get(str(field.get("key") or ""), "")


def _template_field_labels(schema: dict[str, Any]) -> dict[str, str]:
    fields = [
        _clean_field(schema.get("primary_field")),
        _clean_field(schema.get("aggregate_field")),
        *[_clean_field(field) for field in schema.get("custom_fields", []) if isinstance(field, dict)],
        *[_clean_field(field) for field in schema.get("platform_required_fields", []) if isinstance(field, dict)],
    ]
    return {
        field["key"]: field["label"]
        for field in fields
        if field.get("key") and field.get("label")
    }


def _template_hierarchy_role(field: dict[str, Any]) -> str:
    if field.get("required_when"):
        return "条件采集"
    role = str(field.get("relation_role") or "").strip()
    role_labels = {
        "aggregate": "聚合字段",
        "task_object": "任务对象",
        "task_detail": "任务核心",
        "replacement_device": "主设备更换",
        "old_device": "旧设备",
        "accessory_replace_confirm": "附属设备确认",
        "accessory_new_device": "附属设备更换",
        "evidence_photo": "照片证据",
        "supporting_field": "辅助字段",
    }
    if role in role_labels:
        return role_labels[role]
    source = str(field.get("source") or "").strip()
    source_labels = {
        "import": "导入字段",
        "field_collection": "现场采集",
        "review": "审阅补录",
        "system": "平台生成",
    }
    return source_labels.get(source, "业务字段")


def _template_parent_label(schema: dict[str, Any], field: dict[str, Any]) -> str:
    parent_key = str(field.get("parent_key") or "").strip()
    if not parent_key:
        return ""
    return _template_field_labels(schema).get(parent_key, parent_key)


def _template_condition_hint(schema: dict[str, Any], field: dict[str, Any]) -> str:
    required_when = field.get("required_when") if isinstance(field.get("required_when"), dict) else {}
    field_key = str(required_when.get("field_key") or "").strip()
    if not field_key:
        return ""
    equals = required_when.get("equals")
    if isinstance(equals, list):
        value = "/".join(str(item).strip() for item in equals if str(item).strip())
    else:
        value = str(equals or "").strip()
    parent_label = _template_field_labels(schema).get(field_key, field_key)
    return f"{parent_label}={value} 时采集" if value else f"{parent_label} 触发时采集"


def _example_value(field: dict[str, Any], index: int) -> str:
    key = str(field.get("key") or "")
    data_type = str(field.get("data_type") or "")
    if key in {"terminal_no", "terminal"}:
        return "TT-001"
    if key in {"station_area", "area", "area_no"}:
        return "台区-001"
    if key in {"meter_no", "work_item"}:
        return "370100000001"
    if data_type == "datetime":
        return "2026-06-30 09:30:00"
    if data_type in {"number", "duration"}:
        return "15"
    return f"示例值{index}"


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="E8F1FF")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width, 12), 34)


def _template_data_rows(workbook_content: bytes) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    workbook = load_workbook(BytesIO(workbook_content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_clean_cell(value) for value in rows[0]]
    data_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows[1:], start=2)
        if not _is_empty_row(row) and not _is_field_code_row(row)
    ]
    return headers, data_rows


def _platform_fill_policies(schema: dict[str, Any], fields: list[dict[str, Any]], template_type: str) -> list[dict[str, str]]:
    if template_type != "external_completed":
        return []
    field_rows = _field_rows_for_template(schema, fields, template_type)
    policies: list[dict[str, str]] = []
    for field in field_rows:
        rule = _platform_fill_rule(field, template_type)
        if not rule:
            continue
        policies.append({"field_key": field["key"], "field_label": field["label"], "rule": rule})
    return policies


def _append_template_instruction_sheet(
    sheet,
    schema: dict[str, Any],
    fields: list[dict[str, Any]],
    field_rows: list[dict[str, Any]],
    template_type: str,
) -> None:
    template_name = "初始接入模板" if template_type == "initial_work_orders" else "系统外已完成模板"
    sheet.append(["模板填写说明", template_name, "", "", ""])
    sheet.append(["说明类型", "字段编码", "字段名称", "层级/规则", "填写说明"])
    sheet.append(
        [
            "模板用途",
            template_type,
            template_name,
            "字段层级",
            "请按 template 工作表填写数据；本页用于解释字段归属、父字段、条件采集和平台生成规则。",
        ]
    )
    sheet.append(
        [
            "填报范围",
            "",
            "",
            "字段范围",
            "初始接入只填写清单可提供字段；系统外已完成可填写已施工能提供的数据，平台字段按上传时补齐。",
        ]
    )
    replacement_instruction = _template_replacement_hierarchy_instruction(field_rows)
    if replacement_instruction:
        mode, note = replacement_instruction
        sheet.append(["更换层级", "", "", mode, note])

    template_keys = {str(field.get("key") or "") for field in fields}
    for field in field_rows:
        rule = _platform_fill_rule(field, template_type)
        parent_label = _template_parent_label(schema, field)
        condition_hint = _template_condition_hint(schema, field)
        hierarchy_role = _template_hierarchy_role(field)
        if rule:
            sheet.append(
                [
                    "上传时平台生成",
                    field.get("key", ""),
                    field.get("label", ""),
                    "平台补齐",
                    rule,
                ]
            )
            continue
        note_parts: list[str] = []
        if parent_label:
            note_parts.append(f"父字段：{parent_label}")
        if condition_hint:
            note_parts.append(f"条件采集：{condition_hint}")
        if field.get("key") not in template_keys:
            note_parts.append("说明字段，不在 template 填写区")
        if not note_parts:
            note_parts.append("按字段格式填写")
        sheet.append(
            [
                "字段层级",
                field.get("key", ""),
                field.get("label", ""),
                hierarchy_role,
                "；".join(note_parts),
            ]
        )


def _template_replacement_hierarchy_instruction(field_rows: list[dict[str, Any]]) -> tuple[str, str] | None:
    roles = {str(field.get("relation_role") or "").strip() for field in field_rows}
    has_replacement_fields = bool(
        roles
        & {
            "replacement_device",
            "old_device",
            "accessory_replace_confirm",
            "accessory_new_device",
            "evidence_photo",
        }
    )
    if not has_replacement_fields:
        return None
    if "replacement_device" in roles:
        return (
            "主设备更换后确认附属设备",
            "先记录旧主设备和新主设备，再确认通讯模块、SIM卡等附属设备是否更换。",
        )
    return (
        "任务对象下更换附属设备",
        "任务对象保持不变，模块、采集器等附属设备挂在任务对象下，并保留旧件/照片证据。",
    )


def _validation_report(
    project_id: str,
    template_type: str,
    fields: list[dict[str, Any]],
    expected_labels: list[str],
    data_rows: list[tuple[int, tuple[Any, ...]]],
    items: list[dict[str, Any]],
) -> dict[str, Any]:
    error_count = sum(1 for item in items if item["severity"] == "error")
    warning_count = sum(1 for item in items if item["severity"] == "warning")
    return {
        "project_id": project_id,
        "template_type": template_type,
        "status": "failed" if error_count else "warning" if warning_count else "passed",
        "summary": {
            "total_rows": len(data_rows),
            "error_count": error_count,
            "warning_count": warning_count,
        },
        "expected_headers": expected_labels or [field["label"] for field in fields],
        "items": items,
    }


def _required_headers(
    fields: list[dict[str, Any]],
    primary: dict[str, Any],
    aggregate: dict[str, Any],
    template_type: str,
) -> list[dict[str, Any]]:
    required_keys = {primary.get("key"), aggregate.get("key")}
    if template_type == "initial_work_orders":
        required_keys.update(field.get("key") for field in fields if field.get("source") == "import" and field.get("required"))
    return [field for field in fields if field.get("key") in required_keys]


def _recommended_headers(
    fields: list[dict[str, Any]],
    primary: dict[str, Any],
    aggregate: dict[str, Any],
    template_type: str,
) -> list[dict[str, Any]]:
    if template_type != "external_completed":
        return []
    blocked_keys = {primary.get("key"), aggregate.get("key")}
    return [
        field
        for field in fields
        if field.get("required") and field.get("key") not in blocked_keys and not _platform_fill_rule(field, template_type)
    ]


def _validate_conditional_template_fields(
    *,
    items: list[dict[str, Any]],
    fields: list[dict[str, Any]],
    row_number: int,
    row_values: dict[str, Any],
    header_indexes: dict[str, int],
    template_type: str,
) -> None:
    if template_type != "external_completed":
        return
    fields_by_key = {field.get("key"): field for field in fields if field.get("key")}
    for field in fields:
        required_when = field.get("required_when")
        if not isinstance(required_when, dict):
            continue
        trigger_field = fields_by_key.get(str(required_when.get("field_key") or ""))
        if not trigger_field:
            continue
        trigger_label = trigger_field.get("label", "")
        trigger_value = _clean_cell(row_values.get(trigger_label))
        if not _required_when_matches(trigger_value, required_when):
            continue
        field_label = field.get("label", "")
        missing_value = field_label not in header_indexes or not _clean_cell(row_values.get(field_label))
        if not missing_value:
            continue
        items.append(
            _report_item(
                "warning",
                "missing_conditional_field",
                row_number,
                field.get("key", ""),
                field_label,
                f"{trigger_label}={trigger_value} 后建议补齐：{field_label}",
            )
        )


def _required_when_matches(value: Any, required_when: dict[str, Any]) -> bool:
    normalized_value = _clean_cell(value)
    equals = required_when.get("equals")
    if isinstance(equals, list):
        return normalized_value in {_clean_cell(item) for item in equals}
    return normalized_value == _clean_cell(equals)


def _report_item(
    severity: str,
    code: str,
    row: int | None,
    field_key: str,
    field_label: str,
    message: str,
    value: str = "",
) -> dict[str, Any]:
    return {
        "severity": severity,
        "code": code,
        "row": row,
        "field_key": field_key,
        "field_label": field_label,
        "message": message,
        "value": value,
    }


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _value_at(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(not _clean_cell(value) for value in row)


def _is_field_code_row(row: tuple[Any, ...]) -> bool:
    return any(_clean_cell(value).startswith(("字段编码:", "瀛楁")) for value in row)


def _is_valid_number(value: Any) -> bool:
    if isinstance(value, (int, float)):
        return True
    try:
        float(_clean_cell(value))
    except ValueError:
        return False
    return True


def _is_valid_datetime(value: Any) -> bool:
    if isinstance(value, (datetime, date, time)):
        return True
    text = _clean_cell(value)
    if not text:
        return True
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False
