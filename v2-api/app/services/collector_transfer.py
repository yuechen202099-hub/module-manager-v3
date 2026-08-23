from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID
from uuid import uuid4
from xml.etree.ElementTree import ParseError as ElementTreeParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.domain.collector_transfer import (
    CollectorScanDecision,
    CollectorScanDecisionKind,
    MeterSource,
    PoolInsufficientError,
    build_terminal_snapshots,
    decide_collector_scan,
    normalize_identifier,
    plan_random_assignments,
)
from app.models import (
    AuditLog,
    CollectorAssignment,
    CollectorImportRow,
    CollectorMeterItem,
    CollectorPhoto,
    CollectorRequirement,
    CollectorRequirementMeter,
    CollectorScanEvent,
    CollectorTransferRun,
    CollectorTransferTerminal,
    CollectorWorkbenchItem,
    MaterialGroup,
    Photo,
    PhysicalCollector,
    Project,
)
from app.services.photo_storage import resolve_photo_for_response


_WORKBOOK_PARSE_ERRORS: tuple[type[Exception], ...] = (
    BadZipFile,
    InvalidFileException,
    OSError,
    EOFError,
    KeyError,
    ElementTreeParseError,
)
try:
    from lxml.etree import XMLSyntaxError as LxmlXMLSyntaxError
except ImportError:
    pass
else:
    _WORKBOOK_PARSE_ERRORS += (LxmlXMLSyntaxError,)


class CollectorAllocationConflictError(ValueError):
    """A database uniqueness backstop rejected an allocation after the service acquired its locks."""


class CollectorPhotoConflictError(ValueError):
    """Photo content is already bound to a different physical collector."""


_MISSING_TERMINAL_PREFIX = "__missing_terminal__:"


@dataclass(frozen=True, slots=True)
class MeterSourceProjection:
    sources: tuple[MeterSource, ...]
    diagnostics: tuple[dict[str, str], ...]


def _raw_value(raw_data: object, *keys: str) -> str:
    if not isinstance(raw_data, Mapping):
        return ""
    for key in keys:
        value = normalize_identifier(raw_data.get(key))
        if value:
            return value
    return ""


def meter_sources_from_groups(groups: Iterable[object], photos: Iterable[object]) -> MeterSourceProjection:
    photos_by_group: dict[str, list[object]] = defaultdict(list)
    for photo in photos:
        if getattr(photo, "is_active", True) is False:
            continue
        photos_by_group[normalize_identifier(getattr(photo, "group_id", ""))].append(photo)

    sources: list[MeterSource] = []
    diagnostics: list[dict[str, str]] = []
    for group in groups:
        group_id = normalize_identifier(getattr(group, "id", ""))
        terminal_code = normalize_identifier(getattr(group, "terminal", ""))
        if not terminal_code:
            diagnostics.append(
                {"group_id": group_id, "code": "terminal_missing", "message": "终端地址码为空"}
            )
            terminal_code = f"{_MISSING_TERMINAL_PREFIX}{group_id}"
        group_photos = photos_by_group.get(group_id, [])
        module_meter = next(
            (photo for photo in group_photos if normalize_identifier(getattr(photo, "category", "")) == "module_meter"),
            None,
        )
        after_box = next(
            (photo for photo in group_photos if normalize_identifier(getattr(photo, "category", "")) == "after_box"),
            None,
        )
        raw_data = getattr(group, "raw_data", {})
        collector_no = next(
            (
                normalize_identifier(getattr(photo, "collector", ""))
                for photo in group_photos
                if normalize_identifier(getattr(photo, "collector", ""))
            ),
            "",
        ) or _raw_value(raw_data, "collector", "采集器", "采集器号", "construction_collector")
        module_no = (
            normalize_identifier(getattr(module_meter, "asset_no", ""))
            if module_meter is not None
            else ""
        ) or next(
            (
                normalize_identifier(getattr(photo, "asset_no", ""))
                for photo in group_photos
                if normalize_identifier(getattr(photo, "asset_no", ""))
            ),
            "",
        ) or _raw_value(raw_data, "module_asset_no", "模块资产编号", "模块号", "construction_module_asset_no")
        meter_no = normalize_identifier(getattr(group, "display_meter_no", ""))
        source_diagnostics: list[str] = []
        if not meter_no:
            source_diagnostics.append("meter_missing")
        if not collector_no:
            source_diagnostics.append("collector_missing")
        if module_meter is None:
            source_diagnostics.append("module_meter_photo_missing")
        if after_box is None:
            source_diagnostics.append("after_box_photo_missing")
        for code in source_diagnostics:
            diagnostics.append({"group_id": group_id, "code": code, "message": code})
        sources.append(
            MeterSource(
                group_id=group_id,
                terminal_code=terminal_code,
                installation_address=normalize_identifier(getattr(group, "installation_address", "")),
                meter_no=meter_no,
                collector_no=collector_no,
                module_no=module_no,
                module_meter_photo_id=(
                    normalize_identifier(getattr(module_meter, "id", "")) if module_meter is not None else None
                ),
                after_box_photo_id=(
                    normalize_identifier(getattr(after_box, "id", "")) if after_box is not None else None
                ),
            )
        )
    return MeterSourceProjection(sources=tuple(sources), diagnostics=tuple(diagnostics))


def read_collector_numbers_from_workbook(content: bytes) -> tuple[tuple[int, str], ...]:
    if not content:
        raise ValueError("Excel 文件为空")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [normalize_identifier(value) for value in next(rows, ())]
        header_indexes = {header: index for index, header in enumerate(headers) if header}
        preferred_indexes = [
            header_indexes[header]
            for header in ("采集器", "采集器号", "扫码内容")
            if header in header_indexes
        ]
        if not preferred_indexes:
            raise ValueError("Excel 缺少采集器、采集器号或扫码内容列")
        result: list[tuple[int, str]] = []
        for row_number, row in enumerate(rows, start=2):
            collector_no = next(
                (
                    normalize_identifier(row[index])
                    for index in preferred_indexes
                    if index < len(row) and normalize_identifier(row[index])
                ),
                "",
            )
            result.append((row_number, collector_no))
    except _WORKBOOK_PARSE_ERRORS as exc:
        raise ValueError("Excel 文件无法解析") from exc
    return tuple(result)


def collector_no_from_photo_filename(filename: str) -> str:
    candidate = normalize_identifier(filename)
    if not candidate or "/" in candidate or "\\" in candidate:
        return ""
    return normalize_identifier(Path(candidate).stem)


def _uuid(value: object, field_name: str) -> UUID:
    try:
        return UUID(normalize_identifier(value))
    except (ValueError, TypeError, AttributeError) as exc:
        raise ValueError(f"{field_name} is invalid") from exc


def _photo_response(photo: Photo | CollectorPhoto | None) -> dict[str, object] | None:
    if photo is None:
        return None
    payload = {
        "id": str(photo.id),
        "image_url": normalize_identifier(getattr(photo, "image_url", "")),
        "object_key": normalize_identifier(getattr(photo, "object_key", "")),
        "storage_type": normalize_identifier(getattr(photo, "storage_type", "")),
        "storage_key": normalize_identifier(getattr(photo, "storage_key", "")),
        "storage_bucket": normalize_identifier(getattr(photo, "storage_bucket", "")),
        "sha256": normalize_identifier(getattr(photo, "sha256", "")),
        "content_type": normalize_identifier(getattr(photo, "content_type", "")),
    }
    return resolve_photo_for_response(payload)


class PostgresCollectorTransferService:
    """Transactional collector-transfer use cases scoped to one authenticated team."""

    def __init__(self, *, session: Session, team_id: str, actor: str) -> None:
        self.session = session
        self.team_id = normalize_identifier(team_id)
        self.actor = normalize_identifier(actor) or "admin"
        if not self.team_id:
            raise ValueError("team_id is required")

    def _run(self, run_id: str, *, lock: bool = False) -> CollectorTransferRun:
        query = select(CollectorTransferRun).where(
            CollectorTransferRun.id == _uuid(run_id, "run_id"),
            CollectorTransferRun.team_id == self.team_id,
        )
        if lock:
            query = query.with_for_update()
        record = self.session.scalar(query)
        if record is None:
            raise KeyError(run_id)
        return record

    def _audit(
        self,
        *,
        action: str,
        entity_type: str,
        entity_id: UUID | None,
        project_id: UUID | None = None,
        payload: dict[str, object] | None = None,
    ) -> None:
        self.session.add(
            AuditLog(
                team_id=self.team_id,
                actor_username=self.actor,
                project_id=project_id,
                action=action,
                entity_type=entity_type,
                entity_id=entity_id,
                payload=payload or {},
            )
        )

    def _run_summary(self, run: CollectorTransferRun) -> dict[str, object]:
        return {
            "id": str(run.id),
            "project_id": str(run.project_id),
            "name": run.name,
            "status": run.status,
            **dict(run.stats or {}),
            "diagnostics": list(run.diagnostics or []),
            "created_at": run.created_at.isoformat() if run.created_at else None,
        }

    def create_run(self, *, project_id: str, name: str) -> dict[str, object]:
        project_uuid = _uuid(project_id, "project_id")
        project = self.session.scalar(
            select(Project).where(Project.id == project_uuid, Project.team_id == self.team_id)
        )
        if project is None:
            raise KeyError(project_id)

        groups = list(
            self.session.scalars(
                select(MaterialGroup)
                .where(MaterialGroup.project_id == project_uuid, MaterialGroup.team_id == self.team_id)
                .order_by(MaterialGroup.terminal, MaterialGroup.display_meter_no, MaterialGroup.id)
            ).all()
        )
        group_ids = [group.id for group in groups]
        photos = (
            list(
                self.session.scalars(
                    select(Photo)
                    .where(
                        Photo.team_id == self.team_id,
                        Photo.group_id.in_(group_ids),
                        Photo.is_active.is_(True),
                    )
                    .order_by(Photo.group_id, Photo.sort_order, Photo.id)
                ).all()
            )
            if group_ids
            else []
        )
        projection = meter_sources_from_groups(groups, photos)
        snapshots = build_terminal_snapshots(projection.sources)
        run = CollectorTransferRun(
            team_id=self.team_id,
            project_id=project_uuid,
            name=normalize_identifier(name) or "采集器盘点",
            status="inventory",
            diagnostics=list(projection.diagnostics),
            stats={},
        )
        self.session.add(run)
        self.session.flush()

        source_by_group = {source.group_id: source for source in projection.sources}
        diagnostics_by_group: dict[str, list[dict[str, str]]] = defaultdict(list)
        for diagnostic in projection.diagnostics:
            diagnostics_by_group[diagnostic["group_id"]].append(diagnostic)

        requirement_count = 0
        blocked_terminal_count = 0
        for terminal_index, snapshot in enumerate(snapshots):
            terminal_diagnostics = [
                diagnostic
                for meter_source in snapshot.meters
                for diagnostic in diagnostics_by_group.get(meter_source.group_id, [])
            ]
            terminal_status = "blocked" if terminal_diagnostics else "ready"
            if terminal_status == "blocked":
                blocked_terminal_count += 1
            terminal = CollectorTransferTerminal(
                run_id=run.id,
                team_id=self.team_id,
                terminal_code=snapshot.terminal_code,
                installation_address=snapshot.installation_address,
                status=terminal_status,
                meter_count=len(snapshot.meters),
                collector_requirement_count=len(snapshot.collector_requirements),
                diagnostics=terminal_diagnostics,
            )
            self.session.add(terminal)
            self.session.flush()

            meters_by_group: dict[str, CollectorMeterItem] = {}
            for meter_index, source in enumerate(snapshot.meters):
                source = source_by_group[source.group_id]
                meter_item = CollectorMeterItem(
                    run_id=run.id,
                    terminal_id=terminal.id,
                    team_id=self.team_id,
                    source_group_id=_uuid(source.group_id, "source_group_id"),
                    meter_no=source.meter_no,
                    meter_barcode=source.meter_no,
                    module_no=source.module_no,
                    module_barcode=source.module_no,
                    original_collector_no=source.collector_no,
                    module_meter_photo_id=(
                        _uuid(source.module_meter_photo_id, "module_meter_photo_id")
                        if source.module_meter_photo_id
                        else None
                    ),
                    after_box_photo_id=(
                        _uuid(source.after_box_photo_id, "after_box_photo_id")
                        if source.after_box_photo_id
                        else None
                    ),
                    sort_order=meter_index,
                    diagnostics=diagnostics_by_group.get(source.group_id, []),
                )
                self.session.add(meter_item)
                self.session.flush()
                meters_by_group[source.group_id] = meter_item
                self.session.add(
                    CollectorWorkbenchItem(
                        run_id=run.id,
                        terminal_id=terminal.id,
                        team_id=self.team_id,
                        item_kind="meter_install",
                        source_key=str(meter_item.id),
                        meter_item_id=meter_item.id,
                        status="pending",
                        sort_order=meter_index,
                    )
                )

            for requirement_index, snapshot_requirement in enumerate(snapshot.collector_requirements):
                requirement = CollectorRequirement(
                    run_id=run.id,
                    terminal_id=terminal.id,
                    team_id=self.team_id,
                    original_collector_no=snapshot_requirement.original_collector_no,
                    status="unmatched" if terminal_status != "blocked" else "blocked",
                    sort_order=requirement_index,
                    diagnostics=terminal_diagnostics,
                )
                self.session.add(requirement)
                self.session.flush()
                requirement_count += 1
                for group_id in snapshot_requirement.meter_group_ids:
                    self.session.add(
                        CollectorRequirementMeter(
                            requirement_id=requirement.id,
                            meter_item_id=meters_by_group[group_id].id,
                        )
                    )

        run.stats = {
            "terminal_count": len(snapshots),
            "meter_count": len(projection.sources),
            "collector_requirement_count": requirement_count,
            "blocked_terminal_count": blocked_terminal_count,
            "direct_match_count": 0,
            "pool_available_count": 0,
            "assignment_count": 0,
        }
        self._audit(
            action="collector_transfer.run_created",
            entity_type="collector_transfer_run",
            entity_id=run.id,
            project_id=project_uuid,
            payload=dict(run.stats),
        )
        self.session.commit()
        self.session.refresh(run)
        return self._run_summary(run)

    def list_runs(self, *, project_id: str | None = None) -> list[dict[str, object]]:
        query = select(CollectorTransferRun).where(CollectorTransferRun.team_id == self.team_id)
        if project_id:
            query = query.where(CollectorTransferRun.project_id == _uuid(project_id, "project_id"))
        rows = self.session.scalars(query.order_by(CollectorTransferRun.created_at.desc())).all()
        return [self._run_summary(row) for row in rows]

    def _active_collector_photo(self, physical_collector_id: UUID) -> CollectorPhoto | None:
        return self.session.scalar(
            select(CollectorPhoto)
            .where(
                CollectorPhoto.team_id == self.team_id,
                CollectorPhoto.physical_collector_id == physical_collector_id,
                CollectorPhoto.is_active.is_(True),
            )
            .order_by(CollectorPhoto.created_at.desc(), CollectorPhoto.id.desc())
        )

    def _locked_physical_collector_for_scan(
        self,
        *,
        run: CollectorTransferRun,
        collector_no: str,
    ) -> PhysicalCollector:
        query = (
            select(PhysicalCollector)
            .where(
                PhysicalCollector.team_id == self.team_id,
                PhysicalCollector.collector_no == collector_no,
            )
            .with_for_update()
        )
        physical = self.session.scalar(query)
        if physical is not None:
            return physical

        candidate = PhysicalCollector(
            team_id=self.team_id,
            collector_no=collector_no,
            pool_status="awaiting_photo",
            first_seen_run_id=run.id,
        )
        try:
            with self.session.begin_nested():
                self.session.add(candidate)
                self.session.flush()
        except IntegrityError:
            physical = self.session.scalar(query)
            if physical is None:
                raise
            return physical
        return candidate

    def _create_assignment(
        self,
        *,
        run: CollectorTransferRun,
        requirement: CollectorRequirement,
        physical: PhysicalCollector,
        photo: CollectorPhoto,
        assignment_mode: str,
    ) -> tuple[CollectorAssignment, bool]:
        candidate = CollectorAssignment(
            run_id=run.id,
            team_id=self.team_id,
            requirement_id=requirement.id,
            physical_collector_id=physical.id,
            collector_photo_id=photo.id,
            assignment_mode=assignment_mode,
            status="reserved",
        )
        try:
            with self.session.begin_nested():
                self.session.add(candidate)
                self.session.flush()
        except IntegrityError as exc:
            winner = self.session.scalar(
                select(CollectorAssignment)
                .where(
                    or_(
                        CollectorAssignment.requirement_id == requirement.id,
                        CollectorAssignment.physical_collector_id == physical.id,
                    )
                )
                .with_for_update()
            )
            if (
                winner is not None
                and winner.run_id == run.id
                and winner.requirement_id == requirement.id
                and winner.physical_collector_id == physical.id
                and winner.collector_photo_id == photo.id
                and winner.assignment_mode == assignment_mode
                and winner.status in {"reserved", "used"}
            ):
                return winner, False
            self.session.rollback()
            raise CollectorAllocationConflictError("collector assignment conflicts with an existing allocation") from exc
        return candidate, True

    def _ensure_removal_workbench_item(
        self,
        *,
        run: CollectorTransferRun,
        requirement: CollectorRequirement,
        assignment: CollectorAssignment,
    ) -> None:
        existing = self.session.scalar(
            select(CollectorWorkbenchItem).where(
                CollectorWorkbenchItem.run_id == run.id,
                CollectorWorkbenchItem.item_kind == "collector_removal",
                CollectorWorkbenchItem.source_key == str(requirement.id),
            )
        )
        if existing is None:
            meter_count = int(
                self.session.scalar(
                    select(func.count(CollectorMeterItem.id)).where(
                        CollectorMeterItem.terminal_id == requirement.terminal_id
                    )
                )
                or 0
            )
            self.session.add(
                CollectorWorkbenchItem(
                    run_id=run.id,
                    terminal_id=requirement.terminal_id,
                    team_id=self.team_id,
                    item_kind="collector_removal",
                    source_key=str(requirement.id),
                    requirement_id=requirement.id,
                    assignment_id=assignment.id,
                    status="pending",
                    sort_order=meter_count + requirement.sort_order,
                )
            )
        else:
            existing.assignment_id = assignment.id

    def scan_collector(self, *, run_id: str, collector_no: str) -> dict[str, object]:
        run = self._run(run_id, lock=True)
        normalized_no = normalize_identifier(collector_no)
        if not normalized_no:
            raise ValueError("collector_no is required")
        physical = self._locked_physical_collector_for_scan(run=run, collector_no=normalized_no)
        physical.last_scanned_at = datetime.now(UTC)

        existing_assignment = self.session.scalar(
            select(CollectorAssignment).where(
                CollectorAssignment.physical_collector_id == physical.id,
                CollectorAssignment.status.in_(("reserved", "used")),
            )
        )
        if existing_assignment is not None and existing_assignment.run_id != run.id:
            raise ValueError("该采集器已经在其他批次使用，禁止重复使用")
        photo = self._active_collector_photo(physical.id)

        requirement: CollectorRequirement | None = None
        if existing_assignment is not None:
            requirement = self.session.scalar(
                select(CollectorRequirement).where(CollectorRequirement.id == existing_assignment.requirement_id)
            )
        if requirement is None:
            prior_direct = self.session.scalar(
                select(CollectorScanEvent)
                .where(
                    CollectorScanEvent.run_id == run.id,
                    CollectorScanEvent.physical_collector_id == physical.id,
                    CollectorScanEvent.requirement_id.is_not(None),
                )
                .order_by(CollectorScanEvent.created_at.desc(), CollectorScanEvent.id.desc())
            )
            if prior_direct is not None:
                requirement = self.session.scalar(
                    select(CollectorRequirement).where(CollectorRequirement.id == prior_direct.requirement_id)
                )
        if requirement is None:
            requirement = self.session.scalar(
                select(CollectorRequirement)
                .where(
                    CollectorRequirement.run_id == run.id,
                    CollectorRequirement.team_id == self.team_id,
                    CollectorRequirement.original_collector_no == normalized_no,
                    CollectorRequirement.status == "unmatched",
                )
                .order_by(CollectorRequirement.terminal_id, CollectorRequirement.sort_order, CollectorRequirement.id)
                .with_for_update()
            )

        if existing_assignment is not None:
            decision = CollectorScanDecision(
                kind=CollectorScanDecisionKind.ASSIGNMENT_REUSE,
                requirement_id=str(existing_assignment.requirement_id),
                requires_photo=False,
                add_to_pool=False,
            )
        else:
            decision = decide_collector_scan(
                collector_no=normalized_no,
                unmatched_requirements=({str(requirement.id): requirement.original_collector_no} if requirement else {}),
                has_reusable_photo=photo is not None,
            )
        if requirement is not None and existing_assignment is None:
            physical.pool_status = "direct"
            if photo is None:
                requirement.status = "direct_pending_photo"
            else:
                requirement.status = "direct_ready"
                existing_assignment, _created = self._create_assignment(
                    run=run,
                    requirement=requirement,
                    physical=physical,
                    photo=photo,
                    assignment_mode="direct",
                )
                self._ensure_removal_workbench_item(
                    run=run,
                    requirement=requirement,
                    assignment=existing_assignment,
                )
        if requirement is None:
            physical.pool_status = "available" if photo is not None else "awaiting_photo"
            if photo is not None:
                decision = CollectorScanDecision(
                    kind=CollectorScanDecisionKind.POOL_NEEDS_PHOTO,
                    requirement_id=None,
                    requires_photo=False,
                    add_to_pool=True,
                )

        event = CollectorScanEvent(
            run_id=run.id,
            team_id=self.team_id,
            physical_collector_id=physical.id,
            requirement_id=requirement.id if requirement is not None else None,
            scanned_value=normalized_no,
            decision=decision.kind.value,
            requires_photo=decision.requires_photo,
            add_to_pool=decision.add_to_pool,
        )
        self.session.add(event)
        self._audit(
            action="collector_transfer.collector_scanned",
            entity_type="physical_collector",
            entity_id=physical.id,
            project_id=run.project_id,
            payload={
                "run_id": str(run.id),
                "collector_no": normalized_no,
                "decision": decision.kind.value,
                "requirement_id": str(requirement.id) if requirement else None,
            },
        )
        self.session.commit()
        return {
            "collector_id": str(physical.id),
            "collector_no": normalized_no,
            "decision": decision.kind.value,
            "requires_photo": decision.requires_photo,
            "add_to_pool": decision.add_to_pool,
            "pool_status": physical.pool_status,
            "requirement_id": str(requirement.id) if requirement else None,
            "photo": _photo_response(photo),
        }

    def register_photo(
        self,
        *,
        run_id: str,
        collector_id: str,
        original_filename: str,
        stored: Mapping[str, object],
        byte_size: int,
    ) -> dict[str, object]:
        run = self._run(run_id, lock=True)
        physical = self.session.scalar(
            select(PhysicalCollector)
            .where(
                PhysicalCollector.id == _uuid(collector_id, "collector_id"),
                PhysicalCollector.team_id == self.team_id,
            )
            .with_for_update()
        )
        if physical is None:
            raise KeyError(collector_id)
        sha256 = normalize_identifier(stored.get("sha256"))
        photo = self.session.scalar(
            select(CollectorPhoto).where(
                CollectorPhoto.sha256 == sha256,
            )
        )
        if photo is not None and photo.physical_collector_id != physical.id:
            raise CollectorPhotoConflictError(
                "photo content is already bound to another physical collector"
            )
        if photo is None:
            photo = CollectorPhoto(
                team_id=self.team_id,
                physical_collector_id=physical.id,
                sha256=sha256,
                original_filename=normalize_identifier(original_filename) or f"{physical.collector_no}.jpg",
                object_key=normalize_identifier(stored.get("storage_key") or stored.get("url")),
                image_url=normalize_identifier(stored.get("url")) or None,
                storage_type=normalize_identifier(stored.get("storage_type")) or "local_upload",
                content_type=normalize_identifier(stored.get("content_type")) or None,
                byte_size=byte_size,
                captured_at=datetime.now(UTC),
                is_active=True,
            )
            self.session.add(photo)
            try:
                self.session.flush()
            except IntegrityError as exc:
                self.session.rollback()
                raise CollectorPhotoConflictError(
                    "photo content is already bound to another physical collector"
                ) from exc

        direct_event = self.session.scalar(
            select(CollectorScanEvent)
            .where(
                CollectorScanEvent.run_id == run.id,
                CollectorScanEvent.physical_collector_id == physical.id,
                CollectorScanEvent.requirement_id.is_not(None),
            )
            .order_by(CollectorScanEvent.created_at.desc(), CollectorScanEvent.id.desc())
        )
        assignment: CollectorAssignment | None = None
        if direct_event is not None:
            requirement = self.session.scalar(
                select(CollectorRequirement)
                .where(
                    CollectorRequirement.id == direct_event.requirement_id,
                    CollectorRequirement.run_id == run.id,
                )
                .with_for_update()
            )
            if requirement is not None:
                assignment = self.session.scalar(
                    select(CollectorAssignment).where(CollectorAssignment.requirement_id == requirement.id)
                )
                if assignment is None:
                    assignment, _created = self._create_assignment(
                        run=run,
                        requirement=requirement,
                        physical=physical,
                        photo=photo,
                        assignment_mode="direct",
                    )
                    requirement.status = "direct_ready"
                    physical.pool_status = "direct"
                    self._ensure_removal_workbench_item(
                        run=run,
                        requirement=requirement,
                        assignment=assignment,
                    )
        if direct_event is None and physical.pool_status == "awaiting_photo":
            physical.pool_status = "available"
        self._audit(
            action="collector_transfer.photo_registered",
            entity_type="collector_photo",
            entity_id=photo.id,
            project_id=run.project_id,
            payload={"run_id": str(run.id), "collector_id": str(physical.id), "sha256": photo.sha256},
        )
        self.session.commit()
        return {
            "collector_id": str(physical.id),
            "collector_no": physical.collector_no,
            "pool_status": physical.pool_status,
            "photo": _photo_response(photo),
            "assignment_id": str(assignment.id) if assignment is not None else None,
        }

    def allocate(self, *, run_id: str) -> dict[str, object]:
        run = self._run(run_id, lock=True)
        requirements = list(
            self.session.scalars(
                select(CollectorRequirement)
                .where(
                    CollectorRequirement.run_id == run.id,
                    CollectorRequirement.team_id == self.team_id,
                    CollectorRequirement.status == "unmatched",
                )
                .order_by(CollectorRequirement.terminal_id, CollectorRequirement.sort_order, CollectorRequirement.id)
                .with_for_update()
            ).all()
        )
        if not requirements and run.status == "allocated":
            return {
                "run_id": str(run.id),
                "assignment_count": int(dict(run.stats or {}).get("assignment_count") or 0),
                "assignments": [],
            }
        physical_collectors = list(
            self.session.scalars(
                select(PhysicalCollector)
                .where(
                    PhysicalCollector.team_id == self.team_id,
                    PhysicalCollector.pool_status == "available",
                )
                .order_by(PhysicalCollector.id)
                .with_for_update()
            ).all()
        )
        photos_by_collector: dict[str, CollectorPhoto] = {}
        if physical_collectors:
            photos = self.session.scalars(
                select(CollectorPhoto)
                .where(
                    CollectorPhoto.team_id == self.team_id,
                    CollectorPhoto.physical_collector_id.in_([item.id for item in physical_collectors]),
                    CollectorPhoto.is_active.is_(True),
                )
                .order_by(CollectorPhoto.physical_collector_id, CollectorPhoto.created_at.desc(), CollectorPhoto.id.desc())
                .with_for_update()
            ).all()
            for photo in photos:
                photos_by_collector.setdefault(str(photo.physical_collector_id), photo)

        plan = plan_random_assignments(
            requirement_ids=[str(item.id) for item in requirements],
            available_collector_ids=[
                str(item.id) for item in physical_collectors if str(item.id) in photos_by_collector
            ],
        )
        requirements_by_id = {str(item.id): item for item in requirements}
        collectors_by_id = {str(item.id): item for item in physical_collectors}
        result: list[dict[str, str]] = []
        for requirement_id, collector_id in plan:
            requirement = requirements_by_id[requirement_id]
            physical = collectors_by_id[collector_id]
            photo = photos_by_collector[collector_id]
            assignment, created = self._create_assignment(
                run=run,
                requirement=requirement,
                physical=physical,
                photo=photo,
                assignment_mode="random",
            )
            if created or requirement.status == "unmatched":
                requirement.status = "assigned"
            if created or physical.pool_status == "available":
                physical.pool_status = "reserved"
            self._ensure_removal_workbench_item(
                run=run,
                requirement=requirement,
                assignment=assignment,
            )
            result.append(
                {
                    "assignment_id": str(assignment.id),
                    "requirement_id": requirement_id,
                    "original_collector_no": requirement.original_collector_no,
                    "physical_collector_id": collector_id,
                    "final_collector_no": physical.collector_no,
                    "mode": "random",
                }
            )
        if not requirements:
            result = []
        run.status = "allocated"
        stats = dict(run.stats or {})
        stats["assignment_count"] = int(
            self.session.scalar(
                select(func.count(CollectorAssignment.id)).where(CollectorAssignment.run_id == run.id)
            )
            or len(result)
        )
        stats["pool_available_count"] = max(0, len(physical_collectors) - len(result))
        run.stats = stats
        self._audit(
            action="collector_transfer.random_allocated",
            entity_type="collector_transfer_run",
            entity_id=run.id,
            project_id=run.project_id,
            payload={"assignment_count": len(result)},
        )
        self.session.commit()
        return {"run_id": str(run.id), "assignment_count": len(result), "assignments": result}

    def list_workbench(self, *, run_id: str) -> dict[str, object]:
        run = self._run(run_id)
        terminals = self.session.scalars(
            select(CollectorTransferTerminal)
            .where(CollectorTransferTerminal.run_id == run.id, CollectorTransferTerminal.team_id == self.team_id)
            .order_by(CollectorTransferTerminal.terminal_code, CollectorTransferTerminal.id)
        ).all()
        items = self.session.scalars(
            select(CollectorWorkbenchItem).where(
                CollectorWorkbenchItem.run_id == run.id,
                CollectorWorkbenchItem.team_id == self.team_id,
            )
        ).all()
        counts: dict[str, tuple[int, int]] = {}
        for terminal in terminals:
            terminal_items = [item for item in items if item.terminal_id == terminal.id]
            counts[str(terminal.id)] = (
                sum(item.status == "completed" for item in terminal_items),
                len(terminal_items),
            )
        return {
            "run": self._run_summary(run),
            "terminals": [
                {
                    "id": str(terminal.id),
                    "terminal_code": terminal.terminal_code,
                    "installation_address": terminal.installation_address,
                    "status": terminal.status,
                    "meter_count": terminal.meter_count,
                    "collector_requirement_count": terminal.collector_requirement_count,
                    "completed_count": counts[str(terminal.id)][0],
                    "total_count": counts[str(terminal.id)][1],
                    "progress": (
                        round(counts[str(terminal.id)][0] * 100 / counts[str(terminal.id)][1])
                        if counts[str(terminal.id)][1]
                        else 0
                    ),
                    "diagnostics": list(terminal.diagnostics or []),
                }
                for terminal in terminals
            ],
        }

    def terminal_workbench(self, *, run_id: str, terminal_id: str) -> dict[str, object]:
        run = self._run(run_id)
        terminal = self.session.scalar(
            select(CollectorTransferTerminal).where(
                CollectorTransferTerminal.id == _uuid(terminal_id, "terminal_id"),
                CollectorTransferTerminal.run_id == run.id,
                CollectorTransferTerminal.team_id == self.team_id,
            )
        )
        if terminal is None:
            raise KeyError(terminal_id)
        items = self.session.scalars(
            select(CollectorWorkbenchItem)
            .where(
                CollectorWorkbenchItem.terminal_id == terminal.id,
                CollectorWorkbenchItem.team_id == self.team_id,
            )
            .order_by(CollectorWorkbenchItem.sort_order, CollectorWorkbenchItem.id)
        ).all()
        meter_ids = [item.meter_item_id for item in items if item.meter_item_id]
        meters = (
            self.session.scalars(select(CollectorMeterItem).where(CollectorMeterItem.id.in_(meter_ids))).all()
            if meter_ids
            else []
        )
        meters_by_id = {item.id: item for item in meters}
        legacy_photo_ids = [
            photo_id
            for meter_item in meters
            for photo_id in (meter_item.module_meter_photo_id, meter_item.after_box_photo_id)
            if photo_id
        ]
        legacy_photos = (
            self.session.scalars(select(Photo).where(Photo.id.in_(legacy_photo_ids), Photo.is_active.is_(True))).all()
            if legacy_photo_ids
            else []
        )
        legacy_photos_by_id = {item.id: item for item in legacy_photos}
        assignment_ids = [item.assignment_id for item in items if item.assignment_id]
        assignments = (
            self.session.scalars(select(CollectorAssignment).where(CollectorAssignment.id.in_(assignment_ids))).all()
            if assignment_ids
            else []
        )
        assignments_by_id = {item.id: item for item in assignments}
        collector_ids = [item.physical_collector_id for item in assignments]
        collector_photo_ids = [item.collector_photo_id for item in assignments]
        physicals = (
            self.session.scalars(select(PhysicalCollector).where(PhysicalCollector.id.in_(collector_ids))).all()
            if collector_ids
            else []
        )
        collector_photos = (
            self.session.scalars(select(CollectorPhoto).where(CollectorPhoto.id.in_(collector_photo_ids))).all()
            if collector_photo_ids
            else []
        )
        physical_by_id = {item.id: item for item in physicals}
        collector_photo_by_id = {item.id: item for item in collector_photos}

        payload_items: list[dict[str, object]] = []
        for item in items:
            if item.item_kind == "meter_install" and item.meter_item_id in meters_by_id:
                meter_item = meters_by_id[item.meter_item_id]
                payload_items.append(
                    {
                        "id": str(item.id),
                        "kind": item.item_kind,
                        "status": item.status,
                        "meter_no": meter_item.meter_no,
                        "meter_barcode": meter_item.meter_barcode,
                        "module_no": meter_item.module_no,
                        "module_barcode": meter_item.module_barcode,
                        "photos": [
                            {
                                "slot": "module_meter",
                                "label": "电表和模块",
                                "photo": _photo_response(legacy_photos_by_id.get(meter_item.module_meter_photo_id)),
                            },
                            {
                                "slot": "after_box",
                                "label": "改造完成",
                                "photo": _photo_response(legacy_photos_by_id.get(meter_item.after_box_photo_id)),
                            },
                        ],
                    }
                )
            elif item.item_kind == "collector_removal" and item.assignment_id in assignments_by_id:
                assignment = assignments_by_id[item.assignment_id]
                physical = physical_by_id.get(assignment.physical_collector_id)
                payload_items.append(
                    {
                        "id": str(item.id),
                        "kind": item.item_kind,
                        "status": item.status,
                        "collector_no": physical.collector_no if physical else "",
                        "collector_barcode": physical.collector_no if physical else "",
                        "assignment_mode": assignment.assignment_mode,
                        "photos": [
                            {
                                "slot": "collector",
                                "label": "拆旧采集器",
                                "photo": _photo_response(collector_photo_by_id.get(assignment.collector_photo_id)),
                            }
                        ],
                    }
                )
        return {
            "run_id": str(run.id),
            "terminal": {
                "id": str(terminal.id),
                "terminal_code": terminal.terminal_code,
                "installation_address": terminal.installation_address,
                "status": terminal.status,
            },
            "items": payload_items,
        }

    def set_workbench_item_status(self, *, item_id: str, completed: bool) -> dict[str, object]:
        item = self.session.scalar(
            select(CollectorWorkbenchItem)
            .where(
                CollectorWorkbenchItem.id == _uuid(item_id, "item_id"),
                CollectorWorkbenchItem.team_id == self.team_id,
            )
            .with_for_update()
        )
        if item is None:
            raise KeyError(item_id)
        item.status = "completed" if completed else "pending"
        item.completed_at = datetime.now(UTC) if completed else None
        if item.assignment_id:
            assignment = self.session.scalar(
                select(CollectorAssignment)
                .where(CollectorAssignment.id == item.assignment_id)
                .with_for_update()
            )
            if assignment is not None:
                assignment.status = "used" if completed else "reserved"
                assignment.used_at = datetime.now(UTC) if completed else None
                physical = self.session.scalar(
                    select(PhysicalCollector)
                    .where(PhysicalCollector.id == assignment.physical_collector_id)
                    .with_for_update()
                )
                requirement = self.session.scalar(
                    select(CollectorRequirement)
                    .where(CollectorRequirement.id == assignment.requirement_id)
                    .with_for_update()
                )
                if physical is not None:
                    physical.pool_status = "used" if completed else (
                        "direct" if assignment.assignment_mode == "direct" else "reserved"
                    )
                if requirement is not None:
                    requirement.status = "used" if completed else (
                        "direct_ready" if assignment.assignment_mode == "direct" else "assigned"
                    )
        total = int(
            self.session.scalar(
                select(func.count(CollectorWorkbenchItem.id)).where(
                    CollectorWorkbenchItem.terminal_id == item.terminal_id
                )
            )
            or 0
        )
        completed_count = int(
            self.session.scalar(
                select(func.count(CollectorWorkbenchItem.id)).where(
                    CollectorWorkbenchItem.terminal_id == item.terminal_id,
                    CollectorWorkbenchItem.status == "completed",
                )
            )
            or 0
        )
        terminal = self.session.scalar(
            select(CollectorTransferTerminal)
            .where(CollectorTransferTerminal.id == item.terminal_id)
            .with_for_update()
        )
        if terminal is not None:
            terminal.completed_item_count = completed_count
            terminal.status = "completed" if total and completed_count >= total else "in_progress"
        self._audit(
            action="collector_transfer.workbench_item_completed" if completed else "collector_transfer.workbench_item_reopened",
            entity_type="collector_workbench_item",
            entity_id=item.id,
            payload={"completed": completed},
        )
        self.session.commit()
        return {"id": str(item.id), "status": item.status, "completed_at": item.completed_at.isoformat() if item.completed_at else None}

    def record_import_row(
        self,
        *,
        run_id: str,
        batch_id: str,
        row_number: int,
        collector_no: str,
        outcome: str,
        message: str = "",
        payload: dict[str, object] | None = None,
    ) -> None:
        run = self._run(run_id)
        self.session.add(
            CollectorImportRow(
                run_id=run.id,
                team_id=self.team_id,
                batch_id=batch_id,
                row_number=row_number,
                collector_no=normalize_identifier(collector_no),
                outcome=outcome,
                message=message,
                payload=payload or {},
            )
        )

    def import_inventory(
        self,
        *,
        run_id: str,
        rows: Sequence[tuple[int, str]],
        photos_by_collector: Mapping[str, Mapping[str, object]],
        photo_inputs: Sequence[Mapping[str, object]] = (),
    ) -> dict[str, object]:
        self._run(run_id)
        batch_id = uuid4().hex
        outcomes: list[dict[str, object]] = []
        counters = {"inserted": 0, "reused": 0, "needs_photo": 0, "invalid": 0}
        seen: set[str] = set()

        def process_collector(
            collector_no: str,
            photo_payload: Mapping[str, object] | None,
        ) -> tuple[str, str, dict[str, object], bool]:
            try:
                result = self.scan_collector(run_id=run_id, collector_no=collector_no)
                photo_referenced = False
                if photo_payload and result.get("requires_photo"):
                    result = {
                        **result,
                        **self.register_photo(
                            run_id=run_id,
                            collector_id=str(result["collector_id"]),
                            original_filename=str(
                                photo_payload.get("original_filename") or f"{collector_no}.jpg"
                            ),
                            stored=(
                                photo_payload.get("stored")
                                if isinstance(photo_payload.get("stored"), Mapping)
                                else {}
                            ),
                            byte_size=int(photo_payload.get("byte_size") or 0),
                        ),
                    }
                    stored = photo_payload.get("stored")
                    response_photo = result.get("photo")
                    if isinstance(stored, Mapping) and isinstance(response_photo, Mapping):
                        stored_key = normalize_identifier(stored.get("storage_key") or stored.get("url"))
                        response_key = normalize_identifier(
                            response_photo.get("object_key")
                            or response_photo.get("storage_key")
                            or response_photo.get("image_url")
                        )
                        photo_referenced = bool(stored_key and stored_key == response_key)
                if result.get("requires_photo") and not photo_payload:
                    return "needs_photo", "已登记，仍需手机补拍", result, False
                if result.get("decision") == CollectorScanDecisionKind.DIRECT_REUSE.value:
                    return "reused", "同号照片可直接复用", result, photo_referenced
                return "inserted", "已按扫码规则登记", result, photo_referenced
            except (KeyError, ValueError) as exc:
                self.session.rollback()
                return "invalid", str(exc), {}, False

        def persist_diagnostic(
            *,
            row_number: int,
            input_kind: str,
            collector_no: str,
            outcome: str,
            message: str,
            result: Mapping[str, object] | None = None,
            original_filename: str = "",
        ) -> None:
            counters[outcome] += 1
            diagnostic_payload = {**dict(result or {}), "input_kind": input_kind}
            if original_filename:
                diagnostic_payload["original_filename"] = original_filename
            self.record_import_row(
                run_id=run_id,
                batch_id=batch_id,
                row_number=int(row_number),
                collector_no=collector_no,
                outcome=outcome,
                message=message,
                payload=diagnostic_payload,
            )
            self.session.commit()
            response_item: dict[str, object] = {
                "row_number": int(row_number),
                "input_kind": input_kind,
                "collector_no": collector_no,
                "outcome": outcome,
                "message": message,
            }
            if original_filename:
                response_item["original_filename"] = original_filename
            outcomes.append(response_item)

        processed: dict[str, tuple[str, str, dict[str, object], bool]] = {}
        for row_number, raw_collector_no in rows:
            collector_no = normalize_identifier(raw_collector_no)
            if not collector_no:
                outcome = "invalid"
                message = "采集器号为空"
                result: dict[str, object] = {}
                photo_referenced = False
            elif collector_no in seen:
                outcome = "reused"
                message = "批次内重复，已复用前一条结果"
                result = {}
                photo_referenced = False
            else:
                seen.add(collector_no)
                outcome, message, result, photo_referenced = process_collector(
                    collector_no,
                    photos_by_collector.get(collector_no),
                )
                processed[collector_no] = (outcome, message, result, photo_referenced)
            persist_diagnostic(
                row_number=int(row_number),
                input_kind="excel",
                collector_no=collector_no,
                outcome=outcome,
                message=message,
                result=result,
            )

        for photo_input in photo_inputs:
            row_number = int(photo_input.get("row_number") or 0)
            collector_no = normalize_identifier(photo_input.get("collector_no"))
            original_filename = normalize_identifier(photo_input.get("original_filename"))
            status = normalize_identifier(photo_input.get("status"))
            if status == "invalid" or not collector_no:
                outcome = "invalid"
                message = normalize_identifier(photo_input.get("message")) or "照片文件名无效"
                result = {}
            elif status == "duplicate":
                outcome = "reused"
                message = normalize_identifier(photo_input.get("message")) or "同号照片已复用第一张"
                result = {}
            elif collector_no in processed:
                prior_outcome, _prior_message, result, photo_referenced = processed[collector_no]
                if prior_outcome == "invalid":
                    outcome = "invalid"
                    message = "关联采集器登记失败，照片未使用"
                elif photo_referenced:
                    outcome = "inserted"
                    message = "照片已随 Excel 行登记"
                else:
                    outcome = "reused"
                    message = "台账已有可复用照片，本次照片未使用"
            else:
                outcome, message, result, photo_referenced = process_collector(
                    collector_no,
                    photo_input,
                )
                processed[collector_no] = (outcome, message, result, photo_referenced)
            persist_diagnostic(
                row_number=row_number,
                input_kind="photo",
                collector_no=collector_no,
                outcome=outcome,
                message=message,
                result=result,
                original_filename=original_filename,
            )
        return {"batch_id": batch_id, "total": len(outcomes), **counters, "rows": outcomes}
