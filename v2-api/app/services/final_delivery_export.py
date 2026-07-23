from __future__ import annotations

import hashlib
import json
import os
import re
import stat
import threading
from collections import deque
from contextlib import contextmanager
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from uuid import uuid4
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile
from zoneinfo import ZoneInfo


OLD_DEVICE_HEADERS = (
    "终端地址",
    "终端厂家",
    "终端安装地址",
    "表号",
    "采集器设备号",
    "供服中心",
    "改造工程队",
    "旧设备拆除时间",
    "旧设备状态",
)
NEW_DEVICE_HEADERS = ("改造日期", "模块对应的表号", "新模块编号", "备注")
STANDARD_PHOTO_CATEGORIES = {
    "before_box": "表箱整体改造前",
    "collector_barcode": "采集器条形码",
    "module_meter": "模块与电能表",
    "after_box": "表箱整体改造后",
}
STANDARD_PHOTO_ORDER = tuple(STANDARD_PHOTO_CATEGORIES)
PACKAGE_TTL = timedelta(days=7)
PACKAGE_INTEGRITY_SCHEMA_VERSION = 1
MAX_DELIVERY_PHOTO_BYTES = 32 * 1024 * 1024
MAX_DELIVERY_PACKAGE_INPUT_BYTES = 8 * 1024 * 1024 * 1024
LOCAL_TIMEZONE = ZoneInfo("Asia/Shanghai")

_INVALID_WINDOWS_PART = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_TRAILING_WINDOWS_PART = re.compile(r"[ .]+$")
@dataclass
class _PackageLockEntry:
    lock: threading.Lock
    references: int = 0


_PACKAGE_LOCKS: dict[str, _PackageLockEntry] = {}
_PACKAGE_LOCKS_GUARD = threading.Lock()
_ACTIVE_CACHE_PATHS: dict[str, int] = {}
_CACHE_PATH_CONDITION = threading.Condition(threading.Lock())


@dataclass(eq=False)
class DeliveryCacheFileLock:
    handle: Any
    path: Path
    exclusive: bool
    platform_token: Any = field(default=None, repr=False)
    _released: bool = field(default=False, init=False, repr=False)
    _release_lock: threading.Lock = field(default_factory=threading.Lock, init=False, repr=False)

    def release(self) -> None:
        with self._release_lock:
            if self._released:
                return
            self._released = True
        try:
            _unlock_file_handle(self.handle, self.platform_token)
        finally:
            self.handle.close()

    def downgrade_to_shared(self) -> None:
        if not self.exclusive:
            return
        self.platform_token = _downgrade_file_handle(self.handle, self.platform_token)
        self.exclusive = False


class DeliveryPackageValidationError(ValueError):
    def __init__(self, errors: Sequence[Mapping[str, Any]]):
        self.errors = [dict(error) for error in errors]
        super().__init__("Formal delivery package validation failed")


@dataclass(eq=False)
class LeasedDeliveryPackage:
    path: Path
    lease: str
    file_lock: DeliveryCacheFileLock | None = field(default=None, repr=False)
    _released: bool = field(default=False, init=False, repr=False)
    _release_lock: threading.Lock = field(default_factory=threading.Lock, init=False, repr=False)

    def __fspath__(self) -> str:
        return os.fspath(self.path)

    def __getattr__(self, name: str) -> Any:
        return getattr(self.path, name)

    def __eq__(self, other: object) -> bool:
        if isinstance(other, LeasedDeliveryPackage):
            return self.path == other.path
        return self.path == other

    def release(self) -> None:
        with self._release_lock:
            if self._released:
                return
            self._released = True
        release_delivery_cache_path(self.lease)
        if self.file_lock is not None:
            self.file_lock.release()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _first_text(*values: Any) -> str:
    return next((_text(value) for value in values if _text(value)), "")


def _active_photos(group: Mapping[str, Any]) -> list[dict[str, Any]]:
    from app.services.local_simulation import is_valid_photo_evidence

    return [
        dict(photo)
        for photo in group.get("photos", []) or []
        if isinstance(photo, Mapping) and is_valid_photo_evidence(photo)
    ]


def _group_identity(group: Mapping[str, Any]) -> dict[str, str]:
    photos = _active_photos(group)
    return {
        "terminal": _text(group.get("terminal")),
        "meter_no": _first_text(group.get("meter_no"), group.get("display_meter_no")),
        "module_asset_no": _first_text(
            group.get("construction_module_asset_no"),
            group.get("module_asset_no"),
            group.get("asset_no"),
            *(photo.get("module_asset_no") or photo.get("asset_no") for photo in photos),
        ),
        "collector": _first_text(
            group.get("construction_collector"),
            group.get("collector"),
            *(photo.get("collector") for photo in photos),
        ),
        "address": _first_text(group.get("address"), group.get("installation_address")),
    }


def _completed_at(group: Mapping[str, Any]) -> str:
    photos = _active_photos(group)
    return _first_text(
        group.get("client_completed_at"),
        *(photo.get("client_completed_at") for photo in photos),
    )


def _local_date(value: Any) -> str:
    text = _text(value)
    if not text:
        raise ValueError("missing date")
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError("invalid date") from exc
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(LOCAL_TIMEZONE)
    return parsed.date().isoformat()


def group_is_formally_archived(group: Mapping[str, Any]) -> bool:
    status = _text(group.get("status")).lower()
    archive_status = _text(group.get("archive_status")).lower()
    verification = group.get("barcode_verification") or {}
    photos = _active_photos(group)
    return (
        status == "archived"
        or archive_status == "archived"
        or bool(group.get("archived_at"))
        or _text(verification.get("auto_archive_status")).lower() == "archived"
        or (bool(photos) and all(_text(photo.get("archive_status")).lower() == "archived" for photo in photos))
    )


def _is_controlled_cache_path(value: Any) -> bool:
    text = _text(value).replace("\\", "/")
    if not text or text.startswith("/"):
        return False
    path = PurePosixPath(text)
    return not path.is_absolute() and ".." not in path.parts


def delivery_group_readiness(group: Mapping[str, Any]) -> dict[str, Any]:
    identity = _group_identity(group)
    photos = _active_photos(group)
    constructed = len(photos) == 4
    archived = group_is_formally_archived(group)
    cache_path_uncontrolled = any(
        _text(photo.get("delivery_cache_path")) and not _is_controlled_cache_path(photo.get("delivery_cache_path"))
        for photo in photos
    )
    cache_ready = (
        constructed
        and archived
        and not cache_path_uncontrolled
        and all(_text(photo.get("delivery_cache_status")) == "ready" for photo in photos)
        and all(_is_controlled_cache_path(photo.get("delivery_cache_path")) for photo in photos)
    )
    return {
        "constructed": constructed,
        "archived": archived,
        "cache_ready": cache_ready,
        "cache_pending": archived and not cache_ready and not cache_path_uncontrolled,
        "cache_path_uncontrolled": archived and cache_path_uncontrolled,
        "identity_ready": all(identity.values()),
        "valid_photo_count": len(photos),
    }


def _error(group_id: str, code: str, field: str, message: str) -> dict[str, str]:
    return {"group_id": group_id, "code": code, "field": field, "message": message}


def _deduplicated_groups(groups: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for group in groups:
        payload = dict(group)
        group_id = _text(payload.get("id"))
        key = group_id or hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        ).hexdigest()
        if key in seen:
            continue
        seen.add(key)
        result.append(payload)
    return result


def collect_delivery_validation_errors(
    groups: Iterable[Mapping[str, Any]],
    *,
    require_cache: bool = True,
) -> list[dict[str, Any]]:
    from app.services.local_simulation import validate_real_formal_identity_value

    deduplicated = _deduplicated_groups(groups)
    if not deduplicated:
        return [_error("", "no_groups", "scope", "No archived groups are eligible for formal delivery")]
    errors: list[dict[str, str]] = []
    identities: dict[str, dict[str, list[str]]] = {
        field: {} for field in ("terminal", "meter_no", "module_asset_no", "collector")
    }
    for index, group in enumerate(deduplicated, start=1):
        group_id = _text(group.get("id")) or f"group-{index}"
        identity = _group_identity(group)
        if _text(group.get("status")).lower() != "approved":
            errors.append(
                _error(group_id, "review_not_approved", "status", "Current group review state is not approved")
            )
        for field, code in (
            ("terminal", "missing_terminal"),
            ("meter_no", "missing_meter_no"),
            ("module_asset_no", "missing_module_no"),
            ("collector", "missing_collector"),
            ("address", "missing_address"),
        ):
            if not identity[field]:
                errors.append(_error(group_id, code, field, f"Missing required {field}"))
        for field in ("terminal", "meter_no", "module_asset_no", "collector"):
            if not identity[field]:
                continue
            try:
                validate_real_formal_identity_value(identity[field], f"delivery {field}")
            except ValueError:
                errors.append(_error(group_id, "placeholder_identity", field, "Placeholder identity is forbidden"))
            else:
                identities[field].setdefault(identity[field], []).append(group_id)
        if not group_is_formally_archived(group):
            errors.append(_error(group_id, "not_archived", "status", "Group is not archived"))
        try:
            _local_date(_completed_at(group))
        except ValueError:
            errors.append(_error(group_id, "invalid_completed_at", "client_completed_at", "Missing or invalid completion date"))
        photos = _active_photos(group)
        if len(photos) != 4:
            errors.append(_error(group_id, "invalid_photo_count", "photos", "Exactly four active photos are required"))
        categories = [_text(photo.get("category")) for photo in photos]
        if len(categories) != 4 or set(categories) != set(STANDARD_PHOTO_CATEGORIES):
            errors.append(
                _error(group_id, "invalid_photo_categories", "photos", "Four unique standard photo categories are required")
            )
        if require_cache:
            for photo in photos:
                if _text(photo.get("delivery_cache_status")) == "invalid":
                    errors.append(
                        _error(group_id, "delivery_cache_invalid", "photos", "Completed photo cache is invalid")
                    )
                    break
                if _text(photo.get("delivery_cache_status")) != "ready" or not _is_controlled_cache_path(
                    photo.get("delivery_cache_path")
                ):
                    errors.append(
                        _error(group_id, "delivery_cache_pending", "photos", "Completed original photo cache is required")
                    )
                    break
    for field, values in identities.items():
        for value, group_ids in values.items():
            unique_group_ids = sorted(set(group_ids))
            if len(unique_group_ids) < 2:
                continue
            for group_id in unique_group_ids:
                errors.append(
                    _error(group_id, "device_conflict", field, f"Device identifier {value} is used by multiple groups")
                )
    return errors


def validate_delivery_groups(
    groups: Iterable[Mapping[str, Any]],
    *,
    require_cache: bool = True,
) -> list[dict[str, Any]]:
    deduplicated = _deduplicated_groups(groups)
    errors = collect_delivery_validation_errors(deduplicated, require_cache=require_cache)
    if errors:
        raise DeliveryPackageValidationError(errors)
    return deduplicated


def _delivery_remark(group: Mapping[str, Any]) -> str:
    notes = [
        _text(item)
        for item in (
            *(group.get("exception_reasons") or []),
            group.get("exception_note"),
        )
        if _text(item)
    ]
    old_meter_no = _first_text(
        group.get("replacement_old_meter_no"),
        *(photo.get("replacement_old_meter_no") for photo in _active_photos(group)),
    )
    if old_meter_no:
        notes.insert(0, f"换表：旧表号 {old_meter_no}")
    for field in ("replacement_remark", "export_remark", "delivery_export_remark"):
        value = _text(group.get(field))
        if value:
            notes.append(value)
    manually_confirmed = bool(group.get("group_barcode_manual_confirmed")) or _text(
        group.get("auto_archive_source")
    ) == "manual_confirmed"
    if manually_confirmed and "人工确认" not in notes:
        notes.append("人工确认")
    return "；".join(dict.fromkeys(notes))


def _workbook_rows(group: Mapping[str, Any]) -> dict[str, tuple[str, ...]]:
    identity = _group_identity(group)
    return {
        "old": (
            identity["terminal"],
            "",
            identity["address"],
            identity["meter_no"],
            identity["collector"],
            "南大供电服务中心",
            "奕福",
            "",
            "",
        ),
        "new": (
            _local_date(_completed_at(group)),
            identity["meter_no"],
            identity["module_asset_no"],
            _delivery_remark(group),
        ),
    }


def _append_text_row(sheet: Any, values: Sequence[Any]) -> None:
    sheet.append([str(value or "") for value in values])
    for cell in sheet[sheet.max_row]:
        cell.number_format = "@"
        cell.data_type = "s"


def build_delivery_workbook(groups: Iterable[Mapping[str, Any]]) -> bytes:
    from openpyxl import Workbook

    validated = validate_delivery_groups(groups, require_cache=False)
    workbook = Workbook()
    old_sheet = workbook.active
    old_sheet.title = "老设备"
    new_sheet = workbook.create_sheet("新设备")
    old_sheet.append(list(OLD_DEVICE_HEADERS))
    new_sheet.append(list(NEW_DEVICE_HEADERS))
    old_sheet.freeze_panes = "A2"
    new_sheet.freeze_panes = "A2"
    for group in validated:
        rows = _workbook_rows(group)
        _append_text_row(old_sheet, rows["old"])
        _append_text_row(new_sheet, rows["new"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def sanitize_delivery_path_part(value: Any) -> str:
    text = _INVALID_WINDOWS_PART.sub("_", str(value or ""))
    text = _TRAILING_WINDOWS_PART.sub(lambda match: "_" * len(match.group()), text)
    return text or "_"


def _photo_extension(photo: Mapping[str, Any]) -> str:
    for value in (
        photo.get("original_filename"),
        photo.get("source_file"),
        str(photo.get("image_url") or "").split("?", 1)[0],
        photo.get("delivery_cache_path"),
    ):
        suffix = Path(str(value or "")).suffix.lower()
        if re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
            return suffix
    return ".jpg"


def _windows_member_key(path: str) -> str:
    return "/".join(part.rstrip(" .").casefold() for part in path.replace("\\", "/").split("/"))


def _collision_name(path: str, group_id: str, discriminator: str = "") -> str:
    candidate = Path(path)
    suffix = candidate.suffix
    stable_id = sanitize_delivery_path_part(group_id)
    if discriminator:
        stable_id = f"{stable_id}-{hashlib.sha256(discriminator.encode('utf-8')).hexdigest()[:8]}"
    return str(candidate.with_name(f"{candidate.stem}-{stable_id}{suffix}")).replace("\\", "/")


def _delivery_photo_members(groups: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for group in sorted(groups, key=lambda item: _text(item.get("id"))):
        identity = _group_identity(group)
        group_id = _text(group.get("id")) or "group"
        directory = "/".join(
            (
                sanitize_delivery_path_part(identity["terminal"]),
                sanitize_delivery_path_part(
                    f'{identity["meter_no"]}-{identity["module_asset_no"]}-{identity["address"]}'
                ),
            )
        )
        photos = {str(photo.get("category")): photo for photo in _active_photos(group)}
        for category in STANDARD_PHOTO_ORDER:
            photo = photos[category]
            filename = sanitize_delivery_path_part(
                f'{identity["module_asset_no"]}-{STANDARD_PHOTO_CATEGORIES[category]}'
            ) + _photo_extension(photo)
            path = f"{directory}/{filename}"
            candidates.append(
                {
                    "group_id": group_id,
                    "photo": photo,
                    "path": path,
                    "member_key": _windows_member_key(path),
                    "discriminator": f'{group_id}|{category}|{_text(photo.get("id"))}',
                    "category_order": STANDARD_PHOTO_ORDER.index(category),
                }
            )
    collision_counts: dict[str, int] = {}
    for candidate in candidates:
        key = candidate["member_key"]
        collision_counts[key] = collision_counts.get(key, 0) + 1
    assigned: set[str] = set()
    for candidate in sorted(
        candidates,
        key=lambda item: (item["member_key"], item["group_id"].casefold(), item["discriminator"]),
    ):
        path = candidate["path"]
        if collision_counts[candidate["member_key"]] > 1:
            path = _collision_name(path, candidate["group_id"])
        key = _windows_member_key(path)
        if key in assigned:
            path = _collision_name(candidate["path"], candidate["group_id"], candidate["discriminator"])
            key = _windows_member_key(path)
        assigned.add(key)
        candidate["path"] = path
    return sorted(
        candidates,
        key=lambda item: (item["group_id"].casefold(), item["group_id"], item["category_order"]),
    )


def _write_delivery_package_archive(
    output: Any,
    groups: Iterable[Mapping[str, Any]],
    photo_reader: Callable[[dict[str, Any]], bytes | Path],
    *,
    max_photo_bytes: int,
    max_total_input_bytes: int,
) -> None:
    validated = sorted(validate_delivery_groups(groups), key=lambda item: _text(item.get("id")))
    members = _delivery_photo_members(validated)
    errors: list[dict[str, str]] = []
    photo_limit = max(1, int(max_photo_bytes))
    aggregate_limit = max(1, int(max_total_input_bytes))
    total_input_bytes = 0
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("设备清单.xlsx", build_delivery_workbook(validated))
        for member in members:
            photo = member["photo"]
            group_id = member["group_id"]
            expected_sha256 = _text(photo.get("delivery_cache_content_sha256")).lower()
            try:
                source = photo_reader(photo)
                if isinstance(source, bytes):
                    size = len(source)
                    actual_sha256 = hashlib.sha256(source).hexdigest()
                elif isinstance(source, Path):
                    source = source.resolve(strict=True)
                    size = source.stat().st_size
                    digest = hashlib.sha256()
                    with source.open("rb") as cached_photo:
                        while chunk := cached_photo.read(1024 * 1024):
                            digest.update(chunk)
                    actual_sha256 = digest.hexdigest()
                else:
                    raise TypeError("Unsupported completed photo source")
            except DeliveryPackageValidationError as exc:
                errors.extend(exc.errors)
                continue
            except Exception:
                errors.append(
                    _error(group_id, "delivery_cache_invalid", "photos", "Completed photo cache could not be read")
                )
                continue
            if size > photo_limit:
                errors.append(
                    _error(group_id, "delivery_photo_too_large", "photos", "Completed photo exceeds package limit")
                )
                continue
            if total_input_bytes + size > aggregate_limit:
                errors.append(
                    _error(group_id, "delivery_package_too_large", "photos", "Delivery package input exceeds limit")
                )
                continue
            total_input_bytes += size
            if (
                size <= 0
                or not re.fullmatch(r"[0-9a-f]{64}", expected_sha256)
                or actual_sha256 != expected_sha256
            ):
                errors.append(
                    _error(
                        group_id,
                        "delivery_cache_invalid",
                        "photos",
                        "Completed photo cache failed integrity verification",
                    )
                )
                continue
            if isinstance(source, bytes):
                archive.writestr(member["path"], source)
            else:
                archive.write(source, arcname=member["path"])
    if errors:
        raise DeliveryPackageValidationError(errors)


def build_delivery_package(
    groups: Iterable[Mapping[str, Any]],
    photo_reader: Callable[[dict[str, Any]], bytes | Path],
    *,
    max_photo_bytes: int = MAX_DELIVERY_PHOTO_BYTES,
    max_total_input_bytes: int = MAX_DELIVERY_PACKAGE_INPUT_BYTES,
) -> bytes:
    output = BytesIO()
    _write_delivery_package_archive(
        output,
        groups,
        photo_reader,
        max_photo_bytes=max_photo_bytes,
        max_total_input_bytes=max_total_input_bytes,
    )
    return output.getvalue()


def build_delivery_package_file(
    target: Path,
    groups: Iterable[Mapping[str, Any]],
    photo_reader: Callable[[dict[str, Any]], bytes | Path],
    *,
    max_photo_bytes: int = MAX_DELIVERY_PHOTO_BYTES,
    max_total_input_bytes: int = MAX_DELIVERY_PACKAGE_INPUT_BYTES,
) -> Path:
    destination = Path(target)
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        _write_delivery_package_archive(
            destination,
            groups,
            photo_reader,
            max_photo_bytes=max_photo_bytes,
            max_total_input_bytes=max_total_input_bytes,
        )
    except BaseException:
        destination.unlink(missing_ok=True)
        raise
    return destination


def delivery_evidence_fingerprint(groups: Iterable[Mapping[str, Any]]) -> str:
    payload = []
    for group in sorted(_deduplicated_groups(groups), key=lambda item: _text(item.get("id"))):
        payload.append(
            {
                "id": _text(group.get("id")),
                "invalidation_epoch": int(group.get("delivery_package_invalidation_epoch") or 0),
                "workbook_rows": _workbook_rows(group),
                "photos": sorted(
                    (
                        {
                            "id": _text(photo.get("id")),
                            "category": _text(photo.get("category")),
                            "sha256": _first_text(
                                photo.get("delivery_cache_content_sha256"), photo.get("sha256")
                            ).lower(),
                            "cache_version": _text(photo.get("delivery_cache_version")),
                            "extension": _photo_extension(photo),
                        }
                        for photo in _active_photos(group)
                    ),
                    key=lambda item: (item["category"], item["id"]),
                ),
            }
        )
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _expected_delivery_package_members(groups: Sequence[Mapping[str, Any]]) -> list[str]:
    return ["设备清单.xlsx", *(member["path"] for member in _delivery_photo_members(groups))]


def _delivery_package_integrity_path(path: Path) -> Path:
    return path.with_suffix(f"{path.suffix}.integrity.json")


def _package_file_sha256(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as package_file:
        while chunk := package_file.read(1024 * 1024):
            size += len(chunk)
            digest.update(chunk)
    return digest.hexdigest(), size


def _safe_archive_member(info: Any) -> bool:
    name = _text(info.filename)
    member = PurePosixPath(name)
    unix_mode = (int(info.external_attr) >> 16) & 0xFFFF
    return bool(
        name
        and "\\" not in name
        and not name.startswith("/")
        and re.match(r"^[A-Za-z]:", name) is None
        and not info.is_dir()
        and all(part not in {"", ".", ".."} for part in member.parts)
        and not (unix_mode and stat.S_ISLNK(unix_mode))
    )


def _verified_archive_members(path: Path, expected_members: Sequence[str]) -> list[str] | None:
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            members = [info.filename for info in infos]
            if (
                members != list(expected_members)
                or len(members) != len(set(members))
                or not all(_safe_archive_member(info) for info in infos)
                or archive.testzip() is not None
            ):
                return None
            return members
    except (BadZipFile, EOFError, OSError, RuntimeError, ValueError):
        return None


def _delivery_package_integrity_payload(
    path: Path,
    fingerprint: str,
    expected_members: Sequence[str],
) -> dict[str, Any] | None:
    members = _verified_archive_members(path, expected_members)
    if members is None:
        return None
    try:
        package_sha256, package_size = _package_file_sha256(path)
    except OSError:
        return None
    return {
        "schema_version": PACKAGE_INTEGRITY_SCHEMA_VERSION,
        "evidence_fingerprint": fingerprint,
        "package_sha256": package_sha256,
        "package_size": package_size,
        "members": members,
    }


def _write_delivery_package_integrity_proof(path: Path, payload: Mapping[str, Any]) -> None:
    with path.open("x", encoding="utf-8", newline="\n") as proof_file:
        json.dump(dict(payload), proof_file, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        proof_file.write("\n")
        proof_file.flush()
        os.fsync(proof_file.fileno())


def _windows_overlapped() -> Any:
    import ctypes
    from ctypes import wintypes

    class Overlapped(ctypes.Structure):
        _fields_ = [
            ("Internal", ctypes.c_size_t),
            ("InternalHigh", ctypes.c_size_t),
            ("Offset", wintypes.DWORD),
            ("OffsetHigh", wintypes.DWORD),
            ("hEvent", wintypes.HANDLE),
        ]

    return Overlapped()


def _lock_file_handle(handle: Any, *, blocking: bool, exclusive: bool) -> Any | None:
    if os.name == "nt":
        import ctypes
        import msvcrt
        from ctypes import wintypes

        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        overlapped = _windows_overlapped()
        flags = 0x00000002 if exclusive else 0
        if not blocking:
            flags |= 0x00000001
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        ctypes.set_last_error(0)
        acquired = kernel32.LockFileEx(
            wintypes.HANDLE(msvcrt.get_osfhandle(handle.fileno())),
            wintypes.DWORD(flags),
            wintypes.DWORD(0),
            wintypes.DWORD(1),
            wintypes.DWORD(0),
            ctypes.byref(overlapped),
        )
        if acquired:
            return overlapped
        error = ctypes.get_last_error()
        if not blocking and error in {32, 33, 158}:
            return None
        raise OSError(error, "Unable to lock delivery cache file")

    import fcntl

    operation = (fcntl.LOCK_EX if exclusive else fcntl.LOCK_SH) | (0 if blocking else fcntl.LOCK_NB)
    try:
        fcntl.flock(handle.fileno(), operation)
        return True
    except BlockingIOError:
        return None


def _unlock_file_handle(handle: Any, platform_token: Any) -> None:
    if os.name == "nt":
        import ctypes
        import msvcrt
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.UnlockFileEx(
            wintypes.HANDLE(msvcrt.get_osfhandle(handle.fileno())),
            wintypes.DWORD(0),
            wintypes.DWORD(1),
            wintypes.DWORD(0),
            ctypes.byref(platform_token),
        )
        return

    import fcntl

    fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _downgrade_file_handle(handle: Any, platform_token: Any) -> Any:
    if os.name == "nt":
        _unlock_file_handle(handle, platform_token)
        shared_token = _lock_file_handle(handle, blocking=True, exclusive=False)
        if shared_token is None:  # pragma: no cover - blocking acquisition returns a token or raises
            raise RuntimeError("Unable to downgrade delivery cache lock")
        return shared_token

    import fcntl

    fcntl.flock(handle.fileno(), fcntl.LOCK_SH)
    return platform_token


def acquire_delivery_cache_file_lock(
    cache_root: Path | str,
    target: Path | str,
    *,
    blocking: bool = True,
    exclusive: bool = True,
) -> DeliveryCacheFileLock | None:
    root = Path(cache_root).resolve()
    target_key = _cache_path_key(target)
    lock_root = root / ".locks"
    lock_root.mkdir(parents=True, exist_ok=True)
    lock_path = lock_root / f"{hashlib.sha256(target_key.encode('utf-8')).hexdigest()}.lock"
    handle = lock_path.open("a+b")
    try:
        platform_token = _lock_file_handle(handle, blocking=blocking, exclusive=exclusive)
        if platform_token is None:
            handle.close()
            return None
        return DeliveryCacheFileLock(
            handle=handle,
            path=lock_path,
            exclusive=exclusive,
            platform_token=platform_token,
        )
    except BaseException:
        handle.close()
        raise


@contextmanager
def _package_lock(key: str):
    with _PACKAGE_LOCKS_GUARD:
        entry = _PACKAGE_LOCKS.get(key)
        if entry is None:
            entry = _PackageLockEntry(threading.Lock())
            _PACKAGE_LOCKS[key] = entry
        entry.references += 1
    entry.lock.acquire()
    try:
        yield
    finally:
        entry.lock.release()
        with _PACKAGE_LOCKS_GUARD:
            entry.references -= 1
            if entry.references == 0 and _PACKAGE_LOCKS.get(key) is entry:
                _PACKAGE_LOCKS.pop(key, None)


def _cache_path_key(path: Path | str) -> str:
    return str(Path(path).resolve()).casefold()


def reserve_delivery_cache_path(path: Path | str) -> str:
    key = _cache_path_key(path)
    with _CACHE_PATH_CONDITION:
        _reserve_delivery_cache_key_locked(key)
    return key


def _reserve_delivery_cache_key_locked(key: str) -> None:
    _ACTIVE_CACHE_PATHS[key] = _ACTIVE_CACHE_PATHS.get(key, 0) + 1


def release_delivery_cache_path(lease: Path | str) -> None:
    key = str(lease).casefold()
    with _CACHE_PATH_CONDITION:
        count = _ACTIVE_CACHE_PATHS.get(key, 0)
        if count <= 1:
            _ACTIVE_CACHE_PATHS.pop(key, None)
        else:
            _ACTIVE_CACHE_PATHS[key] = count - 1
        _CACHE_PATH_CONDITION.notify_all()


def _cache_path_is_reserved(key: str) -> bool:
    separator = os.sep.casefold()
    return any(
        count > 0 and (key == reserved or key.startswith(f"{reserved.rstrip(separator)}{separator}"))
        for reserved, count in _ACTIVE_CACHE_PATHS.items()
    )


def _fresh_cleanup_cursor() -> dict[str, Any]:
    return {"pending_directories": ["."], "current_directory": None, "after_name": ""}


def _normalized_cleanup_cursor(value: Any) -> dict[str, Any]:
    source = dict(value) if isinstance(value, Mapping) else {}
    pending = [
        str(item)
        for item in source.get("pending_directories", [])
        if isinstance(item, str) and item
    ]
    current = source.get("current_directory")
    if not isinstance(current, str) or not current:
        current = None
    if current is None and not pending:
        pending = ["."]
    return {
        "pending_directories": pending,
        "current_directory": current,
        "after_name": str(source.get("after_name") or ""),
        "cycle_bytes": max(0, int(source.get("cycle_bytes") or 0)),
        "known_total_bytes": max(0, int(source.get("known_total_bytes") or 0)),
    }


def _load_cleanup_cursors(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, OSError, ValueError, TypeError):
        payload = {}
    return {
        "objects": _normalized_cleanup_cursor(payload.get("objects")),
        "packages": _normalized_cleanup_cursor(payload.get("packages")),
    }


def _save_cleanup_cursors(path: Path, cursors: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(f".tmp-{uuid4().hex}")
    try:
        temporary.write_text(
            json.dumps(cursors, ensure_ascii=True, sort_keys=True, separators=(",", ":")),
            encoding="utf-8",
        )
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _bounded_cache_files(
    root: Path,
    *,
    max_entries: int,
    suffix: str | None = None,
    cursor: Mapping[str, Any] | None = None,
) -> tuple[list[Path], int, bool, dict[str, Any]]:
    limit = max(0, int(max_entries))
    state = _normalized_cleanup_cursor(cursor)
    if limit == 0 or not root.exists():
        return [], 0, root.exists(), state
    pending = deque(state["pending_directories"])
    current_directory = state["current_directory"]
    after_name = state["after_name"]
    files: list[Path] = []
    scanned = 0
    completed_cycle = False
    while scanned < limit:
        if current_directory is None:
            if not pending:
                completed_cycle = True
                break
            current_directory = pending.popleft()
            after_name = ""
        directory = root if current_directory == "." else root / current_directory
        try:
            with os.scandir(directory) as entries:
                ordered = sorted(entries, key=lambda entry: (entry.name.casefold(), entry.name))
        except FileNotFoundError:
            current_directory = None
            after_name = ""
            continue
        remaining = [
            entry
            for entry in ordered
            if not after_name or (entry.name.casefold(), entry.name) > (after_name.casefold(), after_name)
        ]
        if not remaining:
            current_directory = None
            after_name = ""
            continue
        processed = 0
        for entry in remaining:
            if scanned >= limit:
                break
            scanned += 1
            processed += 1
            after_name = entry.name
            try:
                if entry.is_dir(follow_symlinks=False):
                    child = Path(entry.path).relative_to(root)
                    pending.append(str(child).replace("\\", "/"))
                elif entry.is_file(follow_symlinks=False):
                    path = Path(entry.path)
                    if suffix is None or path.suffix.casefold() == suffix.casefold():
                        files.append(path)
            except FileNotFoundError:
                continue
        if processed == len(remaining):
            current_directory = None
            after_name = ""
    if current_directory is None and not pending:
        completed_cycle = True
    if completed_cycle:
        pending = deque(["."])
        current_directory = None
        after_name = ""
    state.update(
        {
            "pending_directories": list(pending),
            "current_directory": current_directory,
            "after_name": after_name,
        }
    )
    return files, scanned, not completed_cycle, state


def _scan_cache_tree_with_cursor(
    cache_root: Path,
    tree_name: str,
    tree_root: Path,
    *,
    max_entries: int,
    suffix: str | None = None,
    track_bytes: bool = False,
) -> tuple[list[Path], int, bool, int]:
    cursor_path = cache_root / ".delivery-cleanup-cursor.json"
    cursor_lock = acquire_delivery_cache_file_lock(
        cache_root,
        cursor_path,
        blocking=True,
    )
    if cursor_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
        raise RuntimeError("Unable to lock delivery cleanup cursor")
    try:
        cursors = _load_cleanup_cursors(cursor_path)
        files, scanned, truncated, state = _bounded_cache_files(
            tree_root,
            max_entries=max_entries,
            suffix=suffix,
            cursor=cursors.get(tree_name),
        )
        if track_bytes:
            scanned_bytes = 0
            for path in files:
                try:
                    scanned_bytes += path.stat().st_size
                except FileNotFoundError:
                    continue
            state["cycle_bytes"] = int(state.get("cycle_bytes") or 0) + scanned_bytes
            if not truncated:
                state["known_total_bytes"] = int(state["cycle_bytes"])
                state["cycle_bytes"] = 0
        cursors[tree_name] = state
        _save_cleanup_cursors(cursor_path, cursors)
        return files, scanned, truncated, int(state.get("known_total_bytes") or 0)
    finally:
        cursor_lock.release()


def _reduce_cleanup_known_bytes(cache_root: Path, tree_name: str, removed_bytes: int) -> None:
    removed = max(0, int(removed_bytes))
    if removed == 0:
        return
    cursor_path = cache_root / ".delivery-cleanup-cursor.json"
    cursor_lock = acquire_delivery_cache_file_lock(cache_root, cursor_path, blocking=True)
    if cursor_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
        raise RuntimeError("Unable to lock delivery cleanup cursor")
    try:
        cursors = _load_cleanup_cursors(cursor_path)
        state = _normalized_cleanup_cursor(cursors.get(tree_name))
        state["known_total_bytes"] = max(0, int(state["known_total_bytes"]) - removed)
        state["cycle_bytes"] = max(0, int(state["cycle_bytes"]) - removed)
        cursors[tree_name] = state
        _save_cleanup_cursors(cursor_path, cursors)
    finally:
        cursor_lock.release()


def cleanup_delivery_cache(
    cache_root: Path,
    *,
    max_object_bytes: int,
    groups: Iterable[Mapping[str, Any]],
    now: datetime | None = None,
    max_scan_entries: int = 2_000,
    cleanup_objects: bool = True,
) -> dict[str, int]:
    root = Path(cache_root).resolve()
    current = now or datetime.now(UTC)
    if current.tzinfo is None:
        current = current.replace(tzinfo=UTC)
    referenced: set[str] = set()
    for group in groups:
        if _text(group.get("delivery_cache_status")) != "ready":
            continue
        for photo in _active_photos(group):
            if _text(photo.get("delivery_cache_status")) != "ready":
                continue
            relative = _text(photo.get("delivery_cache_path"))
            if not relative:
                continue
            candidate = (root / relative).resolve()
            try:
                candidate.relative_to(root)
            except ValueError:
                continue
            referenced.add(_cache_path_key(candidate))

    object_root = root / "objects"
    object_files: list[Path] = []
    scanned_object_entries = 0
    object_scan_truncated = False
    object_tree_busy = 0
    known_object_bytes = 0
    object_lock = None
    if cleanup_objects and object_root.exists():
        object_lock = acquire_delivery_cache_file_lock(root, object_root, blocking=False)
        if object_lock is None:
            object_tree_busy = 1
        else:
            object_files, scanned_object_entries, object_scan_truncated, known_object_bytes = _scan_cache_tree_with_cursor(
                root,
                "objects",
                object_root,
                max_entries=max_scan_entries,
                track_bytes=True,
            )
    scanned_object_bytes = sum(path.stat().st_size for path in object_files if path.exists())
    total_bytes = max(scanned_object_bytes, known_object_bytes)
    deleted_objects = 0
    deleted_object_bytes = 0
    protected_objects = 0
    try:
        for path in sorted(object_files, key=lambda item: (item.stat().st_mtime, str(item))):
            key = _cache_path_key(path)
            if key in referenced:
                protected_objects += 1
                continue
            if total_bytes <= max(0, int(max_object_bytes)):
                continue
            with _CACHE_PATH_CONDITION:
                if _cache_path_is_reserved(key):
                    protected_objects += 1
                    continue
                try:
                    size = path.stat().st_size
                    path.unlink()
                except FileNotFoundError:
                    continue
                total_bytes -= size
                deleted_objects += 1
                deleted_object_bytes += size
    finally:
        if object_lock is not None:
            object_lock.release()
    _reduce_cleanup_known_bytes(root, "objects", deleted_object_bytes)

    deleted_packages = 0
    scanned_package_entries = 0
    package_scan_truncated = False
    package_root = root / "packages"
    if package_root.exists():
        expires_before = current.astimezone(UTC) - PACKAGE_TTL
        package_files, scanned_package_entries, package_scan_truncated, _known_package_bytes = _scan_cache_tree_with_cursor(
            root,
            "packages",
            package_root,
            max_entries=max_scan_entries,
            suffix=".zip",
        )
        for path in package_files:
            key = _cache_path_key(path)
            with _CACHE_PATH_CONDITION:
                if _cache_path_is_reserved(key):
                    continue
                build_lock = acquire_delivery_cache_file_lock(
                    root,
                    path.with_suffix(f"{path.suffix}.build"),
                    blocking=False,
                )
                if build_lock is None:
                    continue
                try:
                    file_lock = acquire_delivery_cache_file_lock(root, path, blocking=False)
                    if file_lock is None:
                        continue
                    try:
                        try:
                            modified = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
                            if modified < expires_before:
                                path.unlink()
                                _delivery_package_integrity_path(path).unlink(missing_ok=True)
                                deleted_packages += 1
                        except FileNotFoundError:
                            continue
                    finally:
                        file_lock.release()
                finally:
                    build_lock.release()
    return {
        "deleted_objects": deleted_objects,
        "protected_objects": protected_objects,
        "remaining_object_bytes": total_bytes,
        "deleted_packages": deleted_packages,
        "scanned_object_entries": scanned_object_entries,
        "scanned_package_entries": scanned_package_entries,
        "object_scan_truncated": int(object_scan_truncated),
        "package_scan_truncated": int(package_scan_truncated),
        "object_tree_busy": object_tree_busy,
    }


def _delivery_package_is_reusable(
    path: Path,
    current: datetime,
    fingerprint: str,
    expected_members: Sequence[str],
) -> bool:
    try:
        if not path.is_file():
            return False
        modified = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
        if current.astimezone(UTC) - modified > PACKAGE_TTL:
            return False
        proof = json.loads(_delivery_package_integrity_path(path).read_text(encoding="utf-8"))
        if not isinstance(proof, dict):
            return False
        expected_proof = _delivery_package_integrity_payload(path, fingerprint, expected_members)
        return expected_proof is not None and proof == expected_proof
    except (BadZipFile, EOFError, OSError, RuntimeError, TypeError, ValueError, json.JSONDecodeError):
        return False


def get_or_build_delivery_package(
    scope: str,
    evidence_fingerprint: str,
    *,
    groups: Iterable[Mapping[str, Any]],
    photo_reader: Callable[[dict[str, Any]], bytes | Path],
    cache_root: Path,
    now: datetime | None = None,
    package_builder: Callable[
        [Iterable[Mapping[str, Any]], Callable[[dict[str, Any]], bytes | Path]], bytes
    ] = build_delivery_package,
    package_file_builder: Callable[
        [Path, Iterable[Mapping[str, Any]], Callable[[dict[str, Any]], bytes | Path]], Path
    ]
    | None = None,
) -> LeasedDeliveryPackage:
    validated_groups = validate_delivery_groups(groups)
    current = now or datetime.now(UTC)
    if current.tzinfo is None:
        current = current.replace(tzinfo=UTC)
    provided_fingerprint = _text(evidence_fingerprint).lower()
    if not re.fullmatch(r"[0-9a-f]{64}", provided_fingerprint):
        raise ValueError("Invalid delivery evidence fingerprint")
    fingerprint = delivery_evidence_fingerprint(validated_groups)
    if provided_fingerprint != fingerprint:
        raise ValueError("Delivery evidence fingerprint does not match current delivery evidence")
    expected_members = _expected_delivery_package_members(validated_groups)
    scope_hash = hashlib.sha256(_text(scope).encode("utf-8")).hexdigest()[:24]
    package_dir = Path(cache_root).resolve() / "packages" / scope_hash
    target = package_dir / f"{fingerprint}.zip"
    lock_key = str(target).lower()
    with _package_lock(lock_key):
        target_key = _cache_path_key(target)
        file_lock = acquire_delivery_cache_file_lock(
            cache_root,
            target,
            blocking=True,
            exclusive=False,
        )
        if file_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
            raise RuntimeError("Unable to lock delivery package path")
        try:
            with _CACHE_PATH_CONDITION:
                if _delivery_package_is_reusable(target, current, fingerprint, expected_members):
                    _reserve_delivery_cache_key_locked(target_key)
                    package = LeasedDeliveryPackage(target, target_key, file_lock=file_lock)
                    file_lock = None
                    return package
        finally:
            if file_lock is not None:
                file_lock.release()

        build_lock = acquire_delivery_cache_file_lock(
            cache_root,
            target.with_suffix(f"{target.suffix}.build"),
            blocking=True,
        )
        if build_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
            raise RuntimeError("Unable to lock delivery package build")
        try:
            file_lock = acquire_delivery_cache_file_lock(
                cache_root,
                target,
                blocking=True,
                exclusive=False,
            )
            if file_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
                raise RuntimeError("Unable to lock delivery package path")
            try:
                with _CACHE_PATH_CONDITION:
                    if _delivery_package_is_reusable(target, current, fingerprint, expected_members):
                        _reserve_delivery_cache_key_locked(target_key)
                        package = LeasedDeliveryPackage(target, target_key, file_lock=file_lock)
                        file_lock = None
                        return package
                package_dir.mkdir(parents=True, exist_ok=True)
                temporary = target.with_suffix(f".zip.tmp-{uuid4().hex}")
                integrity_target = _delivery_package_integrity_path(target)
                integrity_temporary = package_dir / f".integrity-{uuid4().hex}.tmp"
                try:
                    try:
                        if package_file_builder is not None:
                            package_file_builder(temporary, validated_groups, photo_reader)
                            with temporary.open("rb+") as output:
                                output.flush()
                                os.fsync(output.fileno())
                        else:
                            content = package_builder(validated_groups, photo_reader)
                            with temporary.open("xb") as output:
                                output.write(content)
                                output.flush()
                                os.fsync(output.fileno())
                        integrity_payload = _delivery_package_integrity_payload(
                            temporary,
                            fingerprint,
                            expected_members,
                        )
                        if integrity_payload is None:
                            raise DeliveryPackageValidationError(
                                [_error("", "delivery_package_invalid", "package", "Built delivery package is invalid")]
                            )
                        _write_delivery_package_integrity_proof(integrity_temporary, integrity_payload)
                        os.utime(temporary, (current.timestamp(), current.timestamp()))
                        os.utime(integrity_temporary, (current.timestamp(), current.timestamp()))
                    finally:
                        if file_lock is not None:
                            file_lock.release()
                            file_lock = None

                    replacement_lock = acquire_delivery_cache_file_lock(cache_root, target, blocking=True)
                    if replacement_lock is None:  # pragma: no cover - blocking acquisition returns a lock or raises
                        raise RuntimeError("Unable to lock delivery package replacement")
                    try:
                        with _CACHE_PATH_CONDITION:
                            while _cache_path_is_reserved(target_key):
                                _CACHE_PATH_CONDITION.wait()
                            temporary.replace(target)
                            integrity_temporary.replace(integrity_target)
                            _reserve_delivery_cache_key_locked(target_key)
                            replacement_lock.downgrade_to_shared()
                            package = LeasedDeliveryPackage(target, target_key, file_lock=replacement_lock)
                            replacement_lock = None
                            return package
                    finally:
                        if replacement_lock is not None:
                            replacement_lock.release()
                finally:
                    temporary.unlink(missing_ok=True)
                    integrity_temporary.unlink(missing_ok=True)
            finally:
                if file_lock is not None:
                    file_lock.release()
        finally:
            build_lock.release()
