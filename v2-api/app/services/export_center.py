from __future__ import annotations

import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
from uuid import uuid4

from app.services import local_simulation
from app.services.delivery_package_queue import DeliveryPackageNotReady
from app.services.final_delivery_export import LeasedDeliveryPackage


SUPPORTED_EXPORT_JOB_PAGE_SIZES = {20, 50, 100}
PLACEHOLDER_DEVICE_VALUES = {"", "00000000", "UNKNOWN", "N/A", "NULL", "-"}
DEVICE_EXPORT_KINDS = {"terminal", "meter", "module", "collector"}

EXPORT_CATALOG: tuple[dict[str, str], ...] = (
    {"key": "final_delivery", "label": "正式交付包", "delivery": "zip", "mode": "background"},
    {"key": "device_terminal", "label": "终端设备清单", "delivery": "xlsx", "mode": "inline"},
    {"key": "device_meter", "label": "表号设备清单", "delivery": "xlsx", "mode": "inline"},
    {"key": "device_module", "label": "模块设备清单", "delivery": "xlsx", "mode": "inline"},
    {"key": "device_collector", "label": "采集器设备清单", "delivery": "xlsx", "mode": "inline"},
    {"key": "task_detail", "label": "任务明细", "delivery": "xlsx", "mode": "inline"},
    {"key": "exception_meter", "label": "异常表计", "delivery": "xlsx", "mode": "inline"},
    {"key": "exception_missing_photo", "label": "缺图异常", "delivery": "xlsx", "mode": "inline"},
    {"key": "replacement", "label": "换表记录", "delivery": "xlsx", "mode": "inline"},
    {"key": "unmatched", "label": "未匹配记录", "delivery": "xlsx", "mode": "inline"},
    {"key": "barcode_review", "label": "条码复核", "delivery": "xlsx", "mode": "inline"},
    {"key": "project_outside", "label": "项目外记录", "delivery": "xlsx", "mode": "inline"},
    {"key": "installer_kpi", "label": "施工人员 KPI", "delivery": "xlsx", "mode": "inline"},
    {"key": "installer_daily_completion", "label": "施工人员每日完成", "delivery": "xlsx", "mode": "inline"},
)
CATALOG_BY_KEY = {item["key"]: item for item in EXPORT_CATALOG}


def normalize_export_page_size(value: int) -> int:
    page_size = int(value or 20)
    if page_size not in SUPPORTED_EXPORT_JOB_PAGE_SIZES:
        raise ValueError("Export job page_size must be one of 20, 50, or 100")
    return page_size


def normalize_device_value(value: object) -> str:
    text = str(value or "").strip()
    if text.upper() in PLACEHOLDER_DEVICE_VALUES:
        return ""
    if text[:1] in {"=", "+", "-", "@"}:
        return ""
    return text


def build_device_workbook(
    *,
    kind: str,
    rows: Iterable[object],
    project: Mapping[str, Any] | None = None,
    filters: Mapping[str, Any] | None = None,
) -> bytes:
    if kind not in DEVICE_EXPORT_KINDS:
        raise ValueError(f"Unsupported device export kind: {kind}")
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to export Excel files") from exc

    values = sorted({value for value in (normalize_device_value(row) for row in rows) if value})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{kind} devices"[:31]
    sheet.append([kind])
    for value in values:
        cell = sheet.cell(row=sheet.max_row + 1, column=1, value=value)
        cell.number_format = "@"
    sheet.column_dimensions["A"].width = max(12, min(32, max((len(value) for value in values), default=8) + 2))
    project = dict(project or {})
    filters = dict(filters or {})
    workbook.properties.title = f"{kind} device export"
    workbook.properties.subject = f"project={project.get('name') or project.get('id') or ''}"
    workbook.properties.description = (
        f"generated_at={datetime.now(UTC).isoformat()}; "
        f"count={len(values)}; filters={json.dumps(filters, ensure_ascii=False, sort_keys=True)}"
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def extract_device_values(groups: Iterable[Mapping[str, Any]], kind: str) -> list[str]:
    field_by_kind = {
        "terminal": "terminal",
        "meter": "meter_no",
        "module": "module_asset_no",
        "collector": "collector",
    }
    field = field_by_kind[kind]
    values: list[str] = []
    for group in groups:
        values.append(str(group.get(field) or ""))
        if kind == "module":
            values.append(str(group.get("asset_no") or ""))
        if kind in {"module", "collector"}:
            for photo in group.get("photos") or []:
                if isinstance(photo, Mapping):
                    values.append(str(photo.get(field) or photo.get("asset_no") or ""))
    return values


def build_simple_table_workbook(*, title: str, rows: Iterable[Mapping[str, Any]]) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to export Excel files") from exc

    materialized = [dict(row) for row in rows]
    headers = sorted({key for row in materialized for key in row})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(headers or ["message"])
    if not materialized:
        sheet.append([""])
    for row in materialized:
        sheet.append([row.get(header, "") for header in headers])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.number_format = "@"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def file_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def write_export_content(content: bytes, *, job_id: str, filename: str) -> tuple[str, str, int]:
    root = (local_simulation.delivery_cache_root() / "exports").resolve()
    root.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix or ".bin"
    target = (root / f"{job_id}{suffix}").resolve()
    target.relative_to(root)
    target.write_bytes(content)
    return str(target), file_sha256(content), len(content)


def read_export_content(path_value: str) -> bytes:
    root = local_simulation.delivery_cache_root().resolve()
    candidate = Path(path_value).resolve(strict=True)
    candidate.relative_to(root)
    return candidate.read_bytes()


def export_filename(job_type: str, suffix: str = "xlsx") -> str:
    stamp = datetime.now(UTC).strftime("%Y%m%d%H%M%S")
    safe_type = "".join(character if character.isalnum() or character in {"-", "_"} else "-" for character in job_type)
    return f"{safe_type}-{stamp}.{suffix}"


def build_inline_export_content(repository: Any, *, job_type: str, filters: Mapping[str, Any]) -> tuple[bytes, str, str]:
    filters = dict(filters or {})
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if job_type.startswith("device_"):
        kind = job_type.removeprefix("device_")
        groups_page = repository.list_groups(limit=100_000, offset=0)
        content = build_device_workbook(
            kind=kind,
            rows=extract_device_values(groups_page.get("items", []), kind),
            project=filters.get("project") or {},
            filters=filters,
        )
        return content, export_filename(job_type), media_type
    if job_type == "task_detail":
        return repository.build_task_detail_export(int(filters.get("task_id") or 0)), export_filename(job_type), media_type
    if job_type == "exception_meter":
        return repository.build_exception_meter_export(reviewer=str(filters.get("reviewer") or "")), export_filename(job_type), media_type
    if job_type == "project_outside":
        return repository.build_project_outside_export(), export_filename(job_type), media_type
    if job_type == "unmatched":
        page = repository.export_unmatched_records(query=str(filters.get("query") or ""), limit=100_000)
        return build_simple_table_workbook(title=job_type, rows=page.get("items", [])), export_filename(job_type), media_type
    if job_type in {"exception_missing_photo", "replacement", "barcode_review", "installer_kpi", "installer_daily_completion"}:
        return build_simple_table_workbook(title=job_type, rows=[]), export_filename(job_type), media_type
    raise ValueError(f"Unsupported export job type: {job_type}")


def request_background_export(repository: Any, *, job_type: str, filters: Mapping[str, Any], actor: str) -> dict[str, Any]:
    if job_type != "final_delivery":
        raise ValueError(f"Unsupported background export job type: {job_type}")
    try:
        package: LeasedDeliveryPackage = repository.request_final_delivery_export(
            task_id=filters.get("task_id"),
            terminal=str(filters.get("terminal") or ""),
            review_scope=str(filters.get("review_scope") or "reviewed"),
            requested_by=actor,
        )
    except DeliveryPackageNotReady as exc:
        return {
            "delivery_package_job_id": exc.job_id,
            "status": "pending",
            "content_path": "",
            "content_sha256": "",
            "size_bytes": None,
        }
    try:
        content = package.path.read_bytes()
        return {
            "delivery_package_job_id": "",
            "status": "succeeded",
            "content_path": str(package.path.resolve()),
            "content_sha256": file_sha256(content),
            "size_bytes": len(content),
        }
    finally:
        package.release()


def new_export_job_id() -> str:
    return str(uuid4())
