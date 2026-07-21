from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


MAX_PRIORITY_IMPORT_ROWS = 5000
PRIORITY_HEADERS = ("终端号", "优先施工")


class PriorityImportError(ValueError):
    pass


@dataclass
class PriorityImportRow:
    row_number: int
    terminal: str
    priority: bool | None
    status: str
    message: str = ""


def build_priority_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "优先施工"
    sheet.append(list(PRIORITY_HEADERS))
    for row_number in range(2, MAX_PRIORITY_IMPORT_ROWS + 2):
        sheet.cell(row_number, 1).number_format = "@"
    validation = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add("B2:B5001")
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def parse_priority_workbook(content: bytes) -> list[PriorityImportRow]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False, keep_links=False)
    except Exception as exc:
        raise PriorityImportError("Invalid workbook") from exc
    sheet = workbook.worksheets[0]
    headers = tuple(str(sheet.cell(1, column).value or "").strip() for column in (1, 2))
    if headers != PRIORITY_HEADERS:
        raise PriorityImportError("Invalid workbook headers")
    if sheet.max_row - 1 > MAX_PRIORITY_IMPORT_ROWS:
        raise PriorityImportError("Workbook exceeds 5000 data rows")

    rows: list[PriorityImportRow] = []
    by_terminal: dict[str, list[PriorityImportRow]] = {}
    for row_number in range(2, sheet.max_row + 1):
        terminal_cell, priority_cell = sheet.cell(row_number, 1), sheet.cell(row_number, 2)
        if terminal_cell.value is None and priority_cell.value is None:
            continue
        terminal = str(terminal_cell.value or "").strip()
        raw_priority = str(priority_cell.value or "").strip()
        if terminal_cell.data_type == "f" or priority_cell.data_type == "f":
            row = PriorityImportRow(row_number, terminal, None, "malformed", "Formula cells are not allowed")
        elif not terminal:
            row = PriorityImportRow(row_number, "", None, "malformed", "Terminal is required")
        elif raw_priority not in {"是", "否"}:
            row = PriorityImportRow(row_number, terminal, None, "malformed", "Priority must be 是 or 否")
        else:
            row = PriorityImportRow(row_number, terminal, raw_priority == "是", "valid")
        rows.append(row)
        if row.status != "malformed":
            by_terminal.setdefault(terminal, []).append(row)

    for terminal_rows in by_terminal.values():
        priorities = {row.priority for row in terminal_rows}
        if len(priorities) > 1:
            for row in terminal_rows:
                row.status = "conflict"
                row.message = "Conflicting priorities for terminal"
        elif len(terminal_rows) > 1:
            for row in terminal_rows[1:]:
                row.status = "duplicate"
                row.message = "Duplicate terminal"
    return rows
